"""Comptabilisation des droits de timbre par SIMTIS (13/10/2026)."""

import csv
import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import AccountingEntry, AuditLog, BankTransaction, ReconciliationMatch
from app.services import frais_service
from app.services.accounting_import_service import EntryLine, _mark_simtis_entries
from app.services.errors import ConflictError
from app.services.frais_service import (
    est_timbre,
    numero_piece,
    regrouper_timbres,
    type_commission,
    type_frais_tva,
    ventilation,
    ventiler_tva,
)
from tests.helpers import bearer, login, make_auth_user
from tests.test_reconciliation_api import account, company, ecriture, operation, statuses

JOUR = date(2026, 4, 23)
VERSEMENT = "DROIT DE TIMBRE SUR VERSEMENT"
REMISE = "DROITS DE TIMBRE SUR REMISE LCN N 2600114"


@pytest.fixture
def comptable(client, reference) -> dict[str, str]:
    make_auth_user(reference, "COMPTABLE", email="comptable@example.com")
    return bearer(login(client, "comptable@example.com"))


@pytest.fixture
def direction(client, reference) -> dict[str, str]:
    make_auth_user(reference, "DIRECTION", email="direction@example.com")
    return bearer(login(client, "direction@example.com"))


def bp(db, compte_timbre: str | None = "61670000"):
    compte = account(db, bank_code="BP")
    compte.journal_sage = "5243,5543"
    compte.journal_frais = "5543"
    compte.compte_comptable = "55410003"
    compte.compte_timbre = compte_timbre
    db.flush()
    return compte


def timbres_du_jour(db, compte, jour=JOUR):
    return [
        operation(db, compte, "-200", jour, REMISE),
        operation(db, compte, "-1", jour, VERSEMENT),
        operation(db, compte, "-1", jour, VERSEMENT),
    ]


class _Lecture:
    """Réponse de `GET /api/frais` réduite à une famille (le tableau unique les mêle)."""

    def __init__(self, status_code: int, data: dict):
        self.status_code = status_code
        self._data = data

    def json(self) -> dict:
        return self._data


def lire(client, headers, db, famille="timbres", **params):
    params = {"company_id": str(company(db).id), "from": "2026-04-01", "to": "2026-04-30", **params}
    response = client.get("/api/frais", params=params, headers=headers)
    if response.status_code != 200 or famille is None:
        return response
    data = response.json()
    return _Lecture(
        200,
        {
            **data,
            "propositions": [p for p in data["propositions"] if p["famille"] == famille],
            "ecritures": [e for e in data["ecritures"] if e["famille"] == famille],
        },
    )


def valider(client, headers, db, *jours, famille="timbres"):
    body = {
        "company_id": company(db).id,
        "jours": [
            {"famille": famille, "bank_account_id": compte.id, "date": str(jour)}
            for compte, jour in jours
        ],
    }
    return client.post("/api/frais/valider", json=body, headers=headers)


def exporter(client, headers, db, famille="timbres", **params):
    params = {"company_id": str(company(db).id), "from": "2026-04-01", "to": "2026-04-30", **params}
    # Un seul export pour toutes les familles (13/10/2026) : `famille` est ignorée
    return client.get("/api/frais/export", params=params, headers=headers)


def lignes_csv(response) -> list[tuple[str, ...]]:
    """Lignes du fichier CSV pour Sage (16/10/2026) : point-virgule, Windows-1252."""
    texte = response.content.decode("cp1252")
    return [tuple(row) for row in csv.reader(io.StringIO(texte), delimiter=";")]


def actions(db) -> list[str]:
    return list(db.scalars(select(AuditLog.action).order_by(AuditLog.id)))


# --- Règles pures ---------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("libelle", "attendu"),
    [
        (VERSEMENT, True),
        (REMISE, True),
        ("droit de timbre sur versement", True),
        ("COMMISSION", False),
        ("TIMBRE FISCAL", False),
        ("REMBOURSEMENT DROIT DE TIMBRE", False),
        (None, False),
    ],
)
def test_a_stamp_duty_is_read_from_the_label(libelle, attendu):
    assert est_timbre(libelle) is attendu


def _tx(id_, compte, jour, montant, libelle=VERSEMENT):
    return BankTransaction(
        id=id_,
        bank_account_id=compte,
        date_operation=jour,
        montant=Decimal(montant),
        libelle=libelle,
    )


def test_stamps_are_grouped_per_account_and_day_without_credits_or_days_in_sage():
    autre_jour = date(2026, 4, 24)
    groupes = regrouper_timbres(
        [
            _tx(1, 1, JOUR, "-1"),
            _tx(2, 1, JOUR, "-200", REMISE),
            _tx(3, 2, JOUR, "-1"),
            _tx(4, 1, JOUR, "1"),  # remboursement : pas un frais
            _tx(5, 1, JOUR, "-50", "COMMISSION"),
            _tx(6, 1, autre_jour, "-1"),
        ],
        {(1, autre_jour)},
    )

    assert [(g.bank_account_id, g.jour, g.total) for g in groupes] == [
        (1, JOUR, Decimal("201")),
        (2, JOUR, Decimal("1")),
    ]
    assert [tx.id for tx in groupes[0].operations] == [1, 2]


# --- Préparer -------------------------------------------------------------------------------------


def test_one_entry_is_proposed_per_day_and_nothing_is_written(client, comptable, db):
    compte = bp(db)
    timbres_du_jour(db, compte)
    operation(db, compte, "-1", date(2026, 4, 24), VERSEMENT)
    operation(db, compte, "-30", JOUR, "COMMISSION")

    response = lire(client, comptable, db)

    assert response.status_code == 200
    body = response.json()
    assert [(p["date"], p["total"], p["nb_operations"]) for p in body["propositions"]] == [
        ("2026-04-24", "1.00", 1),
        ("2026-04-23", "202.00", 3),
    ]
    assert body["propositions"][0]["libelle"] == "DT D'ENREGISTREMENT"
    # Tableau unique : le total compte aussi la COMMISSION de 30 du même jour
    assert body["total_propositions"] == "233.00"
    assert body["ecritures"] == []
    assert db.scalar(select(AccountingEntry).filter_by(origine="SIMTIS")) is None


def test_a_day_already_booked_in_sage_or_a_matched_stamp_is_not_proposed(client, comptable, db):
    compte = bp(db)
    timbres_du_jour(db, compte)
    ecriture(db, compte, "202", JOUR, "DT D'ENREGISTREMENT")
    autre = operation(db, compte, "-1", date(2026, 4, 24), VERSEMENT)
    autre.statut = "Rapprochée"
    db.flush()

    assert lire(client, comptable, db).json()["propositions"] == []


# --- Valider --------------------------------------------------------------------------------------


def test_validating_creates_the_entry_and_matches_it_at_once(client, comptable, db):
    compte = bp(db)
    timbres = timbres_du_jour(db, compte)

    response = valider(client, comptable, db, (compte, JOUR))

    assert response.status_code == 200, response.text
    code = response.json()["codes"][0]
    # N° pièce BANQUE AAMMNNN (16/10/2026) : premier frais BP du mois
    assert code == f"BP{JOUR:%y%m}001"
    entry = db.scalar(select(AccountingEntry).filter_by(origine="SIMTIS"))
    assert (entry.journal, entry.compte, entry.libelle, entry.numero_piece) == (
        "5543",
        "55410003",
        "DT D'ENREGISTREMENT",
        code,
    )
    assert (entry.debit, entry.credit, entry.montant) == (
        Decimal("0.00"),
        Decimal("202.00"),
        Decimal("202.00"),
    )
    match = db.scalar(select(ReconciliationMatch).filter_by(origine="Frais"))
    assert (match.statut, match.origine, match.type) == ("Validée", "Frais", "N-1")
    assert statuses(db, *timbres, entry) == ["Rapprochée"] * 4
    assert "comptabilisation_timbres" in actions(db)

    body = lire(client, comptable, db).json()
    assert body["propositions"] == []
    assert [
        (e["code"], e["numero_piece"], e["nb_operations"], e["exportee_le"])
        for e in body["ecritures"]
    ] == [(match.code, code, 3, None)]


def test_validation_is_all_or_nothing_and_recomputed_by_the_server(client, comptable, db):
    compte = bp(db)
    timbres_du_jour(db, compte)

    response = valider(client, comptable, db, (compte, JOUR), (compte, date(2026, 4, 25)))

    assert response.status_code == 409
    assert "plus de droits de timbre" in response.json()["detail"]
    assert db.scalar(select(AccountingEntry).filter_by(origine="SIMTIS")) is None


def test_validation_needs_the_bank_account_settings(client, comptable, db):
    compte = bp(db)
    compte.journal_frais = None  # deux journaux Sage : lequel ?
    db.flush()
    timbres_du_jour(db, compte)

    response = valider(client, comptable, db, (compte, JOUR))

    assert response.status_code == 409
    assert "journal des frais" in response.json()["detail"]


def test_preparing_needs_the_permission(client, direction, db):
    compte = bp(db)
    timbres_du_jour(db, compte)

    assert lire(client, direction, db).status_code == 200
    assert valider(client, direction, db, (compte, JOUR)).status_code == 403
    assert exporter(client, direction, db).status_code == 403


# --- Export pour Sage ------------------------------------------------------------------------------


def test_export_writes_a_csv_credit_line_and_marks_the_entry(client, comptable, db):
    """16/10/2026 : un CSV au point-virgule, Référence vide ; 17/10/2026 : le timbre au crédit de
    61671, sans ligne banque."""
    compte = bp(db, compte_timbre=None)
    timbres_du_jour(db, compte)
    code = valider(client, comptable, db, (compte, JOUR)).json()["codes"][0]

    response = exporter(client, comptable, db)

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/csv")
    assert ".csv" in response.headers["content-disposition"]
    rows = lignes_csv(response)
    # Format d'import paramétrable de Sage (15/10/2026)
    assert rows[0] == (
        "* Date de pièce",
        "N° pièce",
        "* Code journal",
        "* N° compte général",
        "N° compte tiers",
        "Libellé écriture",
        "Montant",
        "Sens",
        "Mode de règlement",
        "Date d'échéance",
        "Référence",
        "Numéro facture",
        "Montant devise",
        "Cours de la devise",
        "Code ISO devise",
    )
    vides = ("",) * 7
    assert rows[1:] == [
        ("23/04/2026", code, "5543", "61671", "", "DT D'ENREGISTREMENT", "202,00", "Crédit")
        + vides,
    ]
    entry = db.scalar(select(AccountingEntry).filter_by(origine="SIMTIS"))
    db.refresh(entry)
    assert entry.exportee_le is not None
    # Déjà exportée : plus rien de neuf, sauf à demander un nouvel export
    assert exporter(client, comptable, db).status_code == 409
    assert exporter(client, comptable, db, inclure_exportees="true").status_code == 200


# --- Suppression et annulation ---------------------------------------------------------------------


def _supprimer(client, headers, entry_id, motif="Erreur de saisie"):
    return client.request(
        "DELETE", f"/api/frais/ecritures/{entry_id}", json={"motif": motif}, headers=headers
    )


def test_an_entry_not_yet_exported_can_be_deleted(client, comptable, db):
    compte = bp(db)
    timbres = timbres_du_jour(db, compte)
    valider(client, comptable, db, (compte, JOUR))
    entry_id = db.scalar(select(AccountingEntry.id).filter_by(origine="SIMTIS"))

    assert _supprimer(client, comptable, entry_id, motif=" ").status_code == 409
    response = _supprimer(client, comptable, entry_id)

    assert response.status_code == 204
    db.expire_all()
    assert db.get(AccountingEntry, entry_id) is None
    assert statuses(db, *timbres) == ["Non rapprochée"] * 3
    assert "suppression_ecriture_frais" in actions(db)
    assert len(lire(client, comptable, db).json()["propositions"]) == 1


def test_an_exported_entry_cannot_be_deleted_nor_cancelled(client, comptable, db):
    compte = bp(db)
    timbres_du_jour(db, compte)
    valider(client, comptable, db, (compte, JOUR))
    match = db.scalar(select(ReconciliationMatch).filter_by(origine="Frais"))
    entry_id = db.scalar(select(AccountingEntry.id).filter_by(origine="SIMTIS"))

    annulation = client.request(
        "DELETE",
        f"/api/reconciliation/matches/{match.id}",
        json={"motif": "test"},
        headers=comptable,
    )
    assert annulation.status_code == 409
    assert "comptabilisée par SIMTIS" in annulation.json()["detail"]

    exporter(client, comptable, db)
    response = _supprimer(client, comptable, entry_id)

    assert response.status_code == 409
    assert "déjà été exportée" in response.json()["detail"]


# --- Retour par l'export Sage ---------------------------------------------------------------------


def test_a_simtis_entry_coming_back_from_sage_is_a_duplicate(client, comptable, db):
    compte = bp(db)
    timbres_du_jour(db, compte)
    code = valider(client, comptable, db, (compte, JOUR)).json()["codes"][0]

    def ligne(numero, libelle, montant="202.00", piece=None):
        return EntryLine(
            numero=numero,
            statut="Valide",
            motifs=[],
            date_ecriture=JOUR,
            libelle=libelle,
            montant=Decimal(montant),
            numero_piece=piece,
            bank_account_id=compte.id,
        )

    par_libelle = ligne(2, "DT D'ENREGISTREMENT")
    autre_montant = ligne(3, "DT D'ENREGISTREMENT", "15.00")
    _mark_simtis_entries(db, company(db).id, [par_libelle, autre_montant])
    assert par_libelle.statut == "Doublon"
    assert code in par_libelle.motifs[0]
    assert autre_montant.statut == "Valide"

    par_piece = ligne(4, "FRAIS DIVERS", piece=code.lower())
    _mark_simtis_entries(db, company(db).id, [par_piece])
    assert par_piece.statut == "Doublon"


def test_validated_entries_are_listed_with_their_code(db, comptable, client):
    compte = bp(db)
    timbres_du_jour(db, compte)
    valider(client, comptable, db, (compte, JOUR))

    items = frais_service.ecritures(db, company(db).id)

    assert [(item.bank_code, item.nb_operations) for item in items] == [("BP", 3)]


def test_the_reconciliation_screens_show_a_fees_match(client, comptable, db):
    """Une correspondance d'origine « Frais » se lit dans les volets et l'historique."""
    compte = bp(db)
    timbres_du_jour(db, compte)
    valider(client, comptable, db, (compte, JOUR))
    params = {"company_id": str(company(db).id), "du": "2026-04-01", "au": "2026-04-30"}

    for path in ("transactions", "history"):
        response = client.get(f"/api/reconciliation/{path}", params=params, headers=comptable)
        assert response.status_code == 200, (path, response.text)


# --- Commissions (13/10/2026) ---------------------------------------------------------------------

COM_EFFET = "COMMISSION EFFET N 3600839 REMIS PAR SGMB EN FAVEUR DE X"
TVA_EFFET = "TVA COM EFFET N 3600839 REMIS PAR SGMB EN FAVEUR DE X"


@pytest.mark.parametrize(
    ("libelle", "attendu"),
    [
        ("COMMISSION", "commission"),
        (" commission ", "commission"),
        ("TAXE SUR VALEUR AJOUTEE", "tva"),
        (COM_EFFET, "commission_effet"),
        (TVA_EFFET, "tva_commission_effet"),
        # Restent manuels
        # TVA comprise, dans la COM du jour depuis le 14/10/2026
        ("COMMISSION VIR. RTGS EMIS VERS BMCE EN FAVEUR DE X", "commission_rtgs"),
        ("COMMISSION REGLEMENT PRELEVEMENT EN FAVEUR DE X", "commission_prelevement"),
        ("COMMISSION COTISATIONS CNSS", None),
        ("TVA COM VIR. RTGS", None),
        (VERSEMENT, None),
    ],
)
def test_only_the_four_commission_labels_are_taken(libelle, attendu):
    assert type_commission(libelle) == attendu


def commissions(db, compte_tva: str | None = "34552000"):
    compte = bp(db)
    compte.compte_commission = "61470000"
    compte.compte_commission_effet = "61471000"
    compte.compte_tva = compte_tva
    db.flush()
    lignes = [
        operation(db, compte, "-20", JOUR, "COMMISSION"),
        operation(db, compte, "-30", JOUR, "COMMISSION"),
        operation(db, compte, "-5", JOUR, "TAXE SUR VALEUR AJOUTEE"),
        operation(db, compte, "-10", JOUR, COM_EFFET),
        operation(db, compte, "-1", JOUR, TVA_EFFET),
        operation(db, compte, "-165", JOUR, "COMMISSION COTISATIONS CNSS"),
    ]
    return compte, lignes


def test_one_com_entry_per_day_with_a_total_per_type(client, comptable, db):
    commissions(db)

    body = lire(client, comptable, db, "commissions").json()

    assert len(body["propositions"]) == 1
    item = body["propositions"][0]
    assert (item["libelle"], item["total"], item["nb_operations"]) == ("COM", "66.00", 5)
    assert [(d["type"], d["nb_operations"], d["total"]) for d in item["details"]] == [
        ("commission", 2, "50.00"),
        ("tva", 1, "5.00"),
        ("commission_effet", 1, "10.00"),
        ("tva_commission_effet", 1, "1.00"),
    ]
    # Les timbres ne voient pas les commissions
    assert lire(client, comptable, db).json()["propositions"] == []


def test_a_day_with_a_sage_com_entry_is_not_proposed(client, comptable, db):
    compte, _ = commissions(db)
    ecriture(db, compte, "231", JOUR, "COM")

    assert lire(client, comptable, db, "commissions").json()["propositions"] == []


def test_validating_commissions_matches_them_and_export_is_balanced(client, comptable, db):
    compte, lignes = commissions(db)

    response = valider(client, comptable, db, (compte, JOUR), famille="commissions")

    assert response.status_code == 200, response.text
    code = response.json()["codes"][0]
    entry = db.scalar(select(AccountingEntry).filter_by(origine="SIMTIS", libelle="COM"))
    assert (entry.montant, entry.numero_piece) == (Decimal("66.00"), code)
    assert statuses(db, *lignes) == ["Rapprochée"] * 5 + ["Non rapprochée"]
    assert "comptabilisation_commissions" in actions(db)
    # Les écritures de timbres ne listent pas la COM
    assert lire(client, comptable, db).json()["ecritures"] == []

    export = exporter(client, comptable, db, "commissions")
    assert export.status_code == 200, export.text
    rows = lignes_csv(export)
    # Une ligne par type au crédit de son compte, sans ligne banque (17/10/2026)
    assert [(r[3], r[1], r[5], r[7], r[6]) for r in rows[1:]] == [
        ("61714000", code, "COM", "Crédit", "50,00"),
        ("34552142", code, "TVA COM", "Crédit", "5,00"),
        ("61714000", code, "COM", "Crédit", "10,00"),
        ("34552142", code, "TVA COM", "Crédit", "1,00"),
    ]


def test_a_simtis_com_entry_coming_back_from_sage_is_a_duplicate(client, comptable, db):
    compte, _ = commissions(db)
    valider(client, comptable, db, (compte, JOUR), famille="commissions")
    line = EntryLine(
        numero=2,
        statut="Valide",
        motifs=[],
        date_ecriture=JOUR,
        libelle="COM",
        montant=Decimal("66.00"),
        bank_account_id=compte.id,
    )

    _mark_simtis_entries(db, company(db).id, [line])

    assert line.statut == "Doublon"


def test_an_unknown_family_is_refused(client, comptable, db):
    compte = bp(db)
    assert valider(client, comptable, db, (compte, JOUR), famille="agios").status_code == 422


# --- Format du grand livre, export unique (13/10/2026) --------------------------------------------


def test_lines_follow_the_general_ledger_with_one_line_per_type(client, comptable, db):
    compte, _ = commissions(db)
    timbres_du_jour(db, compte)

    tableau = lire(client, comptable, db, None).json()

    # Un seul tableau : pour un même jour, les timbres puis la COM
    assert [p["famille"] for p in tableau["propositions"]] == ["timbres", "commissions"]
    timbre, com = tableau["propositions"]
    assert com["journal"] == "5543"
    assert [
        (ligne["type"], ligne["compte"], ligne["debit"], ligne["credit"]) for ligne in com["lignes"]
    ] == [
        ("commission", "61714000", None, "50.00"),
        ("tva", "34552142", None, "5.00"),
        ("commission_effet", "61714000", None, "10.00"),
        ("tva_commission_effet", "34552142", None, "1.00"),
    ]
    assert [(ligne["type"], ligne["debit"], ligne["credit"]) for ligne in timbre["lignes"]] == [
        ("timbre", None, "202.00"),
    ]

    # Les deux familles se valident ensemble, toutes ou aucune
    body = {
        "company_id": company(db).id,
        "jours": [
            {"famille": "timbres", "bank_account_id": compte.id, "date": str(JOUR)},
            {"famille": "commissions", "bank_account_id": compte.id, "date": str(JOUR)},
        ],
    }
    response = client.post("/api/frais/valider", json=body, headers=comptable)
    assert response.status_code == 200, response.text
    assert response.json()["nb_ecritures"] == 2
    ecrites = lire(client, comptable, db, None).json()["ecritures"]
    assert sorted(e["famille"] for e in ecrites) == ["commissions", "timbres"]
    com_ecrite = next(e for e in ecrites if e["famille"] == "commissions")
    assert [ligne["type"] for ligne in com_ecrite["lignes"]] == [
        "commission",
        "tva",
        "commission_effet",
        "tva_commission_effet",
    ]


def test_one_export_file_holds_stamps_and_commissions(client, comptable, db):
    compte, _ = commissions(db)
    timbres_du_jour(db, compte, date(2026, 4, 22))
    valider(client, comptable, db, (compte, JOUR), famille="commissions")
    valider(client, comptable, db, (compte, date(2026, 4, 22)))

    response = exporter(client, comptable, db)

    assert response.status_code == 200, response.text
    rows = lignes_csv(response)
    assert [r[5] for r in rows[1:]] == ["DT D'ENREGISTREMENT", "COM", "TVA COM", "COM", "TVA COM"]
    assert all(r[7] == "Crédit" for r in rows[1:])
    assert sum(Decimal(r[6].replace(",", ".")) for r in rows[1:]) == 268
    assert exporter(client, comptable, db).status_code == 409


# --- Comptabilisation manuelle (13/10/2026) ------------------------------------------------------


@pytest.mark.parametrize(
    ("total", "taux", "ht", "tva"),
    [
        ("165.00", "10", "150.00", "15.00"),
        ("11.00", "10", "10.00", "1.00"),
        ("100.00", "10", "90.91", "9.09"),
        ("0.05", "20", "0.04", "0.01"),
    ],
)
def test_vat_is_split_from_the_total_and_always_adds_up(total, taux, ht, tva):
    resultat = ventiler_tva(Decimal(total), Decimal(taux))

    assert resultat == (Decimal(ht), Decimal(tva))
    assert sum(resultat) == Decimal(total)


# Un frais qui reste à passer à la main (la RTGS est automatique depuis le 14/10/2026)
RTGS = "FRAIS DE TENUE DE COMPTE"


def manuelle(client, headers, db, operations, **over):
    body = {
        "company_id": company(db).id,
        "transaction_ids": [tx.id for tx in operations],
        "libelle": "COM RTGS",
        **over,
    }
    return client.post("/api/frais/manuelle", json=body, headers=headers)


def test_other_fees_are_listed_apart_from_the_automatic_ones(client, comptable, db):
    compte = bp(db)
    timbres_du_jour(db, compte)
    rtgs = operation(db, compte, "-165", JOUR, RTGS)
    agios = operation(db, compte, "-604.71", JOUR, "AGIOS BANCAIRES")

    body = lire(client, comptable, db, None).json()

    assert {item["id"] for item in body["autres_frais"]} == {rtgs.id, agios.id}


def test_an_automatic_fee_is_never_booked_by_hand(client, comptable, db):
    """Jour où Sage a déjà sa COM : la commission n'est pas proposée en automatique, mais elle ne se
    passe pas non plus à la main (elle se rapproche de l'écriture Sage)."""
    compte = bp(db)
    commission = operation(db, compte, "-20", JOUR, "COMMISSION")
    # Autre montant : la COM de Sage ne couvre pas cette commission, mais le jour est déjà passé
    ecriture(db, compte, "35", JOUR, "COM")

    body = lire(client, comptable, db, None).json()
    response = manuelle(client, comptable, db, [commission])

    assert body["propositions"] == [] and body["autres_frais"] == []
    assert response.status_code == 409
    assert "automatique" in response.json()["detail"]


def test_a_manual_entry_with_vat_is_matched_and_exported_on_two_lines(client, comptable, db):
    compte = bp(db)
    rtgs = operation(db, compte, "-165", JOUR, RTGS)

    response = manuelle(client, comptable, db, [rtgs], taux_tva="10")

    assert response.status_code == 200, response.text
    code = response.json()["codes"][0]
    entry = db.scalar(select(AccountingEntry).filter_by(famille_frais="manuelle"))
    assert (
        entry.libelle,
        entry.montant,
        entry.montant_tva,
        entry.taux_tva,
        entry.date_ecriture,
    ) == (
        "COM RTGS",
        Decimal("165.00"),
        Decimal("15.00"),
        Decimal("10.00"),
        JOUR,
    )
    assert statuses(db, rtgs) == ["Rapprochée"]
    assert "comptabilisation_manuelle" in actions(db)
    ecrite = next(
        e
        for e in lire(client, comptable, db, None).json()["ecritures"]
        if e["famille"] == "manuelle"
    )
    assert [
        (ligne["type_libelle"], ligne["debit"], ligne["credit"]) for ligne in ecrite["lignes"]
    ] == [
        ("Frais", None, "150.00"),
        ("TVA", None, "15.00"),
    ]

    rows = lignes_csv(exporter(client, comptable, db))
    assert [(r[1], r[3], r[5], r[7], r[6]) for r in rows[1:]] == [
        (code, "61714000", "COM RTGS", "Crédit", "150,00"),
        (code, "34552142", "TVA COM RTGS", "Crédit", "15,00"),
    ]


def test_a_manual_entry_without_vat_has_one_line_and_a_refund_goes_to_the_debit(
    client, comptable, db
):
    compte = bp(db)
    remboursement = operation(db, compte, "151110.79", JOUR, "AGIOS BANCAIRES")

    response = manuelle(client, comptable, db, [remboursement], libelle="agios  rembourses")

    assert response.status_code == 200, response.text
    entry = db.scalar(select(AccountingEntry).filter_by(famille_frais="manuelle"))
    assert (entry.libelle, entry.debit, entry.credit, entry.montant) == (
        "AGIOS REMBOURSES",
        Decimal("151110.79"),
        Decimal("0.00"),
        Decimal("-151110.79"),
    )
    ecrite = lire(client, comptable, db, None).json()["ecritures"][0]
    assert [(ligne["compte"], ligne["debit"], ligne["credit"]) for ligne in ecrite["lignes"]] == [
        ("63112000", "151110.79", None),
    ]
    # Sage : montant positif ; un remboursement d'agios BP au débit de 63112000 (17/10/2026)
    rows = lignes_csv(exporter(client, comptable, db))
    assert [(r[3], r[5], r[6], r[7]) for r in rows[1:]] == [
        ("63112000", "AGIOS REMBOURSES", "151110,79", "Débit"),
    ]


def test_fees_of_a_currency_account_are_not_exported_yet(client, comptable, db):
    compte = bp(db)
    timbres_du_jour(db, compte)
    valider(client, comptable, db, (compte, JOUR))
    compte.devise = "EUR"
    db.flush()

    response = exporter(client, comptable, db)

    assert response.status_code == 409
    assert "BP EUR" in response.json()["detail"]
    assert db.scalar(select(AccountingEntry).filter_by(origine="SIMTIS")).exportee_le is None


def test_a_manual_vat_amount_is_kept_and_must_stay_below_the_total(client, comptable, db):
    compte = bp(db)
    rtgs = operation(db, compte, "-165", JOUR, RTGS)

    trop = manuelle(client, comptable, db, [rtgs], montant_tva="165.00")
    ok = manuelle(client, comptable, db, [rtgs], montant_tva="14.99")

    assert trop.status_code == 409
    assert "inférieure au total" in trop.json()["detail"]
    assert ok.status_code == 200, ok.text
    entry = db.scalar(select(AccountingEntry).filter_by(famille_frais="manuelle"))
    assert (entry.montant_tva, entry.taux_tva) == (Decimal("14.99"), None)


def test_a_manual_entry_needs_one_account_free_fees_a_label_and_a_valid_date(client, comptable, db):
    compte = bp(db)
    autre = account(db, bank_code="AWB")
    autre.journal_sage = "BQ9"
    autre.compte_comptable = "5141"
    db.flush()
    rtgs = operation(db, compte, "-165", JOUR, RTGS)
    ailleurs = operation(db, autre, "-50", JOUR, "AGIOS BANCAIRES")
    rapprochee = operation(db, compte, "-30", JOUR, "AGIOS BANCAIRES")
    rapprochee.statut = "Rapprochée"
    virement = operation(db, compte, "-500", JOUR, "VIR CLIENT ATLAS")
    db.flush()

    assert "même compte" in manuelle(client, comptable, db, [rtgs, ailleurs]).json()["detail"]
    for operations in ([rapprochee], [virement]):
        assert (
            "plus à comptabiliser" in manuelle(client, comptable, db, operations).json()["detail"]
        )
    assert manuelle(client, comptable, db, [rtgs], libelle="   ").status_code == 409
    assert "date" in manuelle(client, comptable, db, [rtgs], date="2026-04-01").json()["detail"]
    assert db.scalar(select(AccountingEntry).filter_by(famille_frais="manuelle")) is None


def test_a_manual_entry_can_be_deleted_before_export(client, comptable, db):
    compte = bp(db)
    rtgs = operation(db, compte, "-165", JOUR, RTGS)
    manuelle(client, comptable, db, [rtgs])
    entry_id = db.scalar(select(AccountingEntry.id).filter_by(famille_frais="manuelle"))

    response = _supprimer(client, comptable, entry_id)

    assert response.status_code == 204
    assert statuses(db, rtgs) == ["Non rapprochée"]


def test_manual_booking_needs_the_permission(client, direction, db):
    compte = bp(db)
    rtgs = operation(db, compte, "-165", JOUR, RTGS)

    assert manuelle(client, direction, db, [rtgs]).status_code == 403


# --- Frais déjà inclus dans une écriture Sage (13/10/2026) ---------------------------------------


def virement_et_frais(db, compte, jour=JOUR):
    virement = operation(db, compte, "-10000", jour, "VIREMENT EMIS FOURNISSEUR X")
    commission = operation(db, compte, "-50", jour, "COMMISSION")
    tva = operation(db, compte, "-5", jour, "TAXE SUR VALEUR AJOUTEE")
    return virement, commission, tva


def test_fees_inside_a_sage_entry_are_neither_proposed_nor_booked(client, comptable, db):
    """Seul le groupe « opération principale + frais » couvre un frais (15/10/2026) : un frais
    seul face à une écriture du même montant se comptabilise dans SIMTIS."""
    compte = bp(db)
    _, commission, tva = virement_et_frais(db, compte)
    agios = operation(db, compte, "-30", JOUR, "AGIOS BANCAIRES")
    ecriture(db, compte, "10055", JOUR, "VIR FOURNISSEUR X")
    ecriture(db, compte, "30", date(2026, 4, 25), "AGIOS AVRIL")

    body = lire(client, comptable, db, None).json()
    automatique = valider(client, comptable, db, (compte, JOUR), famille="commissions")

    assert body["propositions"] == []
    inclus = {item["id"]: item["ecriture"]["libelle"] for item in body["frais_inclus"]}
    assert inclus == {commission.id: "VIR FOURNISSEUR X", tva.id: "VIR FOURNISSEUR X"}
    assert [item["id"] for item in body["autres_frais"]] == [agios.id]
    assert automatique.status_code == 409
    assert manuelle(client, comptable, db, [agios]).status_code == 200


def test_a_rejected_pair_or_an_unrelated_entry_does_not_block_a_fee(client, comptable, db):
    from datetime import UTC, datetime

    from app.models import ReconciliationMatch, ReconciliationMatchItem, User

    compte = bp(db)
    autre = account(db, bank_code="AWB")
    agios = operation(db, compte, "-30", JOUR, "AGIOS BANCAIRES")
    rejetee = ecriture(db, compte, "30", JOUR, "AGIOS AVRIL")
    ecriture(db, autre, "30", JOUR, "AGIOS AWB")  # autre compte
    ecriture(db, compte, "30", date(2026, 4, 28), "AGIOS MAI")  # 5 jours d'écart
    db.add(
        ReconciliationMatch(
            company_id=compte.company_id,
            type="1-1",
            statut="Rejetée",
            decide_par_id=db.scalar(select(User.id).filter_by(email="comptable@example.com")),
            decide_le=datetime.now(UTC),
            items=[
                ReconciliationMatchItem(bank_transaction_id=agios.id, montant_affecte=30),
                ReconciliationMatchItem(accounting_entry_id=rejetee.id, montant_affecte=30),
            ],
        )
    )
    db.flush()

    body = lire(client, comptable, db, None).json()

    assert body["frais_inclus"] == []
    assert [item["id"] for item in body["autres_frais"]] == [agios.id]


# --- Écriture « FRAIS » : frais à TVA comprise (14/10/2026) ---------------------------------------


@pytest.mark.parametrize(
    ("libelle", "attendu"),
    [
        # Les commissions à TVA comprise vont dans la COM du jour, pas dans « FRAIS »
        ("COMMISSION VIR. RTGS EMIS VERS BMCE EN FAVEUR DE X", None),
        ("COMMISSION BANCAIRE", None),
        ("FRAIS RECUPERES PAR LA BANQUE", "frais_recuperes"),
        ("FRAIS RÉCUPÉRÉS PAR LA BANQUE", "frais_recuperes"),
        ("FRAIS DE TENUE DE COMPTE", None),
        ("AGIOS BANCAIRES", None),
        ("COMMISSION", None),
    ],
)
def test_fees_with_vat_inside_are_recognised(libelle, attendu):
    assert type_frais_tva(libelle) == attendu


@pytest.mark.parametrize(
    ("libelle", "attendu"),
    [
        ("COMMISSION VIR. RTGS EMIS VERS BMCE EN FAVEUR DE X", "commission_rtgs"),
        ("COMMISSION REGLEMENT PRELEVEMENT EN FAVEUR DE MAROC LEASING", "commission_prelevement"),
        ("COMMISSION BANCAIRE", "commission_bancaire"),
        ("COMMISSION TENUE DE COMPTE", "commission_tenue"),
        ("COMMISSION", "commission"),
        ("COMMISSION COTISATIONS CNSS", None),
        ("FRAIS RECUPERES PAR LA BANQUE", None),
    ],
)
def test_commissions_with_vat_inside_go_to_the_com_entry(libelle, attendu):
    assert type_commission(libelle) == attendu


@pytest.mark.parametrize(
    ("code", "montant", "parts"),
    [
        (
            "commission_rtgs",
            "165.00",
            [("commission_rtgs", "150.00"), ("commission_rtgs_tva", "15.00")],
        ),
        (
            "commission_prelevement",
            "11.00",
            [("commission_prelevement", "10.00"), ("commission_prelevement_tva", "1.00")],
        ),
        (
            "frais_recuperes",
            "355.00",
            [("frais_recuperes", "322.73"), ("frais_recuperes_tva", "32.27")],
        ),
        (
            "frais_recuperes",
            "77.50",
            [
                ("frais_recuperes", "25.00"),
                ("frais_recuperes_tva", "2.50"),
                ("frais_recuperes_timbre", "50.00"),
            ],
        ),
        ("commission", "20.00", [("commission", "20.00")]),
    ],
)
def test_vat_and_stamp_are_split_from_each_fee(code, montant, parts):
    resultat = ventilation(code, Decimal(montant))

    assert resultat == [(nom, Decimal(valeur)) for nom, valeur in parts]
    assert sum(valeur for _, valeur in resultat) == Decimal(montant)


def frais_tva_du_jour(db, compte, jour=JOUR):
    return [
        operation(db, compte, "-165", jour, "COMMISSION VIR. RTGS EMIS VERS BMCE EN FAVEUR DE X"),
        operation(db, compte, "-77.50", jour, "FRAIS RECUPERES PAR LA BANQUE"),
        operation(db, compte, "-11", jour, "COMMISSION REGLEMENT PRELEVEMENT EN FAVEUR DE X"),
        # Un remboursement reste à la main
        operation(db, compte, "73558.97", jour, "COMMISSION BANCAIRE"),
    ]


def test_commissions_go_to_com_and_frais_recuperes_to_frais_with_computed_vat(
    client, comptable, db
):
    compte = bp(db)
    rtgs, recuperes, prelevement, remboursement = frais_tva_du_jour(db, compte)

    tous = lire(client, comptable, db, None).json()
    com, frais = (lire(client, comptable, db, nom).json() for nom in ("commissions", "frais"))

    # COMMISSION… → « COM » ; FRAIS… → « FRAIS » (14/10/2026)
    [ecriture_com] = com["propositions"]
    assert (ecriture_com["libelle"], ecriture_com["total"]) == ("COM", "176.00")
    assert [(ligne["type_libelle"], ligne["credit"]) for ligne in ecriture_com["lignes"]] == [
        ("Commission RTGS", "150.00"),
        ("TVA commission RTGS", "15.00"),
        ("Commission prélèvement", "10.00"),
        ("TVA commission prélèvement", "1.00"),
    ]
    [ecriture_frais] = frais["propositions"]
    assert (ecriture_frais["libelle"], ecriture_frais["total"]) == ("FRAIS", "77.50")
    assert [
        (ligne["type_libelle"], ligne["compte"], ligne["credit"])
        for ligne in ecriture_frais["lignes"]
    ] == [
        ("Frais récupérés", "61473000", "25.00"),
        ("TVA frais récupérés", "34552142", "2.50"),
        ("Timbre frais récupérés", "61671", "50.00"),
    ]
    assert [a["id"] for a in tous["autres_frais"]] == [remboursement.id]

    body = {
        "company_id": company(db).id,
        "jours": [
            {"famille": nom, "bank_account_id": compte.id, "date": str(JOUR)}
            for nom in ("commissions", "frais")
        ],
    }
    response = client.post("/api/frais/valider", json=body, headers=comptable)
    assert response.status_code == 200, response.text
    code_com, code_frais = response.json()["codes"]
    assert statuses(db, rtgs, recuperes, prelevement, remboursement) == [
        "Rapprochée",
        "Rapprochée",
        "Rapprochée",
        "Non rapprochée",
    ]
    assert "comptabilisation_frais" in actions(db)

    rows = lignes_csv(exporter(client, comptable, db))
    assert [(r[1], r[3], r[5], r[6], r[7]) for r in rows[1:]] == [
        (code_com, "61473000", "COM", "150,00", "Crédit"),
        (code_com, "34552142", "TVA COM", "15,00", "Crédit"),
        (code_com, "61473000", "COM", "10,00", "Crédit"),
        (code_com, "34552142", "TVA COM", "1,00", "Crédit"),
        (code_frais, "61473000", "FRAIS", "25,00", "Crédit"),
        (code_frais, "34552142", "TVA FRAIS", "2,50", "Crédit"),
        (code_frais, "61671", "FRAIS", "50,00", "Crédit"),
    ]
    # Même libellé à l'écran (14/10/2026)
    assert [ligne["libelle"] for ligne in ecriture_frais["lignes"]] == [
        "FRAIS",
        "TVA FRAIS",
        "FRAIS",
    ]


def test_a_day_with_a_sage_frais_entry_is_not_proposed(client, comptable, db):
    compte = bp(db)
    frais_tva_du_jour(db, compte)
    ecriture(db, compte, "999", JOUR, "FRAIS")

    assert lire(client, comptable, db, "frais").json()["propositions"] == []


# --- N° de pièce BANQUE AAMMNNN (16/10/2026) -----------------------------------------------------


def test_a_piece_number_is_bank_year_month_and_three_digits():
    assert numero_piece("BP", date(2026, 10, 3), 123) == "BP2610123"
    assert numero_piece("AWB", date(2027, 1, 31), 1) == "AWB2701001"
    with pytest.raises(ConflictError, match="épuisé"):
        numero_piece("BP", date(2026, 10, 3), 1000)


def test_piece_numbers_follow_each_other_within_a_month_and_restart_the_next(client, comptable, db):
    compte = bp(db)
    autre_jour = JOUR + timedelta(days=1)
    mois_suivant = date(JOUR.year, JOUR.month % 12 + 1, 1)
    for jour in (JOUR, autre_jour, mois_suivant):
        timbres_du_jour(db, compte, jour)

    pieces = [
        valider(client, comptable, db, (compte, jour)).json()["codes"][0]
        for jour in (JOUR, autre_jour, mois_suivant)
    ]

    assert pieces == [
        f"BP{JOUR:%y%m}001",
        f"BP{JOUR:%y%m}002",
        f"BP{mois_suivant:%y%m}001",
    ]


def test_each_bank_has_its_own_piece_counter(client, comptable, db):
    compte = bp(db)
    awb = account(db, bank_code="AWB")
    awb.journal_sage = "5141"
    awb.compte_comptable = "55410001"
    db.flush()
    timbres_du_jour(db, compte)
    timbres_du_jour(db, awb)

    response = valider(client, comptable, db, (compte, JOUR), (awb, JOUR))

    assert response.status_code == 200, response.text
    assert sorted(response.json()["codes"]) == [f"AWB{JOUR:%y%m}001", f"BP{JOUR:%y%m}001"]


# --- Propositions modifiées ou écartées (19/10/2026) ----------------------------------------------


def valider_modifie(client, headers, db, compte, **modif):
    body = {
        "company_id": company(db).id,
        "jours": [{"famille": "timbres", "bank_account_id": compte.id, "date": str(JOUR), **modif}],
    }
    return client.post("/api/frais/valider", json=body, headers=headers)


def test_a_proposal_is_validated_with_its_label_date_and_lines_changed(client, comptable, db):
    compte = bp(db)
    timbres_du_jour(db, compte)
    lignes = [
        {"compte": "61671", "libelle": "DT VERSEMENTS", "montant": "2.00", "sens": "Crédit"},
        {"compte": "61473000", "libelle": "DT REMISE LCN", "montant": "200.00", "sens": "Crédit"},
    ]

    response = valider_modifie(
        client,
        comptable,
        db,
        compte,
        libelle="dt avril",
        date_ecriture=str(JOUR + timedelta(days=2)),
        lignes=lignes,
    )

    assert response.status_code == 200, response.text
    assert response.json()["codes"] == [f"BP{JOUR:%y%m}001"]
    entry = db.scalar(select(AccountingEntry).filter_by(origine="SIMTIS"))
    assert (entry.libelle, entry.date_ecriture, entry.montant) == (
        "DT AVRIL",
        JOUR + timedelta(days=2),
        Decimal("202.00"),
    )
    [ecrite] = lire(client, comptable, db).json()["ecritures"]
    assert ecrite["modifiee"] is True
    assert [(g["compte"], g["libelle"], g["credit"]) for g in ecrite["lignes"]] == [
        ("61671", "DT VERSEMENTS", "2.00"),
        ("61473000", "DT REMISE LCN", "200.00"),
    ]
    rows = lignes_csv(exporter(client, comptable, db))
    assert [(r[0], r[3], r[5], r[6], r[7]) for r in rows[1:]] == [
        ("25/04/2026", "61671", "DT VERSEMENTS", "2,00", "Crédit"),
        ("25/04/2026", "61473000", "DT REMISE LCN", "200,00", "Crédit"),
    ]
    audit = db.scalar(select(AuditLog).where(AuditLog.action == "comptabilisation_timbres"))
    assert audit.nouvelle_valeur["modifiee"] is True


def test_changed_lines_must_make_the_total_and_the_date_stay_in_bounds(client, comptable, db):
    compte = bp(db)
    timbres_du_jour(db, compte)
    faux = [{"compte": "61671", "libelle": "DT", "montant": "200.00", "sens": "Crédit"}]

    total = valider_modifie(client, comptable, db, compte, lignes=faux)
    avant = valider_modifie(
        client, comptable, db, compte, date_ecriture=str(JOUR - timedelta(days=1))
    )

    assert total.status_code == 409
    assert "écart 2.00" in total.json()["detail"]
    assert avant.status_code == 409
    assert db.scalar(select(AccountingEntry).filter_by(origine="SIMTIS")) is None


def test_a_dismissed_day_goes_to_the_fees_booked_by_hand_until_proposed_again(
    client, comptable, direction, db
):
    compte = bp(db)
    timbres = timbres_du_jour(db, compte)
    jour = {"famille": "timbres", "bank_account_id": compte.id, "date": str(JOUR)}
    body = {"company_id": company(db).id, "jours": [jour], "motif": "Déjà payé autrement"}

    assert client.post("/api/frais/ecarter", json=body, headers=direction).status_code == 403
    response = client.post("/api/frais/ecarter", json=body, headers=comptable)

    assert response.status_code == 200, response.text
    tableau = lire(client, comptable, db, None).json()
    assert tableau["propositions"] == []
    assert {a["id"] for a in tableau["autres_frais"]} == {tx.id for tx in timbres}
    [ecarte] = tableau["ecartes"]
    assert (ecarte["famille"], ecarte["date"], ecarte["motif"]) == (
        "timbres",
        str(JOUR),
        "Déjà payé autrement",
    )
    # Deux fois : la journée n'est plus proposée
    assert client.post("/api/frais/ecarter", json=body, headers=comptable).status_code == 409
    assert manuelle(client, comptable, db, timbres, libelle="DT").status_code == 200

    # Reproposer : plus rien à proposer (les timbres sont passés à la main), la ligne disparaît
    retour = client.post(
        "/api/frais/reproposer",
        json={"company_id": company(db).id, "jours": [jour]},
        headers=comptable,
    )
    assert retour.status_code == 200
    assert lire(client, comptable, db, None).json()["ecartes"] == []
    assert "ecartement_frais" in actions(db) and "reproposition_frais" in actions(db)


def test_a_dismissed_day_comes_back_when_proposed_again(client, comptable, db):
    compte = bp(db)
    timbres_du_jour(db, compte)
    jour = {"famille": "timbres", "bank_account_id": compte.id, "date": str(JOUR)}
    client.post(
        "/api/frais/ecarter",
        json={"company_id": company(db).id, "jours": [jour]},
        headers=comptable,
    )

    client.post(
        "/api/frais/reproposer",
        json={"company_id": company(db).id, "jours": [jour]},
        headers=comptable,
    )

    [proposition] = lire(client, comptable, db).json()["propositions"]
    assert proposition["total"] == "202.00"


# --- Frais de plus de 200 DH : à la main et à justifier (19/10/2026) -----------------------------


def test_a_fee_above_200_leaves_the_automatic_proposal_but_a_stamp_does_not(client, comptable, db):
    compte = bp(db)
    petite = operation(db, compte, "-150", JOUR, "COMMISSION")
    grande = operation(db, compte, "-250", JOUR, "COMMISSION")
    timbre = operation(db, compte, "-300", JOUR, REMISE)

    tableau = lire(client, comptable, db, None).json()

    totaux = {p["famille"]: p["total"] for p in tableau["propositions"]}
    assert totaux == {"commissions": "150.00", "timbres": "300.00"}
    assert [a["id"] for a in tableau["autres_frais"]] == [grande.id]
    assert valider(client, comptable, db, (compte, JOUR), famille="commissions").status_code == 200
    assert statuses(db, petite, grande, timbre) == [
        "Rapprochée",
        "Non rapprochée",
        "Non rapprochée",
    ]
    # Elle se comptabilise à la main
    assert manuelle(client, comptable, db, [grande], libelle="COM").status_code == 200


def test_fees_above_200_are_exported_for_the_bank(client, comptable, direction, db):
    compte = bp(db)
    operation(db, compte, "-150", JOUR, "COMMISSION")
    grande = operation(db, compte, "-250", JOUR, "COMMISSION", reference="R1")
    operation(db, compte, "-300", JOUR, REMISE)
    agios = operation(db, compte, "-604.71", date(2026, 4, 24), "AGIOS BANCAIRES")
    manuelle(client, comptable, db, [agios], libelle="AGIOS")
    params = {"company_id": str(company(db).id), "from": "2026-04-01", "to": "2026-04-30"}

    response = client.get("/api/frais/justificatif", params=params, headers=direction)

    assert response.status_code == 200, response.text
    rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
    assert rows[0][:8] == (
        "Banque",
        "Compte",
        "Date d'opération",
        "Date de valeur",
        "Libellé",
        "Référence",
        "Montant",
        "Statut",
    )
    assert [(r[0], r[4], r[5], r[6], r[7]) for r in rows[1:3]] == [
        ("BP", "COMMISSION", "R1", 250, "À comptabiliser"),
        ("BP", "AGIOS BANCAIRES", None, 604.71, "Comptabilisé"),
    ]
    assert rows[2][8].startswith("BP2604")
    assert rows[3][0] == "Total BP" and rows[3][6] == 854.71
    assert rows[4][0] == "Total général"
    assert grande.id and "export_frais_banque" in actions(db)
