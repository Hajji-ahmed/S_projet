"""Constructeurs de données de test et vérification des rejets par la base."""

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from itertools import count

import pytest
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.models import (
    AccountingEntry,
    Bank,
    BankAccount,
    BankStatement,
    BankTransaction,
    Company,
    Currency,
    ForecastCategory,
    Role,
    User,
)

_sequence = count(1)


def assert_rejected(db: Session, message: str, *objects: object, sql: str | None = None) -> None:
    """Vérifie que la base REFUSE l'opération, avec une erreur qui cite `message` (nom de contrainte...).

    Les objets sont enregistrés, ou le SQL exécuté, dans un point de sauvegarde annulé au refus :
    la session reste utilisable et les données préparées par le test sont conservées.
    """
    with pytest.raises(IntegrityError) as excinfo:
        with db.begin_nested():
            if objects:
                db.add_all(objects)
                db.flush()
            if sql:
                db.execute(text(sql))
    assert message in str(excinfo.value.orig), f"attendu « {message} » dans : {excinfo.value.orig}"


def save(db: Session, *objects):
    """Enregistre les objets et retourne le premier (ou l'unique)."""
    db.add_all(objects)
    db.flush()
    return objects[0]


def make_currencies(db: Session) -> None:
    for code, libelle in (("MAD", "Dirham"), ("EUR", "Euro"), ("USD", "Dollar")):
        if db.get(Currency, code) is None:
            db.add(Currency(code=code, libelle=libelle))
    db.flush()


def build_company(**over) -> Company:
    n = next(_sequence)
    return Company(**{"code": f"C{n}", "nom": f"Société {n}", **over})


def build_bank(**over) -> Bank:
    n = next(_sequence)
    return Bank(**{"code": f"B{n}", "nom": f"Banque {n}", **over})


def build_account(company: Company, bank: Bank, **over) -> BankAccount:
    n = next(_sequence)
    return BankAccount(
        **{
            "company_id": company.id,
            "bank_id": bank.id,
            "libelle": "Compte test",
            "numero": f"N{n}",
            "devise": "MAD",
            **over,
        }
    )


def build_transaction(statement: BankStatement, **over) -> BankTransaction:
    """Transaction de 100 au crédit. Qui change débit ou crédit doit aussi fournir `montant`."""
    return BankTransaction(
        **{
            "statement_id": statement.id,
            "bank_account_id": statement.bank_account_id,
            "date_operation": date(2026, 9, 24),
            "libelle": "VIR CLIENT ABC",
            "debit": Decimal("0"),
            "credit": Decimal("100"),
            "montant": Decimal("100"),
            "hash_ligne": f"h{next(_sequence)}",
            **over,
        }
    )


def build_entry(company: Company, **over) -> AccountingEntry:
    """Écriture de 100 au crédit. Qui change débit ou crédit doit aussi fournir `montant`."""
    return AccountingEntry(
        **{
            "company_id": company.id,
            "date_ecriture": date(2026, 9, 24),
            "libelle": "REG FACT-458",
            "debit": Decimal("0"),
            "credit": Decimal("100"),
            "montant": Decimal("100"),
            "hash_ligne": f"e{next(_sequence)}",
            **over,
        }
    )


def build_user(**over) -> User:
    n = next(_sequence)
    return User(
        **{
            "nom": f"Utilisateur {n}",
            "email": f"u{n}@example.com",
            "mot_de_passe_hash": "x",
            **over,
        }
    )


@dataclass
class World:
    """Jeu de données minimal : une société, une banque, un compte, un relevé, un utilisateur."""

    company: Company
    bank: Bank
    account: BankAccount
    statement: BankStatement
    user: User
    category: ForecastCategory


def make_world(db: Session) -> World:
    make_currencies(db)
    company = save(db, build_company())
    bank = save(db, build_bank())
    account = save(db, build_account(company, bank))
    statement = save(db, BankStatement(bank_account_id=account.id))
    user = save(db, build_user())
    category = db.scalar(select(ForecastCategory).filter_by(code="TEST"))
    if category is None:
        category = save(db, ForecastCategory(code="TEST", libelle="Test"))
    return World(company, bank, account, statement, user, category)


# --- Authentification ----------------------------------------------------------------------------

TEST_PASSWORD = "Mot-De-Passe-De-Test-2026"
_test_password_hash: str | None = None


def test_password_hash() -> str:
    """Empreinte de TEST_PASSWORD, calculée une seule fois (argon2 est volontairement lent)."""
    global _test_password_hash
    if _test_password_hash is None:
        _test_password_hash = hash_password(TEST_PASSWORD)
    return _test_password_hash


def make_auth_user(db: Session, *role_codes: str, **over) -> User:
    """Utilisateur avec mot de passe TEST_PASSWORD et les rôles demandés (seeds de référence requis)."""
    roles = list(db.scalars(select(Role).where(Role.code.in_(role_codes)))) if role_codes else []
    assert len(roles) == len(role_codes), f"rôles introuvables : {role_codes}"
    user = build_user(mot_de_passe_hash=test_password_hash(), roles=roles, **over)
    return save(db, user)


def login(client, email: str, password: str = TEST_PASSWORD) -> str:
    """Se connecte et retourne le jeton d'accès. Le cookie de session reste dans `client`."""
    response = client.post("/api/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200, response.text
    return response.json()["access_token"]


def bearer(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


# --- Fichiers Excel .xls (ancien format) -----------------------------------------------------------


def xls(*sheets: tuple[str, list[list]]) -> bytes:
    """Classeur .xls (BIFF) : mêmes feuilles et valeurs qu'un .xlsx ; les dates sont de vraies
    dates Excel (format de date), comme dans un fichier de banque."""
    from datetime import date, datetime
    from io import BytesIO

    import xlwt

    workbook = xlwt.Workbook()
    date_style = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    for title, rows in sheets:
        sheet = workbook.add_sheet(title)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if value is None:
                    continue
                if isinstance(value, date | datetime):
                    sheet.write(r, c, value, date_style)
                else:
                    sheet.write(r, c, value)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def big_xlsx(header: list, rows: int, make_row) -> bytes:
    """Classeur .xlsx d'un en-tête et de `rows` lignes, écrit en flux (rapide pour 50 000 lignes)."""
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Feuille")
    sheet.append(header)
    for index in range(rows):
        sheet.append(make_row(index))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
