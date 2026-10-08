"""Import des relevés : analyse et aperçu (P7.1), puis confirmation et enregistrement (P7.2)."""

import json
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from itertools import count

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models import (
    AuditLog,
    BalanceCheck,
    Bank,
    BankAccount,
    BankAccountBalance,
    BankStatement,
    BankTransaction,
    ColumnMapping,
    Company,
    ImportBatch,
    PointageType,
)
from app.repositories import import_repository
from app.services.position_service import business_today
from tests.helpers import bearer, big_xlsx, build_account, login, make_auth_user, save, xls

URL = "/api/statements/import/analyse"
TODAY = business_today()
_numeros = count(1)

HEADER = ["Date opération", "Date valeur", "Libellé", "Débit", "Crédit", "Solde"]
STATEMENT = [
    ["ATTIJARIWAFA BANK - Relevé de compte"],
    ["Compte", "007 780 0001234500000123 45"],
    [],
    HEADER,
    [date(2026, 9, 24), date(2026, 9, 24), "VIR RECU CLIENT ABC REF: VR-001", None, 50000, 1050000],
    ["25/09/2026", "25/09/2026", "CHQ N° 1234567 FOURNISSEUR X", "12 500,50", None, 1037499.5],
    [date(2026, 9, 26), None, "frais   tenue de compte", 150, None, 1037349.5],
    [None, None, "TOTAL", 12650.5, 50000, None],
]


def xlsx(*sheets: tuple[str, list[list]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def tresorerie(client, reference) -> dict[str, str]:
    make_auth_user(reference, "TRESORERIE", email="tresorerie@example.com")
    return bearer(login(client, "tresorerie@example.com"))


@pytest.fixture
def account(db, reference) -> BankAccount:
    company = db.scalar(select(Company).filter_by(code="SIMTIS"))
    bank = db.scalar(select(Bank).filter_by(code="CIH"))
    return save(db, build_account(company, bank, numero=f"RIB-IMPORT-{next(_numeros):06d}"))


def analyse(client, headers, account_id, content, *, nom="releve.xlsx", **form):
    data = {"bank_account_id": str(account_id), **form}
    if isinstance(data.get("mapping"), dict):
        data["mapping"] = json.dumps(data["mapping"])
    files = {"fichier": (nom, content, "application/octet-stream")}
    return client.post(URL, data=data, files=files, headers=headers)


def counts(db) -> tuple[int, ...]:
    return tuple(
        db.scalar(select(func.count()).select_from(model))
        for model in (ImportBatch, BankStatement, BankTransaction, AuditLog)
    )


# --- Lecture d'un relevé standard ----------------------------------------------------------------


def test_statement_is_analysed_without_being_saved(client, tresorerie, account, db):
    before = counts(db)

    response = analyse(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT)))

    assert response.status_code == 200, response.text
    body = response.json()
    assert counts(db) == before  # rien n'est enregistré, pas même l'audit
    assert (body["bank_code"], body["devise"], body["feuille"]) == ("CIH", "MAD", "Relevé")
    assert body["ligne_entete"] == 4  # les lignes de titre au-dessus sont ignorées
    assert body["mapping_source"] == "Détection"
    assert body["mapping"] == {
        "date_operation": 0,
        "date_valeur": 1,
        "libelle": 2,
        "reference": None,
        "debit": 3,
        "credit": 4,
        "montant": None,
        "solde": 5,
        "pointage": None,
        "lettrage_escompte": None,
        "commentaire": None,
        "banque": None,
    }
    assert body["erreurs_mapping"] == []
    assert [column["entete"] for column in body["colonnes"]] == HEADER
    assert body["colonnes"][0]["exemples"] == ["24/09/2026", "25/09/2026", "26/09/2026"]
    assert body["deja_importe"] is False


def pointage_id(db, code: str) -> int:
    return db.scalar(select(PointageType.id).filter_by(code=code))


def test_lines_are_normalised(client, tresorerie, account, db):
    body = analyse(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT))).json()

    first, cheque, fees = body["lignes"]
    assert [line["statut"] for line in body["lignes"]] == ["Valide"] * 3
    assert first | {"hash_ligne": None} == {
        "numero": 5,
        "statut": "Valide",
        "motifs": [],
        "date_operation": "2026-09-24",
        "date_valeur": "2026-09-24",
        "libelle": "VIR RECU CLIENT ABC REF: VR-001",
        "reference": "VR-001",
        "debit": "0.00",
        "credit": "50000.00",
        "montant": "50000.00",
        "solde": "1050000.00",
        "solde_apercu": None,  # le fichier a ses soldes : rien n'est calculé
        "pointage": None,
        # Aucune catégorie nommée dans le libellé : sans pointage, à choisir (08/10/2026)
        "pointage_type_id": None,
        "pointage_libelle": None,
        "pointage_auto": False,
        "lettrage_escompte": None,
        "commentaire": None,
        "hash_ligne": None,
        "doublon_de": None,
    }
    assert (cheque["date_operation"], cheque["debit"], cheque["montant"]) == (
        "2026-09-25",
        "12500.50",
        "-12500.50",
    )
    assert cheque["reference"] == "1234567"
    assert cheque["pointage_libelle"] is None
    assert fees["libelle"] == "FRAIS TENUE DE COMPTE"
    assert (fees["pointage_libelle"], fees["pointage_auto"]) == ("FRAIS", True)
    assert len({line["hash_ligne"] for line in body["lignes"]}) == 3


def test_summary_totals_period_and_balances(client, tresorerie, account):
    resume = analyse(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT))).json()["resume"]

    assert resume == {
        "nb_lignes": 3,
        "nb_valides": 3,
        "nb_erreurs": 0,
        "nb_doublons": 0,
        "nb_ignorees": 1,  # la ligne TOTAL
        "total_debit": "12650.50",
        "total_credit": "50000.00",
        "periode_debut": "2026-09-24",
        "periode_fin": "2026-09-26",
        "solde_ouverture": "1000000.00",
        "solde_cloture": "1037349.50",
        "solde_ouverture_fichier": False,  # déduits des lignes : le fichier n'a pas de SOLDE INITIAL
        "solde_cloture_fichier": False,
        "soldes_coherents": True,
    }


# --- Lignes SOLDE INITIAL / SOLDE FINAL ------------------------------------------------------------

WITH_BALANCE_LINES = [
    HEADER,
    ["01/09/2026", None, "SOLDE INITIAL", None, None, 120000],  # datée, sans débit ni crédit
    ["02/09/2026", None, "VIR CLIENT ATLAS", None, 45000, 165000],
    ["03/09/2026", None, "COMMISSION BANCAIRE", 250, None, 164750],
    [None, None, "SOLDE FINAL", None, None, 164750],
]


def test_opening_balance_line_is_not_an_error(client, tresorerie, account):
    body = analyse(client, tresorerie, account.id, xlsx(("Relevé", WITH_BALANCE_LINES))).json()

    assert [line["libelle"] for line in body["lignes"]] == [
        "VIR CLIENT ATLAS",
        "COMMISSION BANCAIRE",
    ]
    assert {line["statut"] for line in body["lignes"]} == {"Valide"}
    resume = body["resume"]
    assert (resume["nb_erreurs"], resume["nb_ignorees"]) == (0, 2)
    assert (resume["solde_ouverture"], resume["solde_ouverture_fichier"]) == ("120000.00", True)
    assert (resume["solde_cloture"], resume["solde_cloture_fichier"]) == ("164750.00", True)
    assert resume["soldes_coherents"] is True


def test_opening_balance_line_is_checked_against_the_movements(client, tresorerie, account):
    rows = [row[:] for row in WITH_BALANCE_LINES]
    rows[1][5] = 100000  # solde initial faux : 100 000 + 44 750 ≠ 164 750

    resume = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["resume"]

    assert resume["solde_ouverture"] == "100000.00"
    assert resume["soldes_coherents"] is False


def test_opening_balance_line_without_balance_column_still_gives_the_balances(
    client, tresorerie, account
):
    rows = [
        ["Date", "Libellé", "Montant", "Solde"],
        [None, "Ancien solde", None, "1 000,00"],
        ["02/09/2026", "VIR CLIENT", "500", None],
        [None, "Nouveau solde", None, "1 500,00"],
    ]

    resume = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["resume"]

    assert (resume["solde_ouverture"], resume["solde_cloture"]) == ("1000.00", "1500.00")
    assert resume["soldes_coherents"] is True


def test_confirmed_statement_keeps_the_opening_balance_of_the_file(client, tresorerie, account, db):
    body = confirm(client, tresorerie, account.id, xlsx(("Relevé", WITH_BALANCE_LINES))).json()

    assert (body["nb_importees"], body["solde_ouverture"]) == (2, "120000.00")
    assert db.get(BankStatement, body["statement_id"]).solde_ouverture == Decimal("120000.00")


@pytest.mark.parametrize("empty", [0, "0", "0,00", "-", " "])
def test_opening_balance_line_with_zero_amounts_is_not_an_operation(
    client, tresorerie, account, empty
):
    """Beaucoup de banques écrivent 0 ou « - » dans Débit / Crédit sur les lignes de solde."""
    rows = [
        HEADER,
        ["01/09/2026", None, "SOLDE INITIAL", empty, empty, 145000],
        ["02/09/2026", None, "VIR CLIENT ATLAS", None, 45000, 190000],
        [None, None, "NOUVEAU SOLDE", empty, empty, 190000],
    ]

    body = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()

    assert [line["libelle"] for line in body["lignes"]] == ["VIR CLIENT ATLAS"]
    assert (body["resume"]["nb_erreurs"], body["resume"]["nb_ignorees"]) == (0, 2)
    assert (body["resume"]["solde_ouverture"], body["resume"]["solde_ouverture_fichier"]) == (
        "145000.00",
        True,
    )
    assert body["resume"]["solde_cloture"] == "190000.00"
    assert body["resume"]["soldes_coherents"] is True


def test_balance_line_with_a_real_amount_stays_an_operation(client, tresorerie, account):
    rows = [HEADER, ["01/09/2026", None, "REPORT FRAIS AOUT", 150, None, 144850]]

    lines = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["lignes"]

    assert [(line["libelle"], line["debit"], line["statut"]) for line in lines] == [
        ("REPORT FRAIS AOUT", "150.00", "Valide")
    ]


def test_newest_first_statement_gives_the_same_balances(client, tresorerie, account):
    rows = [HEADER, *reversed(STATEMENT[4:7])]

    resume = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["resume"]

    assert (resume["solde_ouverture"], resume["solde_cloture"]) == ("1000000.00", "1037349.50")
    assert resume["soldes_coherents"] is True


def test_broken_balance_chain_is_reported(client, tresorerie, account):
    rows = [HEADER, *STATEMENT[4:7]]
    rows[-1] = [*rows[-1][:5], 999]

    resume = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["resume"]

    assert resume["soldes_coherents"] is False


def test_signed_amount_column(client, tresorerie, account):
    rows = [
        ["Date", "Libellé", "Montant"],
        ["24/09/2026", "VIR RECU", "1 000,00"],
        ["25/09/2026", "PRLV", "-250"],
    ]

    body = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()

    assert body["mapping"]["montant"] == 2
    assert [(line["debit"], line["credit"]) for line in body["lignes"]] == [
        ("0.00", "1000.00"),
        ("250.00", "0.00"),
    ]
    assert body["resume"]["solde_ouverture"] is None  # pas de colonne Solde
    assert body["resume"]["soldes_coherents"] is None


# Relevé BP « Format_Different » : une colonne Montant toujours positive, le sens est dans le Solde
UNSIGNED = [
    ["Date opération", "Opération", "Montant", "Solde"],
    [None, "SOLDE INITIAL", None, 185000],
    ["02/09/2026", "VIR RECU ALPHA MODE", 38500, 223500],
    ["03/09/2026", "PRLV FOURNISSEUR TEXTILE", 12800, 210700],
    ["05/09/2026", "COMMISSION BANCAIRE", 175, 210525],
]


def test_unsigned_amount_takes_its_direction_from_the_balance(client, tresorerie, account):
    """Décision du 04/10/2026 : le solde qui baisse du montant prouve un débit."""
    body = analyse(client, tresorerie, account.id, xlsx(("Relevé", UNSIGNED))).json()

    assert [
        (line["libelle"], line["debit"], line["credit"], line["statut"], line["pointage_libelle"])
        for line in body["lignes"]
    ] == [
        ("VIR RECU ALPHA MODE", "0.00", "38500.00", "Valide", None),
        ("PRLV FOURNISSEUR TEXTILE", "12800.00", "0.00", "Valide", None),
        ("COMMISSION BANCAIRE", "175.00", "0.00", "Valide", "COM"),  # synonyme
    ]
    assert body["resume"]["soldes_coherents"] is True
    assert (body["resume"]["total_debit"], body["resume"]["total_credit"]) == (
        "12975.00",
        "38500.00",
    )


def test_unsigned_amount_first_line_uses_the_previous_line_without_opening_balance(
    client, tresorerie, account
):
    rows = [UNSIGNED[0], *UNSIGNED[2:]]  # sans SOLDE INITIAL

    lines = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["lignes"]

    # La 1re ligne n'a pas de solde précédent : son sens n'est pas prouvé, elle passe en erreur
    assert lines[0]["statut"] == "Erreur"
    assert lines[0]["motifs"] == [
        "Sens introuvable : le solde ne permet pas de savoir si c'est un débit ou un crédit."
    ]
    assert [(line["debit"], line["statut"]) for line in lines[1:]] == [
        ("12800.00", "Valide"),
        ("175.00", "Valide"),
    ]


def test_unsigned_amount_with_an_unproven_line_is_never_a_silent_credit(
    client, tresorerie, account
):
    rows = [*UNSIGNED[:4], ["05/09/2026", "VIR INCONNU", 999, 210525]]  # le solde ne colle pas

    lines = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["lignes"]

    assert lines[-1]["statut"] == "Erreur"
    assert lines[-1]["motifs"][0].startswith("Sens introuvable")


def test_newest_first_unsigned_statement_is_read_in_date_order(client, tresorerie, account):
    rows = [UNSIGNED[0], UNSIGNED[1], *reversed(UNSIGNED[2:])]

    lines = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["lignes"]

    assert [(line["libelle"], line["debit"]) for line in lines] == [
        ("COMMISSION BANCAIRE", "175.00"),
        ("PRLV FOURNISSEUR TEXTILE", "12800.00"),
        ("VIR RECU ALPHA MODE", "0.00"),
    ]
    assert all(line["statut"] == "Valide" for line in lines)


def test_signed_amount_column_is_not_reread_from_the_balance(client, tresorerie, account):
    """Une colonne qui contient déjà des montants négatifs est un vrai montant signé."""
    rows = [
        ["Date", "Libellé", "Montant", "Solde"],
        ["02/09/2026", "VIR RECU", 100, 1100],
        ["03/09/2026", "PRLV", -50, 1050],
        ["04/09/2026", "AUTRE", 10, 999],  # solde incohérent : signalé par le résumé, pas ici
    ]

    lines = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["lignes"]

    assert [(line["debit"], line["credit"], line["statut"]) for line in lines] == [
        ("0.00", "100.00", "Valide"),
        ("50.00", "0.00", "Valide"),
        ("0.00", "10.00", "Valide"),
    ]


def test_unsigned_amount_confirmed_with_its_debits(client, tresorerie, account, db):
    body = confirm(client, tresorerie, account.id, xlsx(("Relevé", UNSIGNED))).json()

    rows = db.scalars(
        select(BankTransaction)
        .filter_by(statement_id=body["statement_id"])
        .order_by(BankTransaction.id)
    ).all()
    assert [(row.debit, row.credit, row.montant) for row in rows] == [
        (Decimal("0.00"), Decimal("38500.00"), Decimal("38500.00")),
        (Decimal("12800.00"), Decimal("0.00"), Decimal("-12800.00")),
        (Decimal("175.00"), Decimal("0.00"), Decimal("-175.00")),
    ]


# --- Lignes en erreur et en double -----------------------------------------------------------------


def test_invalid_lines_are_reported_with_their_reasons(client, tresorerie, account):
    future = (TODAY + timedelta(days=2)).strftime("%d/%m/%Y")
    rows = [
        HEADER,
        ["31/02/2026", None, "DATE IMPOSSIBLE", 10, None, None],
        ["24/09/2026", None, None, 10, None, None],
        ["24/09/2026", None, "LES DEUX", 10, 20, None],
        ["24/09/2026", None, "AUCUN MONTANT", None, None, None],
        ["24/09/2026", None, "MONTANT ILLISIBLE", "dix", None, None],
        [future, None, "DANS LE FUTUR", 10, None, None],
        ["24/09/2026", None, "VALIDE", 10, None, None],
    ]

    body = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()

    reasons = {line["numero"]: (line["statut"], line["motifs"]) for line in body["lignes"]}
    assert reasons[2][0] == "Erreur" and "Date illisible" in reasons[2][1][0]
    assert reasons[3] == ("Erreur", ["Libellé manquant."])
    assert reasons[4] == ("Erreur", ["Débit et crédit renseignés sur la même ligne."])
    assert reasons[5] == ("Erreur", ["Ni débit ni crédit."])
    assert reasons[6][0] == "Erreur" and reasons[6][1][0].startswith("Débit : Montant illisible")
    assert reasons[7] == ("Erreur", ["Date d'opération dans le futur."])
    assert reasons[8] == ("Valide", [])
    assert (body["resume"]["nb_erreurs"], body["resume"]["nb_valides"]) == (6, 1)
    assert body["resume"]["total_debit"] == "10.00"  # seules les lignes valides comptent


def test_identical_lines_in_the_file_are_flagged(client, tresorerie, account):
    rows = [HEADER, *[["24/09/2026", None, "FRAIS SMS", 10, None, None]] * 2]

    lines = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["lignes"]

    assert [line["statut"] for line in lines] == ["Valide", "Doublon"]
    assert lines[1]["motifs"] == ["Ligne identique à la ligne 2 du fichier."]
    assert lines[1]["doublon_de"] == 2
    # Si l'utilisateur garde les deux, elles restent distinctes
    assert lines[0]["hash_ligne"] != lines[1]["hash_ligne"]


def test_lines_already_imported_are_flagged(client, tresorerie, account, db):
    content = xlsx(("Relevé", STATEMENT))
    first = analyse(client, tresorerie, account.id, content).json()["lignes"][0]
    statement = save(db, BankStatement(bank_account_id=account.id))
    save(
        db,
        BankTransaction(
            statement_id=statement.id,
            bank_account_id=account.id,
            date_operation=date(2026, 9, 24),
            libelle=first["libelle"],
            debit=0,
            credit=50000,
            montant=50000,
            hash_ligne=first["hash_ligne"],
        ),
    )

    body = analyse(client, tresorerie, account.id, content).json()

    assert body["lignes"][0]["statut"] == "Doublon"
    assert body["lignes"][0]["motifs"] == ["Déjà importée pour ce compte."]
    assert body["resume"]["nb_doublons"] == 1


def test_already_confirmed_file_is_reported(client, tresorerie, account, db):
    content = xlsx(("Relevé", STATEMENT))
    fichier_hash = analyse(client, tresorerie, account.id, content).json()["fichier_hash"]
    save(
        db,
        ImportBatch(
            type="Banque",
            company_id=account.company_id,
            bank_account_id=account.id,
            fichier_nom="ancien.xlsx",
            fichier_hash=fichier_hash,
            statut="Confirmé",
        ),
    )

    assert analyse(client, tresorerie, account.id, content).json()["deja_importe"] is True


def test_bank_column_must_match_the_account(client, tresorerie, account):
    rows = [
        ["Banque", "Date", "Libellé", "Débit"],
        ["CIH BANK", "24/09/2026", "OK", 1],
        ["BMCE", "24/09/2026", "AUTRE BANQUE", 1],
    ]

    lines = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["lignes"]

    assert lines[0]["statut"] == "Valide"
    assert lines[1]["motifs"] == ["Banque « BMCE » différente de celle du compte (CIH)."]


# --- Correspondance des colonnes -----------------------------------------------------------------

UNKNOWN_HEADERS = [["Col1", "Col2", "Col3"], ["24/09/2026", "VIR RECU", 100]]


def test_unrecognised_columns_give_mapping_errors_and_no_lines(client, tresorerie, account):
    body = analyse(client, tresorerie, account.id, xlsx(("Relevé", UNKNOWN_HEADERS))).json()

    assert body["erreurs_mapping"] == [
        "Colonne obligatoire non associée : Date d'opération.",
        "Colonne obligatoire non associée : Libellé.",
        "Associez au moins une colonne de montant : Débit, Crédit ou Montant signé.",
    ]
    assert body["lignes"] == []


def test_user_mapping_is_applied(client, tresorerie, account):
    mapping = {"date_operation": 0, "libelle": 1, "credit": 2}

    body = analyse(
        client, tresorerie, account.id, xlsx(("Relevé", UNKNOWN_HEADERS)), mapping=mapping
    ).json()

    assert body["mapping_source"] == "Utilisateur"
    assert body["erreurs_mapping"] == []
    assert body["lignes"][0]["credit"] == "100.00"


@pytest.mark.parametrize(
    ("mapping", "error"),
    [
        (
            {"date_operation": 0, "libelle": 0, "credit": 2},
            "La colonne A est associée à plusieurs champs.",
        ),
        (
            {"date_operation": 0, "libelle": 1, "credit": 2, "montant": 2},
            "Choisissez Débit et Crédit, ou Montant signé, pas les deux.",
        ),
        (
            {"date_operation": 0, "libelle": 1, "credit": 9},
            "Crédit : colonne inexistante dans le fichier.",
        ),
    ],
)
def test_inconsistent_user_mapping_is_reported(client, tresorerie, account, mapping, error):
    body = analyse(
        client, tresorerie, account.id, xlsx(("Relevé", UNKNOWN_HEADERS)), mapping=mapping
    ).json()

    assert error in body["erreurs_mapping"]
    assert body["lignes"] == []


@pytest.mark.parametrize("mapping", ["{pas du json", '{"inconnu": 1}', '{"debit": -1}', "[1]"])
def test_malformed_mapping_gives_422(client, tresorerie, account, mapping):
    content = xlsx(("Relevé", STATEMENT))

    assert analyse(client, tresorerie, account.id, content, mapping=mapping).status_code == 422


def test_saved_bank_mapping_is_reused(client, tresorerie, account, db):
    save(
        db,
        ColumnMapping(
            type="Banque",
            bank_id=account.bank_id,
            nom="CIH standard",
            mapping={"date_operation": "Jour", "libelle": "Texte", "credit": "Entrées"},
        ),
    )
    rows = [["Jour", "Texte", "Entrées"], ["24/09/2026", "VIR RECU", 100]]

    body = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()

    assert body["mapping_source"] == "Modèle de la banque"
    assert (body["mapping"]["date_operation"], body["mapping"]["credit"]) == (0, 2)
    assert body["lignes"][0]["statut"] == "Valide"


def test_sheet_can_be_chosen(client, tresorerie, account):
    content = xlsx(("Couverture", [["Relevé de septembre"]]), ("Opérations", STATEMENT))

    default = analyse(client, tresorerie, account.id, content).json()
    chosen = analyse(client, tresorerie, account.id, content, feuille="Opérations").json()
    unknown = analyse(client, tresorerie, account.id, content, feuille="Autre")

    assert default["feuilles"] == ["Couverture", "Opérations"]
    assert default["erreurs_mapping"] != []  # la première feuille n'est pas un relevé
    assert chosen["resume"]["nb_valides"] == 3
    assert unknown.status_code == 409
    assert unknown.json() == {"detail": "Feuille introuvable dans le fichier : « Autre »."}


# --- Fichier et compte refusés -------------------------------------------------------------------


@pytest.mark.parametrize(
    ("nom", "content", "detail"),
    [
        ("releve.csv", b"Date;Libelle", "Seuls les fichiers Excel .xlsx ou .xls sont acceptés."),
        ("releve.xlsx", b"", "Le fichier est vide."),
        (
            "releve.xlsx",
            b"pas un classeur",
            "Fichier illisible : ce n'est pas un classeur Excel .xlsx ou .xls valide.",
        ),
        ("releve.xlsx", b"x" * (20 * 1024 * 1024 + 1), "Fichier trop volumineux : 20 Mo au plus."),
    ],
)
def test_refused_file_gives_409(client, tresorerie, account, nom, content, detail):
    response = analyse(client, tresorerie, account.id, content, nom=nom)

    assert response.status_code == 409
    assert response.json() == {"detail": detail}


def test_empty_sheet_gives_409(client, tresorerie, account):
    response = analyse(client, tresorerie, account.id, xlsx(("Vide", [])))

    assert response.status_code == 409


def test_inactive_account_gives_409(client, tresorerie, account, db):
    account.actif = False
    db.flush()

    response = analyse(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT)))

    assert response.status_code == 409
    assert "inactif" in response.json()["detail"]


def test_unknown_account_gives_404(client, tresorerie):
    assert analyse(client, tresorerie, 999999, xlsx(("Relevé", STATEMENT))).status_code == 404


@pytest.mark.parametrize("role", ["DIRECTION", "COMPTABLE"])
def test_roles_without_import_permission_get_403(client, reference, account, role):
    make_auth_user(reference, role, email=f"{role.lower()}@example.com")
    headers = bearer(login(client, f"{role.lower()}@example.com"))

    response = analyse(client, headers, account.id, xlsx(("Relevé", STATEMENT)))

    assert response.status_code == 403


# --- Confirmation (P7.2) ---------------------------------------------------------------------------

CONFIRM_URL = "/api/statements/import/confirm"
CLOSING_DAY = date(2026, 9, 26)


def confirm(client, headers, account_id, content, *, nom="releve.xlsx", lignes=None, **form):
    # Fichier sans soldes : un solde d'ouverture est exigé (08/10/2026) ; 0 par défaut dans les
    # tests, sans effet quand le fichier a ses soldes. `solde_ouverture=None` n'en envoie pas.
    form.setdefault("solde_ouverture", "0")
    if form["solde_ouverture"] is None:
        del form["solde_ouverture"]
    data = {"bank_account_id": str(account_id)}
    for key, value in form.items():
        data[key] = json.dumps(value) if isinstance(value, dict | list) else str(value).lower()
    files = {"fichier": (nom, content, "application/octet-stream")}
    if lignes is not None:
        # Comme le navigateur : un fichier JSON, sans la limite de 1 Mo d'un champ de formulaire
        files["lignes"] = ("lignes.json", json.dumps(lignes).encode(), "application/json")
    return client.post(CONFIRM_URL, data=data, files=files, headers=headers)


def transactions(db, account) -> list[BankTransaction]:
    query = select(BankTransaction).filter_by(bank_account_id=account.id)
    return list(db.scalars(query.order_by(BankTransaction.id)))


def day_balance(db, account, jour=CLOSING_DAY) -> BankAccountBalance | None:
    query = select(BankAccountBalance).filter_by(bank_account_id=account.id, date_solde=jour)
    return db.scalar(query)


def test_confirmation_saves_the_statement(client, tresorerie, account, db):
    response = confirm(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT)))

    assert response.status_code == 201, response.text
    body = response.json()
    assert body | {"import_id": 0, "statement_id": 0} == {
        "import_id": 0,
        "statement_id": 0,
        "bank_account_id": account.id,
        "fichier_nom": "releve.xlsx",
        "nb_importees": 3,
        "nb_erreurs_ecartees": 0,
        "nb_doublons_ecartes": 0,
        "total_debit": "12650.50",
        "total_credit": "50000.00",
        "periode_debut": "2026-09-24",
        "periode_fin": "2026-09-26",
        "solde_ouverture": "1000000.00",
        "solde_cloture": "1037349.50",
        "soldes_coherents": True,
        "controle_solde": None,  # aucun solde enregistré ce jour-là : rien à comparer
        "solde_du_jour": "Créé",
        "modele_enregistre": True,
    }
    batch = db.get(ImportBatch, body["import_id"])
    assert (batch.statut, batch.type, batch.nb_lignes, batch.company_id) == (
        "Confirmé",
        "Banque",
        3,
        account.company_id,
    )
    statement = db.get(BankStatement, body["statement_id"])
    assert statement.import_batch_id == batch.id
    assert (statement.solde_ouverture, statement.solde_cloture) == (
        Decimal("1000000.00"),
        Decimal("1037349.50"),
    )
    rows = transactions(db, account)
    assert [(row.date_operation, row.debit, row.credit, row.montant) for row in rows] == [
        (date(2026, 9, 24), Decimal("0.00"), Decimal("50000.00"), Decimal("50000.00")),
        (date(2026, 9, 25), Decimal("12500.50"), Decimal("0.00"), Decimal("-12500.50")),
        (date(2026, 9, 26), Decimal("150.00"), Decimal("0.00"), Decimal("-150.00")),
    ]
    assert rows[1].reference == "1234567"
    assert all(row.statut == "Non rapprochée" and row.statement_id == statement.id for row in rows)
    balance = day_balance(db, account)
    assert (balance.solde, balance.source, balance.credit_utilise) == (
        Decimal("1037349.50"),
        "Relevé",
        None,
    )
    [entry] = db.scalars(select(AuditLog).filter_by(action="import_releve")).all()
    assert entry.entite_id == str(batch.id)
    assert entry.nouvelle_valeur["nb_lignes"] == 3
    assert entry.nouvelle_valeur["total_debit"] == "12650.50"


def test_same_file_cannot_be_confirmed_twice(client, tresorerie, account, db):
    content = xlsx(("Relevé", STATEMENT))
    confirm(client, tresorerie, account.id, content)

    again = confirm(client, tresorerie, account.id, content, nom="copie.xlsx")
    preview = analyse(client, tresorerie, account.id, content).json()

    assert again.status_code == 409
    assert again.json() == {"detail": "Ce fichier a déjà été importé pour cette société."}
    assert len(transactions(db, account)) == 3
    assert preview["deja_importe"] is True
    assert {line["statut"] for line in preview["lignes"]} == {"Doublon"}


def test_overlapping_statement_imports_only_the_new_lines(client, tresorerie, account, db):
    confirm(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT)))
    new_line = ["27/09/2026", None, "VIR RECU CLIENT DEF", None, 1000, 1038349.5]
    next_rows = [HEADER, STATEMENT[6], new_line]

    body = confirm(client, tresorerie, account.id, xlsx(("Relevé", next_rows))).json()

    assert (body["nb_importees"], body["nb_doublons_ecartes"]) == (1, 1)
    assert len(transactions(db, account)) == 4


ROWS_WITH_ERROR = [*STATEMENT[:7], ["27/09/2026", None, "MONTANT ILLISIBLE", "dix", None, None]]


def test_lines_in_error_block_the_confirmation(client, tresorerie, account, db):
    before = counts(db)

    response = confirm(client, tresorerie, account.id, xlsx(("Relevé", ROWS_WITH_ERROR)))

    assert response.status_code == 409
    assert response.json() == {
        "detail": "1 ligne en erreur : corrigez le fichier, ou confirmez en les écartant."
    }
    assert counts(db) == before


def test_lines_in_error_can_be_discarded(client, tresorerie, account, db):
    response = confirm(
        client, tresorerie, account.id, xlsx(("Relevé", ROWS_WITH_ERROR)), ecarter_erreurs=True
    )

    assert response.status_code == 201
    assert (response.json()["nb_importees"], response.json()["nb_erreurs_ecartees"]) == (3, 1)
    assert db.scalar(select(ImportBatch).filter_by(bank_account_id=account.id)).nb_erreurs == 1


@pytest.mark.parametrize(("kept", "expected"), [([], 1), ([3], 2)])
def test_identical_line_in_the_file_is_kept_only_on_request(
    client, tresorerie, account, db, kept, expected
):
    rows = [HEADER, *[["24/09/2026", None, "FRAIS SMS", 10, None, None]] * 2]

    body = confirm(
        client, tresorerie, account.id, xlsx(("Relevé", rows)), garder_doublons=kept
    ).json()

    assert body["nb_importees"] == expected
    assert body["nb_doublons_ecartes"] == 2 - expected
    assert len({row.hash_ligne for row in transactions(db, account)}) == expected


def test_only_identical_lines_of_the_file_can_be_kept(client, tresorerie, account):
    response = confirm(
        client, tresorerie, account.id, xlsx(("Relevé", STATEMENT)), garder_doublons=[5]
    )

    assert response.status_code == 409
    assert "La ligne 5 n'est pas identique" in response.json()["detail"]


@pytest.mark.parametrize("value", ["[0]", "{}", "pas du json", '["a"]'])
def test_malformed_kept_lines_give_422(client, tresorerie, account, value):
    content = xlsx(("Relevé", STATEMENT))
    data = {"bank_account_id": str(account.id), "garder_doublons": value}
    files = {"fichier": ("releve.xlsx", content, "application/octet-stream")}

    assert client.post(CONFIRM_URL, data=data, files=files, headers=tresorerie).status_code == 422


def test_file_without_any_line_to_import_is_refused(client, tresorerie, account):
    rows = [HEADER, ["24/09/2026", None, "ILLISIBLE", "x", None, None]]

    response = confirm(client, tresorerie, account.id, xlsx(("Relevé", rows)), ecarter_erreurs=True)

    assert response.status_code == 409
    assert response.json() == {"detail": "Aucune ligne à importer dans ce fichier."}


def test_incomplete_mapping_refuses_the_confirmation(client, tresorerie, account):
    response = confirm(client, tresorerie, account.id, xlsx(("Relevé", UNKNOWN_HEADERS)))

    assert response.status_code == 409
    assert response.json()["detail"].startswith("Correspondance des colonnes incomplète")


# --- Contrôle et solde du jour de clôture --------------------------------------------------------


def manual_balance(db, account, solde: str) -> BankAccountBalance:
    return save(
        db,
        BankAccountBalance(
            bank_account_id=account.id,
            date_solde=CLOSING_DAY,
            solde=Decimal(solde),
            credit_utilise=Decimal("300000"),
            source="Saisie",
        ),
    )


def test_matching_recorded_balance_gives_a_conforming_check(client, tresorerie, account, db):
    manual_balance(db, account, "1037349.50")

    body = confirm(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT))).json()

    assert body["controle_solde"] == {
        "statut": "Conforme",
        "date_controle": "2026-09-26",
        "solde_releve": "1037349.50",
        "solde_enregistre": "1037349.50",
        "ecart": "0.00",
        "commentaire": None,
    }
    assert body["solde_du_jour"] == "Corrigé"  # même montant, la source devient « Relevé »
    assert day_balance(db, account).source == "Relevé"


def test_different_recorded_balance_gives_a_gap_and_the_statement_wins(
    client, tresorerie, account, db
):
    manual_balance(db, account, "1000000")

    body = confirm(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT))).json()

    assert (body["controle_solde"]["statut"], body["controle_solde"]["ecart"]) == (
        "Écart",
        "37349.50",
    )
    check = db.scalar(select(BalanceCheck).filter_by(bank_account_id=account.id))
    assert (check.solde_enregistre, check.bank_statement_id) == (
        Decimal("1000000.00"),
        body["statement_id"],
    )
    balance = day_balance(db, account)
    assert (balance.solde, balance.source) == (Decimal("1037349.50"), "Relevé")
    assert balance.credit_utilise == Decimal("300000.00")  # le crédit utilisé saisi est conservé
    [entry] = db.scalars(select(AuditLog).filter_by(action="correction_solde")).all()
    assert entry.ancienne_valeur == {
        "date_solde": "2026-09-26",
        "solde": "1000000.00",
        "source": "Saisie",
    }


def test_inconsistent_statement_gives_a_check_to_verify(client, tresorerie, account, db):
    manual_balance(db, account, "999")
    rows = [HEADER, *STATEMENT[4:7]]
    rows[-1] = [*rows[-1][:5], 999]

    body = confirm(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()

    assert body["soldes_coherents"] is False
    assert body["controle_solde"]["statut"] == "À vérifier"
    assert body["controle_solde"]["commentaire"] is not None


def test_statement_without_balance_column_gets_a_computed_closing_balance(
    client, tresorerie, account, db
):
    """Décision du 08/10/2026 : sans colonne Solde, les soldes sont calculés depuis le solde
    d'ouverture et le solde de clôture devient le solde du jour."""
    rows = [["Date", "Libellé", "Montant"], ["26/09/2026", "VIR RECU", "1 000,00"]]

    body = confirm(
        client, tresorerie, account.id, xlsx(("Relevé", rows)), solde_ouverture="5000000"
    ).json()

    assert (body["solde_du_jour"], body["solde_cloture"]) == ("Créé", "5001000.00")
    assert day_balance(db, account).solde == Decimal("5001000.00")


# --- Modèle de correspondance et pointage --------------------------------------------------------


def test_confirmed_mapping_is_reused_for_the_next_statement(client, tresorerie, account, db):
    mapping = {"date_operation": 0, "libelle": 1, "credit": 2}
    confirm(client, tresorerie, account.id, xlsx(("Relevé", UNKNOWN_HEADERS)), mapping=mapping)
    next_file = [UNKNOWN_HEADERS[0], ["25/09/2026", "VIR RECU 2", 200]]

    body = analyse(client, tresorerie, account.id, xlsx(("Relevé", next_file))).json()

    saved = db.scalar(select(ColumnMapping).filter_by(bank_id=account.bank_id))
    assert saved.mapping == {"date_operation": "Col1", "libelle": "Col2", "credit": "Col3"}
    assert body["mapping_source"] == "Modèle de la banque"
    assert body["lignes"][0]["statut"] == "Valide"


def test_mapping_with_a_column_without_header_is_not_saved(client, tresorerie, account, db):
    rows = [["Date", "Libellé", None], ["24/09/2026", "VIR", 5]]
    mapping = {"date_operation": 0, "libelle": 1, "credit": 2}

    body = confirm(client, tresorerie, account.id, xlsx(("Relevé", rows)), mapping=mapping).json()

    assert body["modele_enregistre"] is False
    assert db.scalar(select(ColumnMapping).filter_by(bank_id=account.bank_id)) is None


def test_pointage_of_the_file_wins_and_is_guessed_otherwise(client, tresorerie, account, db):
    rows = [
        ["Date", "Libellé", "Débit", "Crédit", "Pointage"],
        ["24/09/2026", "VIR INTERNE", 10, None, "tva"],  # le fichier l'emporte
        ["24/09/2026", "DGI ACOMPTE", 20, None, "Virement interne"],  # inconnu → libellé
        ["24/09/2026", "AGIOS D'ECHELLE T3", 30, None, None],  # vide → le plus long nom
        ["24/09/2026", "VIR CLIENT", None, 40, None],  # rien de reconnu → sans pointage
    ]

    preview = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["lignes"]
    confirm(client, tresorerie, account.id, xlsx(("Relevé", rows)))

    assert [(line["pointage_libelle"], line["pointage_auto"]) for line in preview] == [
        ("TVA", False),
        ("DGI", True),  # une valeur inconnue ne bloque jamais la ligne
        ("AGIOS D'ECHELLE", True),
        (None, False),
    ]
    assert {line["statut"] for line in preview} == {"Valide"}
    assert [row.pointage_type_id for row in transactions(db, account)] == [
        pointage_id(db, "TVA"),
        pointage_id(db, "DGI"),
        pointage_id(db, "AGIOS_D_ECHELLE"),
        None,
    ]


@pytest.mark.parametrize("role", ["DIRECTION", "COMPTABLE"])
def test_roles_without_import_permission_cannot_confirm(client, reference, account, db, role):
    make_auth_user(reference, role, email=f"{role.lower()}@example.com")
    headers = bearer(login(client, f"{role.lower()}@example.com"))

    response = confirm(client, headers, account.id, xlsx(("Relevé", STATEMENT)))

    assert response.status_code == 403
    assert transactions(db, account) == []


# --- Consultation (P7.3) ---------------------------------------------------------------------------


def history(client, headers, company_id: int, **params):
    return client.get(
        "/api/statements", params={"company_id": company_id, **params}, headers=headers
    )


def test_history_lists_the_company_statements_newest_first(client, tresorerie, account, db):
    manual_balance(db, account, "1000000")
    first = confirm(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT))).json()
    second_rows = [HEADER, ["27/09/2026", None, "VIR RECU", None, 1000, 1038349.5]]
    second = confirm(
        client, tresorerie, account.id, xlsx(("Relevé", second_rows)), nom="octobre.xlsx"
    ).json()

    items = history(client, tresorerie, account.company_id).json()

    assert [item["id"] for item in items] == [second["statement_id"], first["statement_id"]]
    oldest = items[1]
    assert oldest | {"importe_le": None, "controle_solde": None} == {
        "id": first["statement_id"],
        "import_id": first["import_id"],
        "fichier_nom": "releve.xlsx",
        "importe_le": None,
        "importe_par": oldest["importe_par"],
        "bank_account_id": account.id,
        "bank_code": "CIH",
        "devise": "MAD",
        "compte_libelle": account.libelle,
        "compte_numero": account.numero,
        "periode_debut": "2026-09-24",
        "periode_fin": "2026-09-26",
        "solde_ouverture": "1000000.00",
        "solde_cloture": "1037349.50",
        "nb_lignes": 3,
        "nb_erreurs": 0,
        "nb_doublons": 0,
        "controle_solde": None,
    }
    assert oldest["importe_par"] is not None
    assert oldest["controle_solde"]["statut"] == "Écart"
    assert items[0]["controle_solde"] is None


def test_history_never_shows_the_other_company(client, tresorerie, account, db):
    confirm(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT)))
    other = db.scalar(select(Company).filter_by(code="SOCX"))

    assert history(client, tresorerie, other.id).json() == []


def test_history_can_be_filtered_by_account(client, tresorerie, account, db):
    confirm(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT)))

    same = history(client, tresorerie, account.company_id, bank_account_id=account.id).json()
    other = history(client, tresorerie, account.company_id, bank_account_id=999999).json()

    assert (len(same), other) == (1, [])


def test_history_of_unknown_company_gives_404(client, tresorerie):
    assert history(client, tresorerie, 999999).status_code == 404


def test_history_needs_a_company(client, tresorerie):
    assert client.get("/api/statements", headers=tresorerie).status_code == 422


def test_statement_transactions_are_listed_in_order(client, tresorerie, account, db):
    rows = [
        ["Date", "Libellé", "Débit", "Crédit", "Pointage"],
        ["25/09/2026", "CHQ N° 1234567", "12 500,50", None, "Cheque de banque"],
        ["24/09/2026", "VIR RECU", None, 50000, None],
    ]
    body = confirm(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()

    response = client.get(
        f"/api/statements/{body['statement_id']}/transactions", headers=tresorerie
    )

    assert response.status_code == 200
    first, second = response.json()
    assert (first["date_operation"], first["credit"], first["pointage"]) == (
        "2026-09-24",
        "50000.00",
        None,  # pas de pointage dans le fichier, rien de reconnu dans le libellé
    )
    assert second | {"id": 0} == {
        "id": 0,
        "societe": "Simtis",
        "pointage": "Cheque de banque",
        "banque": "CIH",
        "date_operation": "2026-09-25",
        "date_valeur": None,
        "libelle": "CHQ N° 1234567",
        "debit": "12500.50",
        "credit": "0.00",
        # Pas de colonne Solde : calculé depuis 0 (24/09 +50 000, puis 25/09 −12 500,50)
        "solde": "37499.50",
        "lettrage_escompte": None,
        "commentaire": None,
        "reference": "1234567",
        "montant": "-12500.50",
        "statut": "Non rapprochée",
        "origine": "Fichier",
        "pointage_type_id": pointage_id(db, "CHEQUE_DE_BANQUE"),
    }
    # Les 11 champs du format standard d'abord, dans leur ordre
    assert list(second)[1:12] == [
        "societe",
        "pointage",
        "banque",
        "date_operation",
        "date_valeur",
        "libelle",
        "debit",
        "credit",
        "solde",
        "lettrage_escompte",
        "commentaire",
    ]


def test_unknown_statement_gives_404(client, tresorerie):
    assert client.get("/api/statements/999999/transactions", headers=tresorerie).status_code == 404


# --- Export au format standard -------------------------------------------------------------------


def export(client, headers, statement_id: int):
    return client.get(f"/api/statements/{statement_id}/export", headers=headers)


def test_export_is_the_standard_statement_in_excel(client, tresorerie, account):
    rows = [
        [
            "Date opération",
            "Date valeur",
            "Libellé",
            "Débit",
            "Crédit",
            "Solde",
            "Pointage",
            "Lettrage",
            "Observation",
        ],
        [
            "24/09/2026",
            "24/09/2026",
            "VIR RECU",
            None,
            50000,
            1050000,
            "encaissement client",
            "L-12",
            None,
        ],
        ["25/09/2026", None, "FRAIS", "12,50", None, 1049987.5, "Inconnu", None, "à justifier"],
    ]
    body = confirm(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()

    response = export(client, tresorerie, body["statement_id"])

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert response.headers["content-disposition"] == (
        'attachment; filename="releve_CIH_MAD_2026-09-24_2026-09-25.xlsx"'
    )
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.title == "Relevé standard"
    values = [list(row) for row in sheet.iter_rows(values_only=True)]
    assert values[0] == [
        "Société",
        "Pointage",
        "Banque",
        "Date d'opération",
        "Date de valeur",
        "Libellé",
        "Débit",
        "Crédit",
        "Solde",
        "Lettrage / Escompte",
        "Commentaire",
    ]
    assert values[1] == [
        "Simtis",
        "ENCAISSEMENT CLIENT",
        "CIH",
        datetime(2026, 9, 24),
        datetime(2026, 9, 24),
        "VIR RECU",
        None,  # débit nul : cellule vide
        50000,
        1050000,
        "L-12",
        None,
    ]
    # Pointage inconnu du fichier : déduit du libellé (FRAIS) ; montants en vrais nombres
    assert values[2][:3] == ["Simtis", "FRAIS", "CIH"]
    assert (values[2][6], values[2][8], values[2][10]) == (12.5, 1049987.5, "à justifier")
    assert sheet["D2"].number_format == "DD/MM/YYYY"
    assert sheet["H2"].number_format == "#,##0.00"


def test_export_of_unknown_statement_gives_404(client, tresorerie):
    assert export(client, tresorerie, 999999).status_code == 404


def test_export_follows_the_reading_permissions(client, reference, tresorerie, account):
    body = confirm(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT))).json()
    make_auth_user(reference, "DIRECTION", email="direction@example.com")
    make_auth_user(reference, email="sans.role@example.com")

    direction = export(client, bearer(login(client, "direction@example.com")), body["statement_id"])
    sans_role = export(client, bearer(login(client, "sans.role@example.com")), body["statement_id"])

    assert (direction.status_code, sans_role.status_code) == (200, 403)


@pytest.mark.parametrize("role", ["COMPTABLE", "DIRECTION"])
def test_statements_are_read_with_reconciliation_view(client, reference, account, role):
    make_auth_user(reference, role, email=f"{role.lower()}@example.com")
    headers = bearer(login(client, f"{role.lower()}@example.com"))

    assert history(client, headers, account.company_id).status_code == 200


def test_user_without_statement_permission_gets_403(client, reference, account):
    make_auth_user(reference, email="sans.role@example.com")
    headers = bearer(login(client, "sans.role@example.com"))

    response = history(client, headers, account.company_id)

    assert response.status_code == 403
    assert response.json() == {
        "detail": "Permission requise : statements.import ou reconciliation.view."
    }


# --- Relevé continu d'un compte --------------------------------------------------------------------

SEPTEMBER = [HEADER, *STATEMENT[4:7]]  # 24, 25 et 26/09, solde 1 000 000 → 1 037 349,50
OCTOBER = [
    HEADER,
    STATEMENT[6],  # 26/09 déjà importé : écarté, pas recopié
    ["01/10/2026", None, "VIR RECU CLIENT DEF", None, 2650.5, 1040000],
]


def continuous(client, headers, account_id, **params):
    return client.get(f"/api/statements/accounts/{account_id}", params=params, headers=headers)


def test_each_import_is_appended_to_the_account_statement(client, tresorerie, account):
    confirm(client, tresorerie, account.id, xlsx(("Relevé", SEPTEMBER)), nom="septembre.xlsx")
    confirm(client, tresorerie, account.id, xlsx(("Relevé", OCTOBER)), nom="octobre.xlsx")

    body = continuous(client, tresorerie, account.id).json()

    assert body | {"operations": None} == {
        "bank_account_id": account.id,
        "societe": "Simtis",
        "bank_code": "CIH",
        "devise": "MAD",
        "compte_libelle": account.libelle,
        "compte_numero": account.numero,
        "periode_debut": "2026-09-24",
        "periode_fin": "2026-10-01",
        "nb_operations": 4,
        "total_debit": "12650.50",
        "total_credit": "52650.50",
        "solde_ouverture": "1000000.00",
        "solde_cloture": "1040000.00",
        "operations": None,
    }
    assert [op["date_operation"] for op in body["operations"]] == [
        "2026-09-24",
        "2026-09-25",
        "2026-09-26",
        "2026-10-01",
    ]
    assert body["operations"][0]["banque"] == "CIH"


def test_account_statement_can_be_limited_to_a_period(client, tresorerie, account):
    confirm(client, tresorerie, account.id, xlsx(("Relevé", SEPTEMBER)))
    confirm(client, tresorerie, account.id, xlsx(("Relevé", OCTOBER)), nom="octobre.xlsx")

    body = continuous(client, tresorerie, account.id, **{"from": "2026-09-25", "to": "2026-09-26"})

    assert [op["libelle"] for op in body.json()["operations"]] == [
        "CHQ N° 1234567 FOURNISSEUR X",
        "FRAIS TENUE DE COMPTE",
    ]
    assert body.json()["solde_ouverture"] == "1050000.00"  # avant le chèque du 25/09
    assert body.json()["solde_cloture"] == "1037349.50"


def test_account_statement_never_mixes_another_account(client, tresorerie, account, db):
    other = save(
        db,
        build_account(
            db.scalar(select(Company).filter_by(code="SIMTIS")),
            db.scalar(select(Bank).filter_by(code="BP")),
            numero="RIB-IMPORT-AUTRE",
        ),
    )
    confirm(client, tresorerie, account.id, xlsx(("Relevé", SEPTEMBER)))

    body = continuous(client, tresorerie, other.id).json()

    assert (body["nb_operations"], body["operations"], body["periode_debut"]) == (0, [], None)
    assert (body["solde_ouverture"], body["solde_cloture"]) == (None, None)


def test_account_statement_errors(client, tresorerie, account):
    reversed_dates = continuous(
        client, tresorerie, account.id, **{"from": "2026-10-02", "to": "2026-10-01"}
    )

    assert continuous(client, tresorerie, 999999).status_code == 404
    assert reversed_dates.status_code == 409


def test_account_statement_export(client, tresorerie, account):
    confirm(client, tresorerie, account.id, xlsx(("Relevé", SEPTEMBER)))
    confirm(client, tresorerie, account.id, xlsx(("Relevé", OCTOBER)), nom="octobre.xlsx")

    response = client.get(f"/api/statements/accounts/{account.id}/export", headers=tresorerie)

    assert response.headers["content-disposition"] == (
        'attachment; filename="releve_CIH_MAD_2026-09-24_2026-10-01.xlsx"'
    )
    values = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
    assert values[0][3] == "Date d'opération" and len(values[0]) == 11
    assert [row[3].date() for row in values[1:]] == [
        date(2026, 9, 24),
        date(2026, 9, 25),
        date(2026, 9, 26),
        date(2026, 10, 1),
    ]


def test_empty_account_statement_cannot_be_exported(client, tresorerie, account):
    response = client.get(f"/api/statements/accounts/{account.id}/export", headers=tresorerie)

    assert response.status_code == 409
    assert response.json() == {"detail": "Aucune opération à exporter sur cette période."}


def test_account_statement_follows_the_reading_permissions(client, reference, account):
    make_auth_user(reference, "DIRECTION", email="direction@example.com")
    make_auth_user(reference, email="sans.role@example.com")

    direction = continuous(client, bearer(login(client, "direction@example.com")), account.id)
    sans_role = continuous(client, bearer(login(client, "sans.role@example.com")), account.id)

    assert (direction.status_code, sans_role.status_code) == (200, 403)


# --- Confirmation avec les lignes corrigées de l'aperçu --------------------------------------------

EDITABLE = [
    HEADER,
    ["24/09/2026", "24/09/2026", "VIR RECU CLIENT ABC", None, 50000, 1050000],
    ["25/09/2026", None, "CHQ FOURNISSEUR", "dix", None, 1037499.5],  # montant illisible
    ["26/09/2026", None, "FRAIS SMS", 10, None, 1037489.5],
]


def submitted(client, headers, account, content):
    """Lignes de l'aperçu, au format attendu par `lignes`, avant toute correction."""
    body = analyse(client, headers, account.id, content).json()
    return [
        {
            "numero": line["numero"],
            "date_operation": line["date_operation"],
            "date_valeur": line["date_valeur"],
            "libelle": line["libelle"],
            "reference": line["reference"],
            "debit": line["debit"],
            "credit": line["credit"],
            "solde": line["solde"],
            "pointage_type_id": line["pointage_type_id"],
            "lettrage_escompte": line["lettrage_escompte"],
            "commentaire": line["commentaire"],
        }
        for line in body["lignes"]
    ]


def test_corrected_lines_are_saved_with_their_origin(client, tresorerie, account, db):
    content = xlsx(("Relevé", EDITABLE))
    lines = submitted(client, tresorerie, account, content)
    lines[1] |= {"debit": "12500.50", "credit": "0.00"}  # ligne en erreur corrigée
    lines[2]["commentaire"] = "Frais du mois"  # ligne valide corrigée

    response = confirm(client, tresorerie, account.id, content, lignes=lines)

    assert response.status_code == 201, response.text
    rows = transactions(db, account)
    assert [(row.libelle, row.debit, row.origine) for row in rows] == [
        ("VIR RECU CLIENT ABC", Decimal("0.00"), "Fichier"),
        ("CHQ FOURNISSEUR", Decimal("12500.50"), "Corrigée"),
        ("FRAIS SMS", Decimal("10.00"), "Corrigée"),
    ]
    assert rows[2].commentaire == "Frais du mois"
    [entry] = db.scalars(select(AuditLog).filter_by(action="import_releve")).all()
    corrections = entry.nouvelle_valeur["lignes_corrigees"]
    assert {
        "numero": 4,
        "champ": "commentaire",
        "avant": None,
        "apres": "Frais du mois",
    } in corrections
    assert {"numero": 3, "champ": "debit", "avant": None, "apres": "12500.50"} in corrections


def test_unchecked_file_line_is_not_saved_and_is_audited(client, tresorerie, account, db):
    content = xlsx(("Relevé", EDITABLE))
    lines = submitted(client, tresorerie, account, content)

    confirm(client, tresorerie, account.id, content, lignes=[lines[0], lines[2]])

    assert [row.libelle for row in transactions(db, account)] == [
        "VIR RECU CLIENT ABC",
        "FRAIS SMS",
    ]
    [entry] = db.scalars(select(AuditLog).filter_by(action="import_releve")).all()
    assert entry.nouvelle_valeur["lignes_fichier_non_importees"] == [3]


def test_unchanged_submitted_line_stays_from_file(client, tresorerie, account, db):
    content = xlsx(("Relevé", EDITABLE))
    lines = submitted(client, tresorerie, account, content)

    confirm(client, tresorerie, account.id, content, lignes=[lines[0]])

    assert [row.origine for row in transactions(db, account)] == ["Fichier"]
    [entry] = db.scalars(select(AuditLog).filter_by(action="import_releve")).all()
    assert entry.nouvelle_valeur["lignes_corrigees"] == []


def test_corrected_line_keeps_its_original_fingerprint(client, tresorerie, account, db):
    content = xlsx(("Relevé", EDITABLE))
    lines = submitted(client, tresorerie, account, content)
    lines[0]["libelle"] = "VIR RECU CLIENT ABC CORRIGE"
    confirm(client, tresorerie, account.id, content, lignes=[lines[0]])
    # La banque renvoie plus tard la même opération, dans sa version d'origine
    again = xlsx(("Relevé", [HEADER, EDITABLE[1]]))

    preview = analyse(client, tresorerie, account.id, again).json()

    assert preview["lignes"][0]["statut"] == "Doublon"
    assert preview["lignes"][0]["motifs"] == ["Déjà importée pour ce compte."]


def test_corrected_lines_that_become_identical_are_both_saved(client, tresorerie, account, db):
    rows = [
        HEADER,
        ["24/09/2026", None, "A", "dix", None, None],
        ["24/09/2026", None, "A", "onze", None, None],
    ]
    content = xlsx(("Relevé", rows))
    lines = submitted(client, tresorerie, account, content)
    for line in lines:
        line |= {"debit": "10.00", "credit": "0.00"}

    response = confirm(client, tresorerie, account.id, content, lignes=lines)

    assert response.status_code == 201, response.text
    assert len({row.hash_ligne for row in transactions(db, account)}) == 2


def test_invalid_submitted_line_refuses_everything(client, tresorerie, account, db):
    content = xlsx(("Relevé", EDITABLE))
    lines = submitted(client, tresorerie, account, content)  # ligne 3 toujours illisible
    before = counts(db)

    response = confirm(client, tresorerie, account.id, content, lignes=lines)

    assert response.status_code == 409
    assert response.json()["detail"].startswith("Ligne 3 : Ni débit ni crédit.")
    assert counts(db) == before


@pytest.mark.parametrize("numero", [99, 1])
def test_submitted_line_must_come_from_the_file(client, tresorerie, account, numero):
    content = xlsx(("Relevé", EDITABLE))
    lines = submitted(client, tresorerie, account, content)

    response = confirm(
        client, tresorerie, account.id, content, lignes=[lines[0] | {"numero": numero}]
    )

    assert response.status_code == 409
    assert response.json() == {"detail": f"La ligne {numero} n'est pas une opération du fichier."}


def test_submitted_line_already_imported_is_refused(client, tresorerie, account, db):
    content = xlsx(("Relevé", EDITABLE))
    lines = submitted(client, tresorerie, account, content)
    confirm(
        client, tresorerie, account.id, xlsx(("Relevé", [HEADER, EDITABLE[1]])), nom="autre.xlsx"
    )

    response = confirm(client, tresorerie, account.id, content, lignes=[lines[0]])

    assert response.status_code == 409
    assert response.json() == {"detail": "Ligne 2 : déjà importée pour ce compte."}


def test_submitted_inactive_pointage_is_refused(client, tresorerie, account, db):
    content = xlsx(("Relevé", EDITABLE))
    lines = submitted(client, tresorerie, account, content)
    frais = db.scalar(select(PointageType).filter_by(code="FRAIS"))
    frais.actif = False
    db.flush()

    response = confirm(
        client,
        tresorerie,
        account.id,
        content,
        lignes=[lines[0] | {"pointage_type_id": frais.id}],
    )

    assert response.status_code == 409
    assert response.json() == {"detail": "Ligne 2 : Pointage inconnu ou inactif."}


@pytest.mark.parametrize("extra", [{"garder_doublons": [3]}, {"ecarter_erreurs": True}])
def test_lines_cannot_be_combined_with_the_old_options(client, tresorerie, account, extra):
    content = xlsx(("Relevé", EDITABLE))
    lines = submitted(client, tresorerie, account, content)

    response = confirm(client, tresorerie, account.id, content, lignes=lines[:1], **extra)

    assert response.status_code == 422


@pytest.mark.parametrize(
    ("value", "status"),
    [("[]", 409), ("pas du json", 422), ('[{"numero": 2}, {"numero": 2}]', 422)],
)
def test_malformed_lines(client, tresorerie, account, value, status):
    content = xlsx(("Relevé", EDITABLE))
    data = {"bank_account_id": str(account.id)}
    files = {
        "fichier": ("releve.xlsx", content, "application/octet-stream"),
        "lignes": ("lignes.json", value.encode(), "application/json"),
    }

    response = client.post(CONFIRM_URL, data=data, files=files, headers=tresorerie)

    assert response.status_code == status


# --- Modification après l'import (champs métier) ---------------------------------------------------


def first_transaction(client, tresorerie, account, db) -> BankTransaction:
    """Première opération du relevé STATEMENT, pointée « TVA » (le libellé ne nomme aucune
    catégorie) pour que les tests de modification partent d'un pointage connu."""
    confirm(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT)))
    transaction = transactions(db, account)[0]
    transaction.pointage_type_id = pointage_id(db, "TVA")
    db.flush()
    return transaction


def patch(client, headers, transaction_id, body):
    return client.patch(
        f"/api/statements/transactions/{transaction_id}", json=body, headers=headers
    )


def test_business_fields_can_be_changed_and_are_audited(client, tresorerie, account, db):
    transaction = first_transaction(client, tresorerie, account, db)
    frais = pointage_id(db, "FRAIS")

    response = patch(
        client,
        tresorerie,
        transaction.id,
        {
            "pointage_type_id": frais,
            "lettrage_escompte": " L-12 ",
            "commentaire": "Vu avec la banque",
        },
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert (body["pointage"], body["lettrage_escompte"], body["commentaire"], body["origine"]) == (
        "FRAIS",
        "L-12",
        "Vu avec la banque",
        "Fichier",
    )
    [entry] = db.scalars(select(AuditLog).filter_by(action="modification_operation")).all()
    assert entry.ancienne_valeur == {
        "pointage_type_id": pointage_id(db, "TVA"),
        "lettrage_escompte": None,
        "commentaire": None,
    }
    assert entry.nouvelle_valeur["commentaire"] == "Vu avec la banque"


def test_same_values_write_no_audit(client, tresorerie, account, db):
    transaction = first_transaction(client, tresorerie, account, db)
    body = {
        "pointage_type_id": transaction.pointage_type_id,
        "lettrage_escompte": None,
        "commentaire": "  ",
    }

    response = patch(client, tresorerie, transaction.id, body)

    assert response.status_code == 200
    assert db.scalars(select(AuditLog).filter_by(action="modification_operation")).all() == []


@pytest.mark.parametrize("field", ["debit", "libelle", "date_operation", "solde"])
def test_bank_fields_cannot_be_changed(client, tresorerie, account, db, field):
    transaction = first_transaction(client, tresorerie, account, db)

    response = patch(client, tresorerie, transaction.id, {"commentaire": None, field: "1"})

    assert response.status_code == 422


def test_inactive_pointage_is_refused(client, tresorerie, account, db):
    transaction = first_transaction(client, tresorerie, account, db)
    frais = db.scalar(select(PointageType).filter_by(code="FRAIS"))
    frais.actif = False
    db.flush()

    response = patch(client, tresorerie, transaction.id, {"pointage_type_id": frais.id})

    assert response.status_code == 409
    assert response.json() == {"detail": "Pointage inconnu ou inactif."}


def test_unknown_transaction_gives_404(client, tresorerie):
    assert patch(client, tresorerie, 999999, {"commentaire": "x"}).status_code == 404


@pytest.mark.parametrize("role", ["COMPTABLE", "DIRECTION"])
def test_only_importers_can_change_a_transaction(client, reference, tresorerie, account, db, role):
    transaction = first_transaction(client, tresorerie, account, db)
    make_auth_user(reference, role, email=f"{role.lower()}.patch@example.com")
    headers = bearer(login(client, f"{role.lower()}.patch@example.com"))

    assert patch(client, headers, transaction.id, {"commentaire": "x"}).status_code == 403


def test_transactions_report_their_origin(client, tresorerie, account, db):
    body = confirm(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT))).json()

    rows = client.get(
        f"/api/statements/{body['statement_id']}/transactions", headers=tresorerie
    ).json()

    assert {row["origine"] for row in rows} == {"Fichier"}


# --- Corrections issues de la relecture finale -----------------------------------------------------


def test_untouched_file_error_cannot_be_imported_as_is(client, tresorerie, account, db):
    rows = [HEADER, ["24/09/2026", None, "VIR RECU", None, 10, "abc"]]  # solde illisible
    content = xlsx(("Relevé", rows))
    lines = submitted(client, tresorerie, account, content)
    before = counts(db)

    response = confirm(client, tresorerie, account.id, content, lignes=lines)

    assert response.status_code == 409
    assert response.json() == {"detail": "Ligne 2 : Solde : Montant illisible : « abc »."}
    assert counts(db) == before


def test_file_error_fixed_by_the_user_is_corrected_and_audited(client, tresorerie, account, db):
    rows = [HEADER, ["24/09/2026", None, "VIR RECU", None, 10, "abc"]]
    content = xlsx(("Relevé", rows))
    lines = submitted(client, tresorerie, account, content)
    lines[0]["solde"] = "1010.00"

    response = confirm(client, tresorerie, account.id, content, lignes=lines)

    assert response.status_code == 201, response.text
    [row] = transactions(db, account)
    assert (row.origine, row.solde) == ("Corrigée", Decimal("1010.00"))
    [entry] = db.scalars(select(AuditLog).filter_by(action="import_releve")).all()
    assert {
        "numero": 2,
        "champ": "erreurs_du_fichier",
        "avant": "Solde : Montant illisible : « abc ».",
        "apres": None,
    } in entry.nouvelle_valeur["lignes_corrigees"]


def test_line_of_another_bank_can_never_be_imported(client, tresorerie, account, db):
    rows = [["Banque", "Date", "Libellé", "Débit"], ["BMCE", "24/09/2026", "AUTRE", 1]]
    content = xlsx(("Relevé", rows))
    lines = submitted(client, tresorerie, account, content)
    lines[0]["libelle"] = "AUTRE CORRIGE"

    response = confirm(client, tresorerie, account.id, content, lignes=lines)

    assert response.status_code == 409
    assert response.json() == {
        "detail": "Ligne 2 : Banque « BMCE » différente de celle du compte (CIH)."
    }


def test_corrected_file_error_is_recognised_when_the_bank_resends_it(
    client, tresorerie, account, db
):
    content = xlsx(("Relevé", EDITABLE))
    lines = submitted(client, tresorerie, account, content)
    lines[1] |= {"debit": "12500.50", "credit": "0.00"}
    confirm(client, tresorerie, account.id, content, lignes=lines)
    # La banque renvoie la même opération, cette fois lisible : même valeurs que la correction
    resent = [HEADER, ["25/09/2026", None, "CHQ FOURNISSEUR", "12 500,50", None, 1037499.5]]

    preview = analyse(client, tresorerie, account.id, xlsx(("Relevé", resent))).json()

    assert preview["lignes"][0]["motifs"] == ["Déjà importée pour ce compte."]


def test_large_statement_can_be_confirmed_with_its_lines(client, tresorerie, account, db):
    rows = [HEADER] + [
        [f"{1 + index % 28:02d}/09/2026", None, f"VIR CLIENT {index}", None, 10, None]
        for index in range(1100)
    ]
    content = xlsx(("Relevé", rows))
    lines = submitted(client, tresorerie, account, content)
    for line in lines:
        line["commentaire"] = "x" * 990  # plus de 1 Mo de lignes envoyées

    response = confirm(client, tresorerie, account.id, content, lignes=lines)

    assert response.status_code == 201, response.text
    assert response.json()["nb_importees"] == 1100


def test_operations_expose_their_pointage_id(client, tresorerie, account, db):
    transaction = first_transaction(client, tresorerie, account, db)

    body = continuous(client, tresorerie, account.id).json()

    assert body["operations"][0]["pointage_type_id"] == transaction.pointage_type_id


def test_unchanged_inactive_pointage_does_not_block_other_changes(client, tresorerie, account, db):
    transaction = first_transaction(client, tresorerie, account, db)
    current = db.get(PointageType, transaction.pointage_type_id)
    current.actif = False
    db.flush()

    response = patch(
        client, tresorerie, transaction.id, {"pointage_type_id": current.id, "commentaire": "Vu"}
    )

    assert response.status_code == 200, response.text
    assert (response.json()["pointage"], response.json()["commentaire"]) == (current.libelle, "Vu")


def test_xls_statement_is_analysed_like_the_xlsx(client, tresorerie, account):
    as_xlsx = analyse(client, tresorerie, account.id, xlsx(("Relevé", STATEMENT))).json()
    response = analyse(
        client, tresorerie, account.id, xls(("Relevé", STATEMENT)), nom="RELEVE_BP.xls"
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["mapping"] == as_xlsx["mapping"]
    assert [(line["statut"], line["debit"], line["credit"]) for line in body["lignes"]] == [
        (line["statut"], line["debit"], line["credit"]) for line in as_xlsx["lignes"]
    ]


# --- Lignes ignorées et leur raison (08/10/2026) ----------------------------------------------------


def test_ignored_lines_come_back_with_their_reason(client, tresorerie, account):
    rows = [
        HEADER,
        ["01/09/2026", None, "SOLDE INITIAL", None, None, 120000],
        [date(2026, 9, 2), None, "VIR CLIENT", None, 1000, 121000],
        [None, None, "Sous-titre du relevé", None, None, None],
        [None, None, "TOTAL", None, 1000, None],
        [None, None, "SOLDE FINAL", None, None, 121000],
    ]

    body = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()

    assert [(line["numero"], line["raison"]) for line in body["lignes_ignorees"]] == [
        (2, "Ligne SOLDE INITIAL : donne le solde d'ouverture du fichier"),
        (4, "Ligne de titre ou sans montant"),
        (5, "Ligne de total ou de solde"),
        (6, "Ligne SOLDE FINAL : donne le solde de clôture du fichier"),
    ]
    assert body["lignes_ignorees"][2]["cellules"] == ["", "", "TOTAL", "", "1000"]
    assert body["resume"]["nb_ignorees"] == len(body["lignes_ignorees"])
    assert [line["numero"] for line in body["lignes"]] == [3]


# --- Catégories de pointage et mémoire (08/10/2026) -----------------------------------------------


def test_pointage_is_remembered_per_label_within_the_company(client, tresorerie, account, db):
    """Un libellé déjà pointé par quelqu'un reprend le même pointage, avant les mots-clés."""
    first = [["Date", "Libellé", "Débit", "Crédit"], ["01/09/2026", "PRLV ONEE  ", 500, None]]
    confirm(client, tresorerie, account.id, xlsx(("Relevé", first)))
    [done] = transactions(db, account)
    done.pointage_type_id = pointage_id(db, "REDEVANCE")  # choisi par un utilisateur
    db.flush()
    rows = [
        ["Date", "Libellé", "Débit", "Crédit"],
        ["02/09/2026", "prlv onee", 600, None],  # même libellé, casse et espaces près
        ["02/09/2026", "PRLV ONEE TVA", 700, None],  # autre libellé : mot-clé
    ]

    preview = analyse(client, tresorerie, account.id, xlsx(("Relevé", rows))).json()["lignes"]

    assert [(line["pointage_libelle"], line["pointage_auto"]) for line in preview] == [
        ("REDEVANCE", True),
        ("TVA", True),
    ]


def test_pointage_memory_never_crosses_companies(client, tresorerie, account, db):
    other_company = db.scalar(select(Company).filter_by(code="SOCX"))
    other = save(
        db,
        build_account(other_company, account.bank, numero=f"RIB-SOCX-{next(_numeros):06d}"),
    )
    first = [["Date", "Libellé", "Débit", "Crédit"], ["01/09/2026", "PRLV ONEE", 500, None]]
    confirm(client, tresorerie, other.id, xlsx(("Relevé", first)))
    [done] = transactions(db, other)
    done.pointage_type_id = pointage_id(db, "REDEVANCE")
    db.flush()

    preview = analyse(client, tresorerie, account.id, xlsx(("Relevé", first))).json()["lignes"]

    assert preview[0]["pointage_libelle"] is None


# --- Gros fichiers : 50 000 lignes au plus (08/10/2026) -------------------------------------------


def _big_statement(rows: int) -> bytes:
    return big_xlsx(
        ["Date", "Libellé", "Débit", "Crédit"],
        rows,
        lambda index: ["01/09/2026", f"VIR CLIENT {index}", None, index + 1],
    )


def test_a_statement_of_50000_lines_is_analysed(client, tresorerie, account):
    response = analyse(client, tresorerie, account.id, _big_statement(50_000))

    assert response.status_code == 200, response.text
    body = response.json()
    assert len(body["lignes"]) == 50_000
    assert {line["statut"] for line in body["lignes"]} == {"Valide"}


def test_a_statement_of_50001_lines_is_refused(client, tresorerie, account):
    response = analyse(client, tresorerie, account.id, _big_statement(50_001))

    assert response.status_code == 409
    assert response.json() == {"detail": "Fichier trop long : 50 000 lignes au plus par relevé."}


def test_already_imported_lines_are_found_across_hash_batches(
    client, tresorerie, account, monkeypatch
):
    """Les empreintes sont cherchées par paquets : un doublon est trouvé dans chacun."""
    monkeypatch.setattr(import_repository, "HASH_BATCH", 2)
    content = _big_statement(5)
    confirm(client, tresorerie, account.id, content)

    body = analyse(client, tresorerie, account.id, content).json()

    assert [line["motifs"] for line in body["lignes"]] == [["Déjà importée pour ce compte."]] * 5


# --- Relevés sans soldes : solde calculé (08/10/2026) ---------------------------------------------

NO_BALANCE = [
    ["Date", "Libellé", "Débit", "Crédit"],
    ["02/09/2026", "VIR CLIENT", None, 1000],
    ["03/09/2026", "CHQ 12", 250, None],
    ["03/09/2026", "VIR CLIENT 2", None, 50],
]


def test_balances_are_computed_from_the_opening_balance(client, tresorerie, account, db):
    body = analyse(
        client, tresorerie, account.id, xlsx(("Relevé", NO_BALANCE)), solde_ouverture="5 000 000"
    ).json()

    assert body["soldes_calcules"] is True
    assert [line["solde_apercu"] for line in body["lignes"]] == [
        "5001000.00",
        "5000750.00",
        "5000800.00",
    ]
    assert {line["solde"] for line in body["lignes"]} == {None}  # jamais renvoyé par l'écran
    assert (body["resume"]["solde_ouverture"], body["resume"]["solde_cloture"]) == (
        "5000000.00",
        "5000800.00",
    )

    confirm(client, tresorerie, account.id, xlsx(("Relevé", NO_BALANCE)), solde_ouverture="5000000")

    saved = transactions(db, account)
    assert [(row.solde, row.solde_calcule) for row in saved] == [
        (Decimal("5001000.00"), True),
        (Decimal("5000750.00"), True),
        (Decimal("5000800.00"), True),
    ]
    assert day_balance(db, account, date(2026, 9, 3)).solde == Decimal("5000800.00")


def test_a_statement_from_newest_to_oldest_is_computed_in_date_order(client, tresorerie, account):
    rows = [NO_BALANCE[0], NO_BALANCE[3], NO_BALANCE[2], NO_BALANCE[1]]

    body = analyse(
        client, tresorerie, account.id, xlsx(("Relevé", rows)), solde_ouverture="0"
    ).json()

    assert [(line["libelle"], line["solde_apercu"]) for line in body["lignes"]] == [
        ("VIR CLIENT 2", "800.00"),
        ("CHQ 12", "750.00"),
        ("VIR CLIENT", "1000.00"),
    ]


def test_without_any_known_balance_the_opening_must_be_entered(client, tresorerie, account):
    content = xlsx(("Relevé", NO_BALANCE))

    body = analyse(client, tresorerie, account.id, content).json()
    response = confirm(client, tresorerie, account.id, content, solde_ouverture=None)

    assert (body["solde_ouverture_propose"], body["solde_ouverture_source"]) == (None, "À saisir")
    assert {line["solde_apercu"] for line in body["lignes"]} == {None}
    assert response.status_code == 409
    assert "Indiquez le solde d'ouverture" in response.json()["detail"]


def test_the_opening_comes_from_the_last_imported_operation(client, tresorerie, account):
    first = [["Date", "Libellé", "Crédit", "Solde"], ["01/09/2026", "VIR", 10, 4200]]
    confirm(client, tresorerie, account.id, xlsx(("Relevé", first)))

    body = analyse(client, tresorerie, account.id, xlsx(("Relevé", NO_BALANCE))).json()

    assert body["solde_ouverture_propose"] == "4200.00"
    assert (
        body["solde_ouverture_source"] == "Solde de la dernière opération importée, le 01/09/2026"
    )
    assert body["lignes"][0]["solde_apercu"] == "5200.00"


def test_the_opening_comes_from_a_saved_balance_otherwise(client, tresorerie, account, db):
    save(
        db,
        BankAccountBalance(
            bank_account_id=account.id,
            date_solde=date(2026, 8, 31),
            solde=Decimal("700"),
            source="Saisie",
        ),
    )

    body = analyse(client, tresorerie, account.id, xlsx(("Relevé", NO_BALANCE))).json()

    assert body["solde_ouverture_source"] == "Solde du jour enregistré le 31/08/2026"
    assert body["lignes"][0]["solde_apercu"] == "1700.00"


def test_the_file_opening_line_wins_and_the_user_can_override_it(client, tresorerie, account):
    with_opening = [
        ["Date", "Libellé", "Débit", "Crédit", "Solde"],
        [None, "SOLDE INITIAL", None, None, 300],
        ["02/09/2026", "VIR CLIENT", None, 1000, None],
    ]

    proposed = analyse(client, tresorerie, account.id, xlsx(("Relevé", with_opening))).json()
    overridden = analyse(
        client, tresorerie, account.id, xlsx(("Relevé", with_opening)), solde_ouverture="100"
    ).json()

    assert proposed["solde_ouverture_source"] == "Ligne SOLDE INITIAL du fichier"
    assert proposed["lignes"][0]["solde_apercu"] == "1300.00"
    assert overridden["lignes"][0]["solde_apercu"] == "1100.00"


def test_balances_of_the_file_are_never_recomputed(client, tresorerie, account):
    body = analyse(
        client, tresorerie, account.id, xlsx(("Relevé", STATEMENT)), solde_ouverture="0"
    ).json()

    assert body["soldes_calcules"] is False
    assert body["lignes"][0]["solde"] == "1050000.00"


def test_unreadable_opening_balance_gives_422(client, tresorerie, account):
    response = analyse(
        client, tresorerie, account.id, xlsx(("Relevé", NO_BALANCE)), solde_ouverture="beaucoup"
    )

    assert response.status_code == 422


def test_computed_balances_do_not_change_the_line_or_its_fingerprint(
    client, tresorerie, account, db
):
    """Les lignes renvoyées par l'écran (sans solde) ne sont pas « corrigées » et le même fichier
    réimporté est reconnu comme déjà importé."""
    content = xlsx(("Relevé", NO_BALANCE))
    lines = submitted(client, tresorerie, account, content)

    confirm(client, tresorerie, account.id, content, lignes=lines, solde_ouverture="100")
    again = xlsx(("Relevé", [*NO_BALANCE, ["04/09/2026", "VIR NOUVEAU", None, 5]]))
    body = analyse(client, tresorerie, account.id, again).json()

    assert {row.origine for row in transactions(db, account)} == {"Fichier"}
    assert [line["statut"] for line in body["lignes"]] == ["Doublon"] * 3 + ["Valide"]
    # Le nouveau départ est le dernier solde calculé enregistré (100 + 800)
    assert body["solde_ouverture_propose"] == "900.00"
