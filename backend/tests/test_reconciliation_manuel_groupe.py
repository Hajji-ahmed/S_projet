"""Rapprochement manuel de plusieurs lignes (P13, 15/10/2026) : 1-N, N-1, N-N."""

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import AuditLog, BankTransaction, ReconciliationMatch
from app.services.reconciliation_groupes import type_ensemble
from tests.helpers import bearer, login, make_auth_user
from tests.test_reconciliation_api import account, company, ecriture, operation


@pytest.fixture
def comptable(client, reference) -> dict[str, str]:
    make_auth_user(reference, "COMPTABLE", email="comptable@example.com")
    return bearer(login(client, "comptable@example.com"))


def grouper(client, headers, operations, ecritures, **over):
    body = {
        "transaction_ids": [tx.id for tx in operations],
        "ecriture_ids": [e.id for e in ecritures],
        **over,
    }
    return client.post("/api/reconciliation/matches", json=body, headers=headers)


@pytest.mark.parametrize(
    ("nb_ops", "nb_ecr", "attendu"),
    [(1, 1, "1-1"), (1, 3, "1-N"), (2, 1, "N-1"), (2, 2, "N-N")],
)
def test_the_type_follows_the_number_of_lines(nb_ops, nb_ecr, attendu):
    assert type_ensemble(nb_ops, nb_ecr) == attendu


def test_one_operation_for_three_entries_is_validated_with_a_code(client, comptable, db):
    compte = account(db)
    remise = operation(db, compte, montant="12000", libelle="REMISE CHEQUES 0173813")
    cheques = [
        ecriture(db, compte, montant="-4000", libelle=f"Client {n}", numero_piece=f"P{n}")
        for n in range(3)
    ]

    response = grouper(client, comptable, [remise], cheques, commentaire="Remise du 24/09")

    assert response.status_code == 201, response.text
    body = response.json()
    assert (body["type"], body["statut"], body["origine"]) == ("1-N", "Validée", "Manuelle")
    assert body["code"].startswith("AWB-")
    assert len(body["ecritures"]) == 3
    assert (body["total_operations"], body["total_ecritures"], body["ecart"]) == (
        "12000.00",
        "-12000.00",
        "0.00",
    )
    assert "1 opération = 12 000,00 ↔ 3 écritures = 12 000,00" in body["explication"]
    for ligne in [remise, *cheques]:
        db.refresh(ligne)
        assert ligne.statut == "Rapprochée"
    audit = db.scalar(select(AuditLog).where(AuditLog.action == "rapprochement_manuel"))
    assert len(audit.nouvelle_valeur["ecritures"]) == 3


def test_several_operations_for_several_entries(client, comptable, db):
    compte = account(db)
    ops = [
        operation(db, compte, montant="-700", libelle="VIR A"),
        operation(db, compte, montant="-300", libelle="VIR B"),
    ]
    ecrs = [
        ecriture(db, compte, montant="600", libelle="Fournisseur X", numero_piece="F1"),
        ecriture(db, compte, montant="400", libelle="Fournisseur Y", numero_piece="F2"),
    ]

    response = grouper(client, comptable, ops, ecrs)

    assert response.status_code == 201, response.text
    assert response.json()["type"] == "N-N"
    assert len(response.json()["operations"]) == 2

    # Annuler un N-N libère toutes ses lignes
    annule = client.request(
        "DELETE",
        f"/api/reconciliation/matches/{response.json()['id']}",
        json={"motif": "Erreur"},
        headers=comptable,
    )
    assert annule.status_code == 200, annule.text
    for ligne in [*ops, *ecrs]:
        db.refresh(ligne)
        assert ligne.statut == "Non rapprochée"


def test_totals_must_balance_and_nothing_changes_otherwise(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte, montant="1000")
    ecrs = [
        ecriture(db, compte, montant="-600", numero_piece="A"),
        ecriture(db, compte, montant="-300", numero_piece="B"),
    ]

    response = grouper(client, comptable, [tx], ecrs)

    assert response.status_code == 409
    assert "écart de 100.00" in response.json()["detail"]
    assert db.scalar(select(ReconciliationMatch)) is None
    db.refresh(tx)
    assert tx.statut == "Non rapprochée"


def test_lines_of_two_accounts_are_refused(client, comptable, db):
    awb, bmce = account(db), account(db, bank_code="BMCE")
    ops = [operation(db, awb, montant="500"), operation(db, bmce, montant="500")]
    entry = ecriture(db, awb, montant="-1000")

    response = grouper(client, comptable, ops, [entry])

    assert response.status_code == 409
    assert "plusieurs comptes" in response.json()["detail"]


def test_an_already_reconciled_line_is_refused(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte, montant="1000")
    premiere = ecriture(db, compte, montant="-1000", numero_piece="A")
    assert grouper(client, comptable, [tx], [premiere]).status_code == 201
    autre = operation(db, compte, montant="500", libelle="VIR RECU")
    seconde = ecriture(db, compte, montant="-500", numero_piece="B")

    # Totaux équilibrés (1 500), mais l'opération et la première écriture sont déjà rapprochées
    response = grouper(client, comptable, [tx, autre], [premiere, seconde])

    assert response.status_code == 409
    assert "déjà rapprochée" in response.json()["detail"]


def test_mixed_types_are_refused(client, comptable, db):
    compte = account(db)
    cheque = operation(db, compte, montant="1000", libelle="ENCAISSEMENT CHEQUE N 1234567")
    ecrs = [
        ecriture(db, compte, montant="-500", libelle="EAR1° AXK 1234567", numero_piece="A"),
        ecriture(db, compte, montant="-500", libelle="ESP CLIENT", numero_piece="B"),
    ]

    response = grouper(client, comptable, [cheque], ecrs)

    assert response.status_code == 409
    assert "types différents" in response.json()["detail"]


def test_limits_and_payload_forms(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte, montant="1000")
    entry = ecriture(db, compte, montant="-1000")
    beaucoup = [
        operation(db, compte, montant="1", libelle=f"VIR {n}", jour=date(2026, 9, 1))
        for n in range(51)
    ]

    trop = grouper(client, comptable, beaucoup, [entry])
    vide = client.post(
        "/api/reconciliation/matches",
        json={"transaction_ids": [], "ecriture_ids": [entry.id]},
        headers=comptable,
    )
    melange = client.post(
        "/api/reconciliation/matches",
        json={"transaction_id": tx.id, "ecriture_ids": [entry.id]},
        headers=comptable,
    )

    assert trop.status_code == 409 and "au plus 50" in trop.json()["detail"]
    assert vide.status_code == 409
    assert melange.status_code == 422
    assert db.scalar(select(BankTransaction.statut).where(BankTransaction.id == tx.id)) == (
        "Non rapprochée"
    )


def test_the_lettrage_export_lists_every_line_of_a_group(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte, montant="900")
    ecrs = [
        ecriture(db, compte, montant="-400", numero_piece="A"),
        ecriture(db, compte, montant="-500", numero_piece="B"),
    ]
    code = grouper(client, comptable, [tx], ecrs).json()["code"]

    response = client.get(
        "/api/reconciliation/export",
        params={"company_id": str(company(db).id), "from": "2026-09-01", "to": "2026-09-30"},
        headers=comptable,
    )

    rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
    lignes = [r for r in rows[1:] if r[0] == code]
    assert [r[1] for r in lignes] == ["Banque", "Sage", "Sage"]
    assert sum(Decimal(str(r[7] or 0)) for r in lignes) == sum(
        Decimal(str(r[8] or 0)) for r in lignes
    )
