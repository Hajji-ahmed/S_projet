"""Rapprochement 1→1 (P11) : lancement, décisions, lectures, permissions et audit."""

from datetime import date
from decimal import Decimal
from itertools import count

import pytest
from sqlalchemy import select

from app.models import (
    AccountingEntry,
    AuditLog,
    Bank,
    BankAccount,
    BankStatement,
    BankTransaction,
    Company,
    ReconciliationMatch,
    ReconciliationMatchItem,
    ReconciliationRule,
)
from app.schemas.reconciliation import criteres
from app.services.position_service import business_today
from app.services.reconciliation_service import grille
from tests.helpers import bearer, build_account, login, make_auth_user, save

_seq = count(1)
JOUR = date(2026, 9, 24)
PERIODE = {"du": "2026-09-01", "au": "2026-09-30"}


@pytest.fixture
def comptable(client, reference) -> dict[str, str]:
    make_auth_user(reference, "COMPTABLE", email="comptable@example.com")
    return bearer(login(client, "comptable@example.com"))


@pytest.fixture
def direction(client, reference) -> dict[str, str]:
    make_auth_user(reference, "DIRECTION", email="direction@example.com")
    return bearer(login(client, "direction@example.com"))


def company(db, code: str = "SIMTIS") -> Company:
    return db.scalar(select(Company).filter_by(code=code))


def account(db, company_code: str = "SIMTIS", bank_code: str = "AWB") -> BankAccount:
    bank = db.scalar(select(Bank).filter_by(code=bank_code))
    return save(
        db, build_account(company(db, company_code), bank, numero=f"RIB-RAP-{next(_seq):06d}")
    )


def operation(db, compte: BankAccount, montant="50000", jour=JOUR, libelle="VIR ABC", **over):
    statement = db.scalar(select(BankStatement).filter_by(bank_account_id=compte.id))
    if statement is None:
        statement = save(db, BankStatement(bank_account_id=compte.id))
    montant = Decimal(montant)
    return save(
        db,
        BankTransaction(
            statement_id=statement.id,
            bank_account_id=compte.id,
            date_operation=jour,
            libelle=libelle,
            debit=max(-montant, Decimal(0)),
            credit=max(montant, Decimal(0)),
            montant=montant,
            hash_ligne=f"t{next(_seq)}",
            **over,
        ),
    )


def ecriture(
    db, compte: BankAccount, montant="-50000", jour=JOUR, libelle="Règlement ABC", **over
) -> AccountingEntry:
    """Au sens de Sage : un montant négatif est un crédit du compte banque (argent qui sort)."""
    montant = Decimal(montant)
    values = {"numero_piece": "REG458", "tiers": "ABC", **over}
    return save(
        db,
        AccountingEntry(
            company_id=compte.company_id,
            bank_account_id=compte.id,
            journal="BQ1",
            compte="5141",
            date_ecriture=jour,
            libelle=libelle,
            debit=max(-montant, Decimal(0)),
            credit=max(montant, Decimal(0)),
            montant=montant,
            hash_ligne=f"e{next(_seq)}",
            **values,
        ),
    )


def run(client, headers, db, code: str = "SIMTIS", **over):
    body = {"company_id": company(db, code).id, **PERIODE, **over}
    return client.post("/api/reconciliation/run", json=body, headers=headers)


def proposals(client, headers, db, **params):
    params = {"company_id": str(company(db).id), **params}
    return client.get("/api/reconciliation/proposals", params=params, headers=headers)


def statuses(db, *rows) -> list[str]:
    for row in rows:
        db.refresh(row)
    return [row.statut for row in rows]


def audit_actions(db) -> list[str]:
    return list(db.scalars(select(AuditLog.action).order_by(AuditLog.id)))


# --- Lancement du moteur --------------------------------------------------------------------------


def test_example_of_the_functional_architecture_is_proposed_then_validated(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte)
    entry = ecriture(db, compte)

    response = run(client, comptable, db)

    assert response.status_code == 200, response.text
    assert response.json()["nb_propositions"] == 1
    assert statuses(db, tx, entry) == ["À vérifier", "À vérifier"]
    [proposal] = proposals(client, comptable, db).json()["correspondances"]
    assert (proposal["statut"], proposal["origine"], proposal["score"]) == (
        "Proposée",
        "Automatique",
        "100.00",
    )
    # Forte, mais jamais validée sans un utilisateur
    assert proposal["forte"] is True
    # N° de chèque / référence, montant, date et libellé (décisions du 07 et du 08/10/2026)
    assert [(c["code"], c["points"]) for c in proposal["criteres"]] == [
        ("reference", "0.00"),
        ("montant", "50.00"),
        ("date", "30.00"),
        ("libelle", "20.00"),
    ]
    assert proposal["operation"]["id"] == tx.id
    assert proposal["ecriture"]["numero_piece"] == "REG458"

    validated = client.post(
        f"/api/reconciliation/matches/{proposal['id']}/validate", headers=comptable
    )

    assert validated.status_code == 200, validated.text
    assert validated.json()["statut"] == "Validée"
    assert validated.json()["valide_par"]
    assert statuses(db, tx, entry) == ["Rapprochée", "Rapprochée"]
    assert audit_actions(db)[-2:] == ["rapprochement_lance", "validation_rapprochement"]


def test_the_engine_never_validates_by_itself(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte, libelle="VIR REG458 ABC")
    ecriture(db, compte)

    run(client, comptable, db)

    [proposal] = proposals(client, comptable, db).json()["correspondances"]
    assert proposal["score"] == "100.00"
    assert proposal["forte"] is True
    assert proposal["statut"] == "Proposée"
    assert statuses(db, tx) == ["À vérifier"]


def test_close_candidates_give_no_proposal_and_lines_to_check(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte)
    first, second = ecriture(db, compte), ecriture(db, compte)

    body = run(client, comptable, db).json()

    assert (body["nb_propositions"], body["nb_operations_ambigues"]) == (0, 1)
    assert statuses(db, tx, first, second) == ["À vérifier"] * 3


def test_entries_of_another_company_are_never_compared(client, comptable, db):
    simtis, tefil = account(db, "SIMTIS"), account(db, "SOCX")
    tx = operation(db, simtis)
    ecriture(db, tefil)

    assert run(client, comptable, db).json()["nb_propositions"] == 0
    assert statuses(db, tx) == ["Non rapprochée"]


def test_entries_of_another_bank_account_are_not_compared(client, comptable, db):
    awb, bp = account(db, bank_code="AWB"), account(db, bank_code="BP")
    operation(db, awb)
    ecriture(db, bp)

    assert run(client, comptable, db).json()["nb_propositions"] == 0


def test_running_again_replaces_pending_proposals_and_keeps_decisions(client, comptable, db):
    compte = account(db)
    operation(db, compte, jour=date(2026, 9, 10))
    ecriture(db, compte, jour=date(2026, 9, 10))
    kept_tx = operation(db, compte, montant="700", libelle="VIR XYZ")
    ecriture(db, compte, montant="-700", libelle="Règlement XYZ", tiers="XYZ", numero_piece="P7")
    run(client, comptable, db)
    pending = proposals(client, comptable, db).json()["correspondances"]
    to_keep = next(p for p in pending if p["operation"]["id"] == kept_tx.id)
    client.post(f"/api/reconciliation/matches/{to_keep['id']}/validate", headers=comptable)

    second = run(client, comptable, db).json()

    assert second["nb_propositions"] == 1
    assert len(proposals(client, comptable, db).json()["correspondances"]) == 1
    validated = proposals(client, comptable, db, statut="Validée").json()["correspondances"]
    assert [p["id"] for p in validated] == [to_keep["id"]]


def test_rejected_pair_is_not_proposed_again(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte)
    entry = ecriture(db, compte)
    run(client, comptable, db)
    [proposal] = proposals(client, comptable, db).json()["correspondances"]

    rejected = client.post(
        f"/api/reconciliation/matches/{proposal['id']}/reject",
        json={"commentaire": "Pas le même client"},
        headers=comptable,
    )

    assert rejected.status_code == 200, rejected.text
    assert rejected.json()["statut"] == "Rejetée"
    assert statuses(db, tx, entry) == ["Non rapprochée", "Non rapprochée"]
    assert run(client, comptable, db).json()["nb_propositions"] == 0
    assert statuses(db, tx) == ["Non rapprochée"]


def test_run_checks_period_and_account(client, comptable, db):
    other = account(db, "SOCX")

    inverted = run(client, comptable, db, du="2026-09-30", au="2026-09-01")
    # 5 ans au plus par lancement (08/10/2026) : plus d'un an passe, une année mal saisie non
    over_a_year = run(client, comptable, db, du="2025-01-01", au="2026-09-30")
    five_years = run(client, comptable, db, du="2021-10-01", au="2026-09-30")
    too_long = run(client, comptable, db, du="0026-01-01", au="2026-09-30")
    foreign = run(client, comptable, db, bank_account_id=other.id)

    assert inverted.status_code == 409
    assert over_a_year.status_code == 200, over_a_year.text
    assert five_years.status_code == 200, five_years.text
    assert too_long.status_code == 409
    assert too_long.json() == {"detail": "La période ne peut pas dépasser 5 ans."}
    assert foreign.status_code == 404


# --- Décisions -----------------------------------------------------------------------------------


def _proposed(client, headers, db, **over):
    compte = account(db)
    tx = operation(db, compte, **over)
    entry = ecriture(db, compte)
    run(client, headers, db)
    [proposal] = proposals(client, headers, db).json()["correspondances"]
    return proposal, tx, entry


def test_a_decision_needs_a_pending_proposal(client, comptable, db):
    proposal, _, _ = _proposed(client, comptable, db)
    url = f"/api/reconciliation/matches/{proposal['id']}"
    client.post(f"{url}/validate", headers=comptable)

    assert client.post(f"{url}/validate", headers=comptable).status_code == 409
    assert client.post(f"{url}/reject", headers=comptable).status_code == 409
    assert (
        client.post("/api/reconciliation/matches/999999/validate", headers=comptable).status_code
        == 404
    )


def test_cancelling_a_validated_match_needs_a_reason_and_frees_the_lines(client, comptable, db):
    proposal, tx, entry = _proposed(client, comptable, db)
    url = f"/api/reconciliation/matches/{proposal['id']}"

    pending = client.request("DELETE", url, json={"motif": "Erreur"}, headers=comptable)
    client.post(f"{url}/validate", headers=comptable)
    empty = client.request("DELETE", url, json={"motif": ""}, headers=comptable)
    blank = client.request("DELETE", url, json={"motif": "   "}, headers=comptable)
    cancelled = client.request(
        "DELETE", url, json={"motif": "Mauvaise écriture"}, headers=comptable
    )

    assert pending.status_code == 409
    assert empty.status_code == 422
    assert blank.status_code == 409
    assert cancelled.status_code == 200, cancelled.text
    assert (cancelled.json()["statut"], cancelled.json()["commentaire"]) == (
        "Annulée",
        "Mauvaise écriture",
    )
    assert statuses(db, tx, entry) == ["Non rapprochée", "Non rapprochée"]
    assert audit_actions(db)[-1] == "annulation_rapprochement"


def _strong_and_weak(client, comptable, db):
    compte = account(db)
    strong_tx = operation(db, compte, libelle="VIR REG458 ABC")
    ecriture(db, compte)
    # Même montant, deux jours d'écart, libellés différents : 50 + 15 + quelques points, < 90
    weak_tx = operation(db, compte, montant="700", jour=date(2026, 9, 10), libelle="VIR XYZ")
    ecriture(
        db,
        compte,
        montant="-700",
        jour=date(2026, 9, 12),
        libelle="Paiement divers",
        tiers=None,
        numero_piece="P7",
    )
    run(client, comptable, db)
    pending = proposals(client, comptable, db).json()["correspondances"]
    strong = [p["id"] for p in pending if p["forte"]]
    weak = [p["id"] for p in pending if not p["forte"]]
    assert (len(strong), len(weak)) == (1, 1)
    return compte, strong_tx, weak_tx, strong, weak


BATCH = "/api/reconciliation/matches/validate-batch"


def test_batch_validation_accepts_the_selection_strong_and_weak(client, comptable, db):
    _, strong_tx, weak_tx, strong, weak = _strong_and_weak(client, comptable, db)

    accepted = client.post(BATCH, json={"ids": strong + weak}, headers=comptable)

    assert accepted.status_code == 200, accepted.text
    assert accepted.json() == {"nb_validees": 2}
    assert statuses(db, strong_tx, weak_tx) == ["Rapprochée", "Rapprochée"]
    assert proposals(client, comptable, db).json()["correspondances"] == []
    logs = db.scalars(
        select(AuditLog).where(AuditLog.action == "validation_rapprochement").order_by(AuditLog.id)
    ).all()
    assert [(log.nouvelle_valeur["en_lot"], log.nouvelle_valeur["forte"]) for log in logs] == [
        (True, True),
        (True, False),
    ]


def test_batch_validation_is_all_or_nothing(client, comptable, db):
    compte, strong_tx, weak_tx, strong, weak = _strong_and_weak(client, comptable, db)
    client.post(f"/api/reconciliation/matches/{weak[0]}/reject", headers=comptable)

    refused = client.post(BATCH, json={"ids": strong + weak}, headers=comptable)

    assert refused.status_code == 409
    assert f"n° {weak[0]}" in refused.json()["detail"]
    assert statuses(db, strong_tx) == ["À vérifier"]


def test_batch_validation_refuses_different_amounts(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte, libelle="VIR REG458 ABC")
    ecriture(db, compte, montant="-49000")
    run(client, comptable, db)
    [proposal] = proposals(client, comptable, db).json()["correspondances"]

    refused = client.post(BATCH, json={"ids": [proposal["id"]]}, headers=comptable)

    assert refused.status_code == 409
    assert "Montant différent" in refused.json()["detail"]
    assert statuses(db, tx) == ["À vérifier"]


def test_batch_validation_needs_the_permission(client, direction, comptable, db):
    *_, strong, weak = _strong_and_weak(client, comptable, db)

    assert client.post(BATCH, json={"ids": strong}, headers=direction).status_code == 403


# --- Rapprochement manuel ------------------------------------------------------------------------


def manual(client, headers, tx, entry, **over):
    body = {"transaction_id": tx.id, "ecriture_id": entry.id, **over}
    return client.post("/api/reconciliation/matches", json=body, headers=headers)


def test_manual_match_is_validated_at_once(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte, libelle="ONEE FACTURE")
    entry = ecriture(db, compte, jour=date(2026, 9, 20), libelle="Electricité", tiers=None)

    response = manual(client, comptable, tx, entry, commentaire="Facture de septembre")

    assert response.status_code == 201, response.text
    body = response.json()
    assert (body["statut"], body["origine"], body["commentaire"]) == (
        "Validée",
        "Manuelle",
        "Facture de septembre",
    )
    assert statuses(db, tx, entry) == ["Rapprochée", "Rapprochée"]
    assert audit_actions(db)[-1] == "rapprochement_manuel"


def test_manual_match_refuses_different_amounts_companies_and_accounts(client, comptable, db):
    awb, bp, tefil = account(db), account(db, bank_code="BP"), account(db, "SOCX")
    tx = operation(db, awb)

    amount = manual(client, comptable, tx, ecriture(db, awb, montant="-49000"))
    same_sign = manual(client, comptable, tx, ecriture(db, awb, montant="50000"))
    other_account = manual(client, comptable, tx, ecriture(db, bp))
    other_company = manual(client, comptable, tx, ecriture(db, tefil))

    assert [r.status_code for r in (amount, same_sign, other_account, other_company)] == [409] * 4
    assert "deux sociétés" in other_company.json()["detail"]


def test_manual_match_replaces_a_pending_proposal(client, comptable, db):
    proposal, tx, proposed_entry = _proposed(client, comptable, db)
    chosen = ecriture(db, account_of(db, tx), jour=date(2026, 9, 23), libelle="Autre")

    response = manual(client, comptable, tx, chosen)

    assert response.status_code == 201, response.text
    replaced = db.get(ReconciliationMatch, proposal["id"])
    db.refresh(replaced)
    assert replaced.statut == "Rejetée"
    assert statuses(db, tx, chosen, proposed_entry) == [
        "Rapprochée",
        "Rapprochée",
        "Non rapprochée",
    ]


def test_manual_match_refuses_an_already_matched_line(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte)
    first, second = ecriture(db, compte), ecriture(db, compte)
    manual(client, comptable, tx, first)

    response = manual(client, comptable, tx, second)

    assert response.status_code == 409
    assert "déjà rapprochée" in response.json()["detail"]


def account_of(db, tx: BankTransaction) -> BankAccount:
    return db.get(BankAccount, tx.bank_account_id)


# --- Lectures ------------------------------------------------------------------------------------


def test_transactions_pane_counts_statuses_and_shows_the_match(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte)
    ecriture(db, compte)
    # Pas un frais bancaire : un frais compterait dans « À comptabiliser » (12/10/2026)
    operation(db, compte, montant="-80", libelle="CHQ 1234 FOURNISSEUR")
    run(client, comptable, db)

    response = client.get(
        "/api/reconciliation/transactions",
        params={"company_id": company(db).id, "bank_account_id": compte.id},
        headers=comptable,
    )

    body = response.json()
    assert response.status_code == 200
    assert body["total"] == 2
    assert body["par_statut"] == {
        "Non rapprochée": 1,
        "À vérifier": 1,
        "Rapprochée": 0,
        "Écart": 0,
    }
    proposed = next(o for o in body["operations"] if o["id"] == tx.id)
    assert proposed["correspondance"]["statut"] == "Proposée"
    assert proposed["bank_code"] == "AWB"


def test_candidates_are_sorted_by_score(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte)
    best = ecriture(db, compte)
    other = ecriture(db, compte, jour=date(2026, 9, 27), libelle="Divers", tiers=None)
    ecriture(db, compte, montant="50000")  # même sens : jamais candidate

    response = client.get(
        "/api/reconciliation/candidates", params={"transaction_id": tx.id}, headers=comptable
    )

    body = response.json()
    assert response.status_code == 200
    assert [c["ecriture"]["id"] for c in body["candidats"]] == [best.id, other.id]
    assert body["seuil_fort"] == "80.00"


# --- Permissions ---------------------------------------------------------------------------------


def test_consultation_can_read_but_not_decide(client, direction, db):
    compte = account(db)
    tx = operation(db, compte)
    entry = ecriture(db, compte)

    assert proposals(client, direction, db).status_code == 200
    assert run(client, direction, db).status_code == 403
    assert manual(client, direction, tx, entry).status_code == 403


def test_one_match_can_be_read_with_its_score_detail(client, direction, comptable, db):
    proposal, tx, _ = _proposed(client, comptable, db)

    response = client.get(f"/api/reconciliation/matches/{proposal['id']}", headers=direction)
    missing = client.get("/api/reconciliation/matches/999999", headers=direction)

    assert response.status_code == 200
    assert response.json()["operation"]["id"] == tx.id
    assert [c["code"] for c in response.json()["criteres"]] == [
        "reference",
        "montant",
        "date",
        "libelle",
    ]
    assert missing.status_code == 404


def test_a_proposal_with_different_amounts_cannot_be_validated(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte, libelle="VIR REG458 ABC")
    entry = ecriture(db, compte, montant="-49000")
    run(client, comptable, db)
    [proposal] = proposals(client, comptable, db).json()["correspondances"]

    response = client.post(
        f"/api/reconciliation/matches/{proposal['id']}/validate", headers=comptable
    )

    assert proposal["criteres"][1] == {"code": "montant", "libelle": "Montant", "points": "0.00"}
    assert response.status_code == 409
    assert "Montant différent" in response.json()["detail"]
    assert statuses(db, tx, entry) == ["À vérifier", "À vérifier"]


# --- Historique ----------------------------------------------------------------------------------


def history(client, headers, db, **params):
    params = {"company_id": str(company(db).id), **params}
    return client.get("/api/reconciliation/history", params=params, headers=headers)


def test_history_lists_every_decision_with_who_and_when(client, comptable, direction, db):
    compte = account(db)
    for montant, jour in (("100", 10), ("200", 11), ("300", 12)):
        operation(db, compte, montant=montant, jour=date(2026, 9, jour), libelle=f"VIR {montant}")
        ecriture(
            db,
            compte,
            montant=f"-{montant}",
            jour=date(2026, 9, jour),
            libelle=f"VIR {montant}",
            numero_piece=f"P{montant}",
            tiers=None,
        )
    run(client, comptable, db)
    pending = {
        p["operation"]["montant"]: p["id"]
        for p in proposals(client, comptable, db).json()["correspondances"]
    }
    url = "/api/reconciliation/matches"
    client.post(f"{url}/{pending['100.00']}/validate", headers=comptable)
    client.post(f"{url}/{pending['200.00']}/reject", json={"commentaire": "Non"}, headers=comptable)
    client.post(f"{url}/{pending['300.00']}/validate", headers=comptable)
    client.request(
        "DELETE", f"{url}/{pending['300.00']}", json={"motif": "Erreur"}, headers=comptable
    )

    response = history(client, direction, db, **{"from": "2026-09-01", "to": "2026-09-30"})

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["total"] == 3
    assert body["par_statut"] == {"Validée": 1, "Rejetée": 1, "Annulée": 1}
    # La plus récente d'abord : l'annulation a été décidée en dernier
    assert [d["statut"] for d in body["decisions"]] == ["Annulée", "Rejetée", "Validée"]
    assert all(d["decide_par"] and d["decide_le"] for d in body["decisions"])
    cancelled = body["decisions"][0]
    assert (cancelled["commentaire"], cancelled["valide_par"]) == (
        "Erreur",
        cancelled["decide_par"],
    )
    assert history(client, direction, db, statut="Rejetée").json()["total"] == 1
    assert proposals(client, comptable, db).json()["correspondances"] == []


def test_history_keeps_manual_matches_and_ignores_pending_proposals(client, comptable, db):
    compte = account(db)
    tx, entry = operation(db, compte, montant="900"), ecriture(db, compte, montant="-900")
    operation(db, compte)
    ecriture(db, compte)
    run(client, comptable, db)
    client.post(
        "/api/reconciliation/matches",
        json={"transaction_id": tx.id, "ecriture_id": entry.id},
        headers=comptable,
    )

    body = history(client, comptable, db).json()

    assert ("Manuelle", "Validée") in [(d["origine"], d["statut"]) for d in body["decisions"]]
    assert "Proposée" not in {d["statut"] for d in body["decisions"]}
    # Une proposition en attente n'est pas une décision : elle n'apparaît pas
    assert body["total"] == sum(body["par_statut"].values())


def test_history_never_mixes_companies_and_checks_rights(client, comptable, reference, db):
    tefil = account(db, "SOCX")
    tx, entry = operation(db, tefil, montant="900"), ecriture(db, tefil, montant="-900")
    client.post(
        "/api/reconciliation/matches",
        json={"transaction_id": tx.id, "ecriture_id": entry.id},
        headers=comptable,
    )
    make_auth_user(reference, email="sans-droit@example.com")
    no_right = bearer(login(client, "sans-droit@example.com"))

    assert history(client, comptable, db).json()["total"] == 0
    assert (
        history(client, comptable, db, company_id=str(company(db, "SOCX").id)).json()["total"] == 1
    )
    assert history(client, no_right, db).status_code == 403


def test_batch_validation_never_mixes_companies(client, comptable, db):
    simtis, tefil = account(db), account(db, "SOCX")
    for compte in (simtis, tefil):
        operation(db, compte, libelle="VIR REG458 ABC")
        ecriture(db, compte)
    run(client, comptable, db)
    run(client, comptable, db, code="SOCX")
    ids = [p["id"] for p in proposals(client, comptable, db).json()["correspondances"]] + [
        p["id"]
        for p in proposals(client, comptable, db, company_id=str(company(db, "SOCX").id)).json()[
            "correspondances"
        ]
    ]

    response = client.post(BATCH, json={"ids": ids}, headers=comptable)

    assert len(ids) == 2
    assert response.status_code == 409
    assert "une seule société" in response.json()["detail"]


# --- Grille ---------------------------------------------------------------------------------------


def test_a_disabled_criterion_is_worth_zero_and_can_be_enabled_again(reference):
    seeded = grille(reference)
    rule = reference.scalar(select(ReconciliationRule).filter_by(code="TIERS"))
    rule.actif = True
    reference.flush()

    # Tiers désactivé (07/10/2026) ; n° de chèque / référence réactivé (08/10/2026)
    assert (seeded.reference, seeded.tiers) == (Decimal("40"), Decimal("0"))
    assert (seeded.montant, seeded.date, seeded.libelle) == (
        Decimal("50"),
        Decimal("30"),
        Decimal("20"),
    )
    assert grille(reference).tiers == Decimal("5")


def test_an_older_score_still_shows_the_points_it_got():
    shown = criteres({"reference": "40.00", "montant": "30.00"}, ("montant", "date", "libelle"))

    assert [(c.code, str(c.points)) for c in shown] == [
        ("reference", "40.00"),
        ("montant", "30.00"),
        ("date", "0.00"),
        ("libelle", "0.00"),
    ]


# --- « À vérifier » regroupé dans Propositions (07/10/2026) ---------------------------------------


def test_pane_can_hide_lines_to_check_but_still_counts_them(client, comptable, db):
    compte = account(db)
    operation(db, compte)
    ecriture(db, compte)
    alone = operation(db, compte, montant="-80", libelle="FRAIS")
    run(client, comptable, db)
    params = {"company_id": company(db).id, "sans_a_verifier": "true"}

    body = client.get("/api/reconciliation/transactions", params=params, headers=comptable).json()

    assert [o["id"] for o in body["operations"]] == [alone.id]
    assert body["total"] == 1
    assert body["par_statut"]["À vérifier"] == 1


def test_ambiguous_lines_come_with_their_candidates(client, comptable, direction, db):
    compte = account(db)
    tx = operation(db, compte)
    first, second = ecriture(db, compte), ecriture(db, compte)
    proposed_tx = operation(db, compte, montant="900", jour=date(2026, 9, 10), libelle="VIR 900")
    ecriture(db, compte, montant="-900", jour=date(2026, 9, 10), libelle="VIR 900")
    run(client, comptable, db)
    url = "/api/reconciliation/ambiguous"

    response = client.get(url, params={"company_id": company(db).id}, headers=direction)

    assert response.status_code == 200, response.text
    [ambigue] = response.json()["ambigues"]
    assert ambigue["operation"]["id"] == tx.id
    assert {c["ecriture"]["id"] for c in ambigue["candidats"]} == {first.id, second.id}
    assert proposed_tx.id != ambigue["operation"]["id"]

    chosen = client.post(
        "/api/reconciliation/matches",
        json={"transaction_id": tx.id, "ecriture_id": first.id},
        headers=comptable,
    )
    after = client.get(url, params={"company_id": company(db).id}, headers=direction).json()

    assert chosen.status_code == 201
    assert statuses(db, tx, first) == ["Rapprochée", "Rapprochée"]
    assert after["ambigues"] == []


def test_ambiguous_lines_never_mix_companies_and_need_the_right(client, comptable, reference, db):
    tefil = account(db, "SOCX")
    operation(db, tefil)
    ecriture(db, tefil)
    ecriture(db, tefil)
    run(client, comptable, db, code="SOCX")
    make_auth_user(reference, email="sans-droit-2@example.com")
    no_right = bearer(login(client, "sans-droit-2@example.com"))
    url = "/api/reconciliation/ambiguous"

    simtis = client.get(url, params={"company_id": company(db).id}, headers=comptable).json()
    socx = client.get(url, params={"company_id": company(db, "SOCX").id}, headers=comptable).json()
    refused = client.get(url, params={"company_id": company(db).id}, headers=no_right)

    assert simtis["ambigues"] == []
    assert len(socx["ambigues"]) == 1
    assert refused.status_code == 403


# --- Ambiguës : vrai total, pagination, relance après un choix manuel (08/10/2026) ----------------


def test_ambiguous_lines_give_the_real_total_and_are_paged(client, comptable, db):
    compte = account(db)
    for numero in range(55):
        operation(db, compte, montant=str(1000 + numero), statut="À vérifier")
    url = "/api/reconciliation/ambiguous"
    params = {"company_id": company(db).id}

    first = client.get(url, params=params, headers=comptable).json()
    second = client.get(url, params={**params, "page": 2}, headers=comptable).json()

    assert (first["total"], first["page"], first["taille"], len(first["ambigues"])) == (
        55,
        1,
        50,
        50,
    )
    assert (second["total"], second["page"], len(second["ambigues"])) == (55, 2, 5)
    seen = {a["operation"]["id"] for a in first["ambigues"] + second["ambigues"]}
    assert len(seen) == 55


def test_a_manual_match_frees_the_competing_operation(client, comptable, db):
    compte = account(db)
    # Deux opérations identiques convoitent la même écriture E1 : aucune proposition
    tx1, tx2 = operation(db, compte), operation(db, compte)
    e1 = ecriture(db, compte)
    e2 = ecriture(db, compte, jour=date(2026, 9, 26))
    body = run(client, comptable, db).json()
    assert body["nb_propositions"] == 0

    response = manual(client, comptable, tx1, e1)

    assert response.status_code == 201, response.text
    [pending] = proposals(client, comptable, db).json()["correspondances"]
    assert (pending["operation"]["id"], pending["ecriture"]["id"]) == (tx2.id, e2.id)
    assert statuses(db, tx1, e1, tx2, e2) == [
        "Rapprochée",
        "Rapprochée",
        "À vérifier",
        "À vérifier",
    ]
    ambigues = client.get(
        "/api/reconciliation/ambiguous", params={"company_id": company(db).id}, headers=comptable
    ).json()
    assert ambigues["total"] == 0
    [log] = db.scalars(select(AuditLog).filter_by(action="rapprochement_manuel")).all()
    assert log.nouvelle_valeur["relance"]["propositions"] == 1


def test_a_strong_proposal_with_a_close_competitor_is_flagged(client, comptable, db):
    compte = account(db)
    operation(db, compte)
    best = ecriture(db, compte)
    ecriture(db, compte, jour=date(2026, 9, 25))
    body = run(client, comptable, db).json()
    assert (body["nb_propositions"], body["nb_operations_ambigues"]) == (1, 0)

    [proposal] = proposals(client, comptable, db).json()["correspondances"]

    assert proposal["ecriture"]["id"] == best.id
    assert (proposal["forte"], proposal["score_second"], proposal["concurrente_proche"]) == (
        True,
        "92.50",
        True,
    )


# --- Validation en lot : jusqu'à 5 000 propositions (09/10/2026) ----------------------------------


def test_a_batch_of_more_than_500_proposals_is_validated(client, comptable, db):
    compte = account(db)
    for i in range(600):
        jour = date(2026, 9, 1 + i % 30)
        operation(db, compte, montant=str(1000 + i), jour=jour, libelle=f"VIR {i}")
        ecriture(db, compte, montant=str(-(1000 + i)), jour=jour, libelle=f"VIR {i}", tiers=None)
    run(client, comptable, db)
    ids = [p["id"] for p in proposals(client, comptable, db).json()["correspondances"]]
    assert len(ids) == 600

    response = client.post(BATCH, json={"ids": ids}, headers=comptable)

    assert response.status_code == 200, response.text
    assert response.json()["nb_validees"] == 600


def test_a_batch_above_the_limit_is_refused(client, comptable):
    response = client.post(BATCH, json={"ids": list(range(1, 10002))}, headers=comptable)

    assert response.status_code == 422


def test_a_manual_match_between_two_types_is_refused(client, comptable, db):
    compte = account(db)
    tx = operation(db, compte, libelle="VERSEMENT DEPLACE PAR UN TIERS")
    entry = ecriture(db, compte, libelle="EAR1° AXK 173813")

    response = manual(client, comptable, tx, entry)

    assert response.status_code == 409
    assert response.json() == {
        "detail": "Une opération Espèces ne se rapproche pas d'une écriture Chèque."
    }


# --- Frais bancaires « À comptabiliser » (12/10/2026) ---------------------------------------------


def test_bank_fees_are_counted_apart_as_to_be_booked(client, comptable, db):
    compte = account(db)
    operation(db, compte, montant="-1", libelle="DROIT DE TIMBRE SUR VERSEMENT")
    operation(db, compte, montant="-2.50", libelle="TAXE SUR VALEUR AJOUTEE")
    operation(db, compte, montant="-25", libelle="COMMISSION")
    operation(db, compte, montant="-900", libelle="PRELEVEMENT ONEE")
    autre = operation(db, compte, montant="50000", libelle="VIR CLIENT ATLAS")
    url = "/api/reconciliation/transactions"
    params = {"company_id": company(db).id}

    tous = client.get(url, params=params, headers=comptable).json()
    frais = client.get(url, params={**params, "a_comptabiliser": "true"}, headers=comptable).json()
    non_rapprochees = client.get(
        url, params={**params, "statut": "Non rapprochée"}, headers=comptable
    ).json()

    assert (tous["a_comptabiliser"], tous["total_a_comptabiliser"]) == (4, "928.50")
    assert tous["par_statut"]["Non rapprochée"] == 1
    assert frais["total"] == 4
    assert [op["id"] for op in non_rapprochees["operations"]] == [autre.id]


def test_a_commercial_transfer_or_documentary_remittance_is_never_a_fee(client, comptable, db):
    """13/10/2026 : leur pointage (INTERET, FRAIS) ne suffit pas, ce ne sont pas des frais ; un
    remboursement d'agios (au crédit) reste à comptabiliser."""
    from app.models import PointageType

    interet = db.scalar(select(PointageType).filter_by(code="INTERET"))
    agios = db.scalar(select(PointageType).filter_by(code="AGIOS"))
    compte = account(db)
    virement = operation(
        db,
        compte,
        montant="-7981.12",
        libelle="VIREMENT COMMERCIAL EMIS",
        pointage_type_id=interet.id,
    )
    remise = operation(
        db,
        compte,
        montant="-8062.45",
        libelle="REMISE DOCUMENTAIRE IMPORT",
        pointage_type_id=interet.id,
    )
    operation(db, compte, montant="151110.79", libelle="AGIOS BANCAIRES", pointage_type_id=agios.id)
    url = "/api/reconciliation/transactions"
    params = {"company_id": company(db).id}

    tous = client.get(url, params=params, headers=comptable).json()
    non_rapprochees = client.get(
        url, params={**params, "statut": "Non rapprochée"}, headers=comptable
    ).json()

    assert (tous["a_comptabiliser"], tous["total_a_comptabiliser"]) == (1, "151110.79")
    assert {op["id"] for op in non_rapprochees["operations"]} == {virement.id, remise.id}


def test_a_reconciled_fee_is_no_longer_to_be_booked(client, comptable, db):
    compte = account(db)
    frais = operation(db, compte, montant="-25", libelle="COMMISSION")
    entry = ecriture(db, compte, montant="25", libelle="COM")

    manual(client, comptable, frais, entry)
    body = client.get(
        "/api/reconciliation/transactions",
        params={"company_id": company(db).id},
        headers=comptable,
    ).json()

    assert body["a_comptabiliser"] == 0
    assert body["par_statut"]["Rapprochée"] == 1


# --- Groupes N -> 1 (12/10/2026) ----------------------------------------------------------------


def _groupe(client, comptable, db):
    compte = account(db)
    cheque = operation(db, compte, montant="-10000", libelle="CHEQUE N 2100843 REMIS PAR ATW")
    com = operation(db, compte, montant="-50", libelle="COMMISSION")
    tva = operation(db, compte, montant="-5", libelle="TAXE SUR VALEUR AJOUTEE")
    entry = ecriture(db, compte, montant="10055", libelle="CHQ 2100843 FOURNISSEUR", tiers=None)
    body = run(client, comptable, db).json()
    return body, (cheque, com, tva), entry


def test_a_group_is_proposed_shown_validated_and_cancelled(client, comptable, db):
    body, lignes, entry = _groupe(client, comptable, db)

    [proposal] = proposals(client, comptable, db).json()["correspondances"]

    assert (body["nb_groupes"], body["nb_propositions"]) == (1, 0)
    assert proposal["type"] == "N-1"
    assert [o["id"] for o in proposal["operations"]] == [lignes[0].id, lignes[1].id, lignes[2].id]
    assert (proposal["total_operations"], proposal["ecart"]) == ("-10055.00", "0.00")
    assert "= 10 055,00" in proposal["explication"]
    assert statuses(db, *lignes, entry) == ["À vérifier"] * 4

    url = f"/api/reconciliation/matches/{proposal['id']}"
    valide = client.post(f"{url}/validate", headers=comptable)
    assert valide.status_code == 200, valide.text
    assert statuses(db, *lignes, entry) == ["Rapprochée"] * 4

    annule = client.request("DELETE", url, json={"motif": "Groupe incorrect"}, headers=comptable)
    assert annule.status_code == 200, annule.text
    assert statuses(db, *lignes, entry) == ["Non rapprochée"] * 4
    actions = audit_actions(db)
    assert "validation_rapprochement" in actions
    assert "annulation_rapprochement" in actions


def test_a_rejected_group_is_not_proposed_again(client, comptable, db):
    _groupe(client, comptable, db)
    [proposal] = proposals(client, comptable, db).json()["correspondances"]

    client.post(f"/api/reconciliation/matches/{proposal['id']}/reject", headers=comptable)
    again = run(client, comptable, db).json()

    assert again["nb_groupes"] == 0


def test_one_to_one_comes_before_groups(client, comptable, db):
    compte = account(db)
    operation(db, compte, montant="-500", libelle="CHEQUE N 3300001")
    operation(db, compte, montant="-5", libelle="COMMISSION")
    ecriture(db, compte, montant="500", libelle="CHQ 3300001", tiers=None)

    body = run(client, comptable, db).json()

    assert (body["nb_propositions"], body["nb_groupes"]) == (1, 0)


# --- Code de rapprochement (13/10/2026) ---------------------------------------------------------


def _valider(client, comptable, match_id):
    response = client.post(f"/api/reconciliation/matches/{match_id}/validate", headers=comptable)
    assert response.status_code == 200, response.text
    return response.json()


def test_a_validated_match_gets_a_code_per_account_and_year(client, comptable, db):
    compte = account(db)
    operation(db, compte, montant="1000", libelle="VIR A")
    ecriture(db, compte, montant="-1000", libelle="VIR A", tiers=None)
    operation(db, compte, montant="2000", libelle="VIR B", jour=date(2026, 9, 25))
    ecriture(db, compte, montant="-2000", libelle="VIR B", jour=date(2026, 9, 25), tiers=None)
    run(client, comptable, db)
    pending = proposals(client, comptable, db).json()["correspondances"]
    assert [p["code"] for p in pending] == [None, None]

    codes = [_valider(client, comptable, p["id"])["code"] for p in pending]

    annee = business_today().year % 100
    assert codes == [f"AWB-{annee:02d}-000001", f"AWB-{annee:02d}-000002"]


def test_the_code_is_on_every_line_and_searchable(client, comptable, db):
    body, lignes, entry = _groupe(client, comptable, db)
    [proposal] = proposals(client, comptable, db).json()["correspondances"]
    code = _valider(client, comptable, proposal["id"])["code"]
    params = {"company_id": company(db).id}

    operations = client.get(
        "/api/reconciliation/transactions", params={**params, "q": code}, headers=comptable
    ).json()["operations"]
    ecritures = client.get(
        "/api/accounting/entries", params={**params, "q": code.lower()}, headers=comptable
    ).json()["ecritures"]

    assert sorted(o["id"] for o in operations) == sorted(t.id for t in lignes)
    assert {o["correspondance"]["code"] for o in operations} == {code}
    assert [(e["id"], e["code_rapprochement"]) for e in ecritures] == [(entry.id, code)]


def test_a_cancelled_match_keeps_its_code_and_a_new_validation_gets_a_new_one(
    client, comptable, db
):
    compte = account(db)
    tx = operation(db, compte, montant="1000", libelle="VIR A")
    entry = ecriture(db, compte, montant="-1000", libelle="VIR A", tiers=None)
    run(client, comptable, db)
    [proposal] = proposals(client, comptable, db).json()["correspondances"]
    premier = _valider(client, comptable, proposal["id"])["code"]

    url = f"/api/reconciliation/matches/{proposal['id']}"
    annule = client.request("DELETE", url, json={"motif": "Erreur"}, headers=comptable).json()
    nouveau = manual(client, comptable, tx, entry).json()["code"]

    assert annule["code"] == premier
    assert nouveau != premier and nouveau.endswith("000002")
    [log] = [a for a in db.scalars(select(AuditLog).filter_by(action="validation_rapprochement"))]
    assert log.nouvelle_valeur["code"] == premier


# --- Export du lettrage (14/10/2026) ---------------------------------------------------------------


def lettrage(client, headers, db, **params):
    params = {"company_id": str(company(db).id), **PERIODE_EXPORT, **params}
    return client.get("/api/reconciliation/export", params=params, headers=headers)


PERIODE_EXPORT = {"from": "2026-09-01", "to": "2026-09-30"}


def test_the_lettrage_exports_validated_matches_grouped_and_balanced(client, comptable, db):
    from io import BytesIO

    from openpyxl import load_workbook

    _, (cheque, com, tva), entry = _groupe(client, comptable, db)
    compte = account_of(db, cheque)
    operation(db, compte, montant="1000", libelle="VIR A", jour=date(2026, 9, 10))
    ecriture(db, compte, montant="-1000", libelle="VIR A", jour=date(2026, 9, 10), tiers=None)
    operation(db, compte, montant="700", libelle="VIR REJ", jour=date(2026, 9, 12))
    ecriture(db, compte, montant="-700", libelle="VIR REJ", jour=date(2026, 9, 12), tiers=None)
    run(client, comptable, db)
    pending = {
        p["operation"]["libelle"]: p["id"]
        for p in proposals(client, comptable, db).json()["correspondances"]
    }
    groupe = _valider(client, comptable, pending["CHEQUE N 2100843 REMIS PAR ATW"])
    simple = _valider(client, comptable, pending["VIR A"])
    client.post(f"/api/reconciliation/matches/{pending['VIR REJ']}/reject", headers=comptable)

    response = lettrage(client, comptable, db)

    assert response.status_code == 200, response.text
    rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
    assert rows[0][:9] == (
        "Code", "Côté", "Banque", "Date", "Journal", "N° pièce", "Libellé", "Débit", "Crédit",
    )  # fmt: skip
    lignes = rows[1:-1]
    # Le 1 → 1 (10/09) d'abord, puis le groupe (24/09) : ses 3 opérations puis l'écriture Sage
    assert [(r[0], r[1], r[6], r[9]) for r in lignes] == [
        (simple["code"], "Banque", "VIR A", "1-1"),
        (simple["code"], "Sage", "VIR A", "1-1"),
        (groupe["code"], "Banque", "CHEQUE N 2100843 REMIS PAR ATW", "N-1"),
        (groupe["code"], "Banque", "COMMISSION", "N-1"),
        (groupe["code"], "Banque", "TAXE SUR VALEUR AJOUTEE", "N-1"),
        (groupe["code"], "Sage", "CHQ 2100843 FOURNISSEUR", "N-1"),
    ]
    for code in (simple["code"], groupe["code"]):
        debit = sum(r[7] or 0 for r in lignes if r[0] == code)
        credit = sum(r[8] or 0 for r in lignes if r[0] == code)
        assert debit == credit
    assert rows[-1][0] == "Total : 2 rapprochements"
    assert rows[-1][7] == rows[-1][8] == 11055
    assert "export_lettrage" in audit_actions(db)

    # Hors période ou autre compte : rien
    vide = lettrage(client, comptable, db, **{"from": "2026-10-01", "to": "2026-10-31"})
    autre = lettrage(client, comptable, db, bank_account_id=str(account(db, bank_code="CIH").id))
    for reponse in (vide, autre):
        rows = list(load_workbook(BytesIO(reponse.content)).active.iter_rows(values_only=True))
        assert rows[-1][0] == "Total : 0 rapprochement"


def test_the_lettrage_needs_the_view_permission_and_a_valid_period(client, comptable, db):
    assert lettrage(client, {}, db).status_code == 401
    assert (
        lettrage(client, comptable, db, **{"from": "2026-09-30", "to": "2026-09-01"}).status_code
        == 409
    )


# --- Filtre « Type d'opération » (14/10/2026) -----------------------------------------------------

LIBELLES_BANQUE = {
    "ENCAISSEMENT CHEQUE N 0173813": "Chèque",
    "ENCAISSEMENT EFFET N 2600114": "Effet",
    "REMISE D'EFFETS N 12": "Effet",
    "VERSEMENT DEPLACE PAR UN TIERS": "Espèces",
    "VIREMENT RECU DE X": "Virement",
    "VIR.RECU DE Y": "Virement",
    "RTGS EMIS VERS BMCE": "Virement",
    "IMPAYE SUR EFFET ESCOMPTE N 2600009 TIRE SUR X": "Impayé",
    "IMP LCN 1099850/18942568": "Impayé",
    "IMPOTS ET TAXES": "Autre",  # « IMP » n'est pas un mot entier
    "COMMISSION": "Autre",
    "VIRGULE SARL PAIEMENT": "Autre",  # « VIR » n'est pas un mot entier
    "AGIOS BANCAIRES": "Autre",
}
LIBELLES_SAGE = {
    "EAR1° AXK 173813": "Chèque",
    "EAR° 2600114 CLIENT": "Effet",
    "ESP N 45/OV": "Espèces",
    "Virement fournisseur": "Virement",
    "VIRT 12": "Virement",
    "IMP EAR1° AXL 917714 AJETTAR": "Impayé",
    "IMP EAR° DA 9515404 EL MOUAFFAK": "Impayé",
    "COM": "Autre",
    "EARTH SARL": "Autre",
}


def test_the_type_filter_reads_labels_exactly_like_the_engine(client, comptable, db):
    from app.services.reconciliation_scoring import type_ecriture, type_operation

    compte = account(db)
    for libelle in LIBELLES_BANQUE:
        operation(db, compte, libelle=libelle)
    for libelle in LIBELLES_SAGE:
        ecriture(db, compte, libelle=libelle)
    base = {"company_id": str(company(db).id)}

    for nom in ("Chèque", "Effet", "Espèces", "Virement", "Impayé", "Autre"):
        ops = client.get(
            "/api/reconciliation/transactions",
            params={**base, "type_operation": nom},
            headers=comptable,
        ).json()["operations"]
        ecritures = client.get(
            "/api/accounting/entries", params={**base, "type_operation": nom}, headers=comptable
        ).json()["ecritures"]
        attendu_banque = {lib for lib, t in LIBELLES_BANQUE.items() if t == nom}
        attendu_sage = {lib for lib, t in LIBELLES_SAGE.items() if t == nom}
        assert {op["libelle"] for op in ops} == attendu_banque, nom
        assert {e["libelle"] for e in ecritures} == attendu_sage, nom
        # Mêmes réponses que le moteur
        assert {(type_operation(lib) or "Autre") for lib in attendu_banque} <= {nom}
        assert {(type_ecriture(lib) or "Autre") for lib in attendu_sage} <= {nom}

    inconnu = client.get(
        "/api/reconciliation/transactions",
        params={**base, "type_operation": "Carte"},
        headers=comptable,
    )
    assert inconnu.status_code == 422


def test_proposals_and_ambiguous_operations_are_filtered_by_type(client, comptable, db):
    compte = account(db)
    operation(db, compte, montant="1000", libelle="ENCAISSEMENT CHEQUE N 0173813")
    ecriture(db, compte, montant="-1000", libelle="EAR1° AXK 173813", tiers=None)
    operation(db, compte, montant="2000", libelle="VIREMENT RECU CLIENT A")
    ecriture(db, compte, montant="-2000", libelle="VIREMENT CLIENT A", tiers=None)
    run(client, comptable, db)

    cheques = proposals(client, comptable, db, type_operation="Chèque").json()["correspondances"]
    virements = proposals(client, comptable, db, type_operation="Virement").json()[
        "correspondances"
    ]
    autres = proposals(client, comptable, db, type_operation="Autre").json()["correspondances"]
    ambigues = client.get(
        "/api/reconciliation/ambiguous",
        params={"company_id": str(company(db).id), "type_operation": "Effet"},
        headers=comptable,
    )

    assert [p["operation"]["libelle"] for p in cheques] == ["ENCAISSEMENT CHEQUE N 0173813"]
    assert [p["operation"]["libelle"] for p in virements] == ["VIREMENT RECU CLIENT A"]
    assert autres == []
    assert ambigues.status_code == 200
    assert ambigues.json()["total"] == 0


# --- Comptes en devise : le relevé se rapproche du montant devise de Sage (14/10/2026) -------------


def test_a_currency_account_is_matched_on_the_sage_currency_amount(client, comptable, db):
    compte = account(db, bank_code="CIH")
    compte.devise = "EUR"
    db.flush()
    # Relevé en EUR : 1 000 EUR reçus ; Sage : débit de 10 850 MAD, soit 1 000 EUR
    tx = operation(db, compte, montant="1000", libelle="VIR RECU CLIENT EUR")
    entry = ecriture(
        db,
        compte,
        montant="-10850",
        libelle="VIR CLIENT EUR",
        tiers=None,
        montant_devise=Decimal("-1000"),
    )
    mad = operation(db, compte, montant="10850", libelle="VIR RECU AUTRE", jour=date(2026, 9, 25))

    run(client, comptable, db)
    [proposal] = proposals(client, comptable, db).json()["correspondances"]

    assert (proposal["operation"]["id"], proposal["ecriture"]["id"]) == (tx.id, entry.id)
    assert proposal["ecriture"]["montant_devise"] == "-1000.00"
    _valider(client, comptable, proposal["id"])
    assert statuses(db, tx, entry, mad) == ["Rapprochée", "Rapprochée", "Non rapprochée"]


def test_a_second_run_does_not_wait_behind_a_running_one(client, comptable, db, monkeypatch):
    """15/10/2026 : un lancement alors qu'un autre tient la société répond 409 aussitôt."""
    from sqlalchemy.exc import OperationalError

    from app.repositories import reconciliation_repository

    class Verrouillee(Exception):
        sqlstate = "55P03"

    def deja_verrouillee(*args, **kwargs):
        raise OperationalError("SELECT … FOR UPDATE NOWAIT", {}, Verrouillee())

    monkeypatch.setattr(reconciliation_repository, "lock_company", deja_verrouillee)

    response = client.post(
        "/api/reconciliation/run",
        json={"company_id": company(db).id, **PERIODE},
        headers=comptable,
    )

    assert response.status_code == 409
    assert "déjà en cours" in response.json()["detail"]


def test_no_entry_matches_an_ambiguous_operation(client, comptable, direction, db):
    """15/10/2026 : « Aucune écriture ne correspond » écarte toutes les candidates d'un coup."""
    compte = account(db)
    tx = operation(db, compte)
    first, second = ecriture(db, compte), ecriture(db, compte)
    run(client, comptable, db)
    url = f"/api/reconciliation/transactions/{tx.id}/aucune-ecriture"

    refus = client.post(url, json={}, headers=direction)
    response = client.post(url, json={"commentaire": "Chèque d'un autre client"}, headers=comptable)

    assert refus.status_code == 403
    assert response.status_code == 200, response.text
    assert sorted(response.json()["ecritures_ecartees"]) == sorted([first.id, second.id])
    assert statuses(db, tx) == ["Non rapprochée"]
    ambigues = client.get(
        "/api/reconciliation/ambiguous", params={"company_id": company(db).id}, headers=direction
    ).json()
    assert ambigues["ambigues"] == []
    audit = db.scalar(select(AuditLog).where(AuditLog.action == "rejet_ambiguite"))
    assert audit.nouvelle_valeur["commentaire"] == "Chèque d'un autre client"

    # Le moteur ne repropose jamais ces écritures ; une nouvelle écriture reste comparée
    run(client, comptable, db)
    assert statuses(db, tx) == ["Non rapprochée"]
    nouvelle = ecriture(db, compte, libelle="VIR ABC", numero_piece="NEW1")
    run(client, comptable, db)
    proposees = client.get(
        "/api/reconciliation/proposals",
        params={"company_id": company(db).id},
        headers=direction,
    ).json()["correspondances"]
    assert [(p["operation"]["id"], p["ecriture"]["id"]) for p in proposees] == [
        (tx.id, nouvelle.id)
    ]

    # L'opération a maintenant une proposition (la nouvelle écriture) : à traiter d'abord
    assert client.post(url, json={}, headers=comptable).status_code == 409


def test_ambiguous_entries_come_with_their_operations(client, comptable, direction, db):
    """15/10/2026 : une écriture ambiguë (plusieurs opérations aussi proches) est listée avec ses
    opérations candidates ; le rapprochement N → 1 qui suit passe par le rapprochement manuel de
    plusieurs lignes (tests de `test_reconciliation_manuel_groupe.py`)."""
    compte = account(db)
    entry = ecriture(db, compte)
    first, second = operation(db, compte), operation(db, compte)
    run(client, comptable, db)
    url = "/api/reconciliation/ambiguous-entries"

    refus = client.get(url, params={"company_id": company(db).id}, headers={})
    response = client.get(url, params={"company_id": company(db).id}, headers=direction)

    assert refus.status_code == 401
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["total"] == 1
    [ambigue] = body["ecritures"]
    assert ambigue["ecriture"]["id"] == entry.id
    assert {c["operation"]["id"] for c in ambigue["candidats"]} == {first.id, second.id}


def test_an_entry_is_ambiguous_only_with_two_operations_at_the_same_score(
    client, comptable, direction, db
):
    """16/10/2026 : une opération qui hésite entre deux écritures rend ces écritures « À
    vérifier », mais chacune n'a qu'une opération : elles ne sont pas des écritures ambiguës
    (l'opération se traite dans « Opérations ambiguës »)."""
    compte = account(db)
    tx = operation(db, compte)
    first, second = ecriture(db, compte), ecriture(db, compte)
    run(client, comptable, db)
    params = {"company_id": company(db).id}

    ecritures = client.get(
        "/api/reconciliation/ambiguous-entries", params=params, headers=direction
    ).json()
    operations = client.get(
        "/api/reconciliation/ambiguous", params=params, headers=direction
    ).json()

    assert statuses(db, first) == statuses(db, second) == ["À vérifier"]
    assert ecritures == {**ecritures, "total": 0, "ecritures": []}
    [ambigue] = operations["ambigues"]
    assert ambigue["operation"]["id"] == tx.id
    assert {c["ecriture"]["id"] for c in ambigue["candidats"]} == {first.id, second.id}


def test_no_operation_matches_an_ambiguous_entry(client, comptable, direction, db):
    compte = account(db)
    entry = ecriture(db, compte)
    first, second = operation(db, compte), operation(db, compte)
    run(client, comptable, db)
    url = f"/api/reconciliation/entries/{entry.id}/aucune-operation"

    assert client.post(url, json={}, headers=direction).status_code == 403
    response = client.post(url, json={"commentaire": "Pas encore en banque"}, headers=comptable)

    assert response.status_code == 200, response.text
    assert sorted(response.json()["operations_ecartees"]) == sorted([first.id, second.id])
    assert statuses(db, entry) == ["Non rapprochée"]
    restantes = client.get(
        "/api/reconciliation/ambiguous-entries",
        params={"company_id": company(db).id},
        headers=direction,
    ).json()
    assert restantes["ecritures"] == []
    audit = db.scalar(select(AuditLog).where(AuditLog.action == "rejet_ambiguite_ecriture"))
    assert audit.nouvelle_valeur["commentaire"] == "Pas encore en banque"
    encore = client.post(url, json={}, headers=comptable)
    assert encore.status_code == 200
    assert encore.json()["operations_ecartees"] == []


def test_proposals_are_rejected_in_a_batch_all_or_nothing(client, comptable, direction, db):
    """15/10/2026 : « Rejeter la sélection » des propositions à vérifier."""
    compte = account(db)
    tx1 = operation(db, compte, montant="100", jour=date(2026, 9, 1), libelle="VIR UN")
    ec1 = ecriture(db, compte, montant="-100", jour=date(2026, 9, 1), libelle="Client un")
    tx2 = operation(db, compte, montant="200", jour=date(2026, 9, 5), libelle="VIR DEUX")
    ec2 = ecriture(db, compte, montant="-200", jour=date(2026, 9, 5), libelle="Client deux")
    run(client, comptable, db)
    par_operation = {
        p["operation"]["id"]: p["id"]
        for p in client.get(
            "/api/reconciliation/proposals",
            params={"company_id": company(db).id},
            headers=direction,
        ).json()["correspondances"]
    }
    ids = [par_operation[tx1.id], par_operation[tx2.id]]
    url = "/api/reconciliation/matches/reject-batch"

    assert client.post(url, json={"ids": ids}, headers=direction).status_code == 403
    client.post(f"/api/reconciliation/matches/{ids[0]}/validate", headers=comptable)
    tout_ou_rien = client.post(url, json={"ids": ids}, headers=comptable)
    assert tout_ou_rien.status_code == 409
    assert statuses(db, tx2, ec2) == ["À vérifier", "À vérifier"]

    response = client.post(url, json={"ids": ids[1:], "commentaire": "Faux"}, headers=comptable)

    assert response.status_code == 200, response.text
    assert response.json() == {"nb_traitees": 1}
    assert statuses(db, tx2, ec2) == ["Non rapprochée", "Non rapprochée"]
    assert statuses(db, tx1, ec1) == ["Rapprochée", "Rapprochée"]
    rejet = db.scalar(
        select(AuditLog)
        .where(AuditLog.action == "rejet_rapprochement")
        .order_by(AuditLog.id.desc())
    )
    assert rejet.nouvelle_valeur["en_lot"] is True
    run(client, comptable, db)
    assert statuses(db, tx2, ec2) == ["Non rapprochée", "Non rapprochée"]


def test_ambiguous_operations_and_entries_are_discarded_in_a_batch(client, comptable, db):
    compte = account(db)
    # Deux opérations ambiguës (deux écritures chacune), un autre jour
    ambigues = []
    for jour in (date(2026, 9, 2), date(2026, 9, 20)):
        ambigues.append(operation(db, compte, jour=jour))
        ecriture(db, compte, jour=jour)
        ecriture(db, compte, jour=jour)
    # Une écriture ambiguë (deux opérations)
    entry = ecriture(db, compte, montant="-700", jour=date(2026, 9, 12), libelle="Fournisseur")
    operation(db, compte, montant="700", jour=date(2026, 9, 12), libelle="VIR A")
    operation(db, compte, montant="700", jour=date(2026, 9, 12), libelle="VIR A")
    run(client, comptable, db)

    ops = client.post(
        "/api/reconciliation/transactions/aucune-ecriture",
        json={"ids": [tx.id for tx in ambigues]},
        headers=comptable,
    )
    ecrs = client.post(
        "/api/reconciliation/entries/aucune-operation", json={"ids": [entry.id]}, headers=comptable
    )
    encore = client.post(
        "/api/reconciliation/transactions/aucune-ecriture",
        json={"ids": [ambigues[0].id]},
        headers=comptable,
    )

    assert ops.status_code == 200, ops.text
    assert ops.json() == {"nb_traitees": 2}
    assert statuses(db, *ambigues) == ["Non rapprochée", "Non rapprochée"]
    assert ecrs.status_code == 200, ecrs.text
    assert statuses(db, entry) == ["Non rapprochée"]
    # Plus rien à écarter pour la première : traitée sans erreur (15/10/2026)
    assert encore.status_code == 200


def test_discarding_operations_frees_the_entries_left_without_candidates(
    client, comptable, direction, db
):
    """15/10/2026 : deux opérations et deux écritures identiques sont toutes ambiguës. Écarter les
    candidates des deux opérations laisse les écritures sans candidate : elles sortent aussi de
    « À vérifier » au lieu de rester dans « Écritures ambiguës » sans rien à choisir."""
    compte = account(db)
    ops = [operation(db, compte), operation(db, compte)]
    ecrs = [ecriture(db, compte), ecriture(db, compte)]
    run(client, comptable, db)
    assert statuses(db, *ops, *ecrs) == ["À vérifier"] * 4

    response = client.post(
        "/api/reconciliation/transactions/aucune-ecriture",
        json={"ids": [tx.id for tx in ops]},
        headers=comptable,
    )

    assert response.status_code == 200, response.text
    assert statuses(db, *ops, *ecrs) == ["Non rapprochée"] * 4
    restantes = client.get(
        "/api/reconciliation/ambiguous-entries",
        params={"company_id": company(db).id},
        headers=direction,
    ).json()
    assert restantes["total"] == 0
    audit = db.scalar(
        select(AuditLog).where(AuditLog.action == "rejet_ambiguite").order_by(AuditLog.id.desc())
    )
    assert sorted(audit.nouvelle_valeur["ecritures_liberees"]) == sorted(e.id for e in ecrs)


def test_a_fee_is_never_proposed_alone_and_stays_to_be_booked(client, comptable, direction, db):
    """15/10/2026 : les frais et leur TVA se comptabilisent dans SIMTIS. Une commission face à
    une écriture « COM » du même montant n'est pas proposée ; elle reste « À comptabiliser »."""
    compte = account(db)
    commission = operation(db, compte, montant="-50", libelle="COMMISSION")
    ecriture(db, compte, montant="50", libelle="COM", numero_piece="C1", tiers=None)

    run(client, comptable, db)

    assert statuses(db, commission) == ["Non rapprochée"]
    proposees = client.get(
        "/api/reconciliation/proposals", params={"company_id": company(db).id}, headers=direction
    ).json()["correspondances"]
    assert proposees == []
    page = client.get(
        "/api/reconciliation/transactions",
        params={"company_id": company(db).id, **PERIODE},
        headers=direction,
    ).json()
    assert page["a_comptabiliser"] == 1


# --- Rapprochement depuis les ambiguës : relance ciblée (16/10/2026) ------------------------------


def manuel_cible(client, headers, tx, entry):
    return client.post(
        "/api/reconciliation/matches",
        json={"transaction_id": tx.id, "ecriture_id": entry.id, "relance_ciblee": True},
        headers=headers,
    )


def test_a_targeted_match_settles_only_the_competing_lines(client, comptable, db):
    """Deux opérations et deux écritures identiques : toutes à égalité. Rapprocher la première
    paire depuis les ambiguës propose la seconde ; une proposition ailleurs ne bouge pas."""
    compte = account(db)
    tx1, tx2 = operation(db, compte), operation(db, compte)
    ec1, ec2 = ecriture(db, compte), ecriture(db, compte)
    autre_tx = operation(db, compte, "700", date(2026, 9, 20), "VIR XYZ")
    autre_ec = ecriture(db, compte, "-700", date(2026, 9, 20), "Règlement XYZ", numero_piece="X1")
    run(client, comptable, db)
    assert statuses(db, tx1, tx2, ec1, ec2) == ["À vérifier"] * 4
    ailleurs = db.scalar(
        select(ReconciliationMatchItem.match_id).filter_by(bank_transaction_id=autre_tx.id)
    )

    response = manuel_cible(client, comptable, tx1, ec1)

    assert response.status_code == 201, response.text
    proposees = {
        (p["operation"]["id"], p["ecriture"]["id"], p["id"])
        for p in proposals(client, comptable, db).json()["correspondances"]
    }
    assert {(t, e) for t, e, _ in proposees} == {(tx2.id, ec2.id), (autre_tx.id, autre_ec.id)}
    assert ailleurs in {i for _, _, i in proposees}
    assert statuses(db, tx2, ec2) == ["À vérifier", "À vérifier"]
    audit = db.scalar(
        select(AuditLog)
        .where(AuditLog.action == "rapprochement_manuel")
        .order_by(AuditLog.id.desc())
    )
    assert audit.nouvelle_valeur["relance"]["ciblee"] is True


def test_a_targeted_match_frees_the_other_tied_entry(client, comptable, db):
    """Une opération à égalité entre deux écritures : l'écriture non choisie repasse « Non
    rapprochée » (plus aucune candidate)."""
    compte = account(db)
    tx = operation(db, compte)
    ec1, ec2 = ecriture(db, compte), ecriture(db, compte)
    run(client, comptable, db)
    assert statuses(db, tx, ec1, ec2) == ["À vérifier"] * 3

    assert manuel_cible(client, comptable, tx, ec1).status_code == 201

    assert statuses(db, tx, ec1, ec2) == ["Rapprochée", "Rapprochée", "Non rapprochée"]


def test_both_panes_filter_by_debit_or_credit(client, comptable, db):
    """19/10/2026 : filtre « Sens » des volets (banque : argent sorti / entré ; Sage : colonne)."""
    compte = account(db)
    sortie = operation(db, compte, montant="-300", libelle="CHEQUE 12")
    entree = operation(db, compte, montant="400", libelle="VIR CLIENT")
    debit = ecriture(db, compte, montant="-400", libelle="Client")
    credit = ecriture(db, compte, montant="300", libelle="Fournisseur", numero_piece="F1")
    base = {"company_id": str(company(db).id)}

    def ids(url, cle, sens):
        reponse = client.get(url, params={**base, "sens": sens}, headers=comptable)
        assert reponse.status_code == 200, reponse.text
        return [item["id"] for item in reponse.json()[cle]]

    assert ids("/api/reconciliation/transactions", "operations", "Débit") == [sortie.id]
    assert ids("/api/reconciliation/transactions", "operations", "Crédit") == [entree.id]
    assert ids("/api/accounting/entries", "ecritures", "Débit") == [debit.id]
    assert ids("/api/accounting/entries", "ecritures", "Crédit") == [credit.id]
    refus = client.get(
        "/api/reconciliation/transactions", params={**base, "sens": "Autre"}, headers=comptable
    )
    assert refus.status_code == 422


def test_a_pending_operation_is_never_compared_nor_matched_by_hand(client, comptable, db):
    """20/10/2026 : une opération en cours (« (*) ») n'est jamais proposée ni rapprochée ; une
    opération remplacée n'est plus listée."""
    compte = account(db)
    en_cours = operation(db, compte, libelle="VIR ABC (*)", en_cours=True)
    entry = ecriture(db, compte)
    remplacee = operation(db, compte, montant="700", libelle="VIR Z (*)", en_cours=True)
    definitive = operation(db, compte, montant="700", libelle="VIR Z")
    remplacee.remplacee_par_id = definitive.id
    db.flush()

    run(client, comptable, db)
    manuel = client.post(
        "/api/reconciliation/matches",
        json={"transaction_id": en_cours.id, "ecriture_id": entry.id},
        headers=comptable,
    )
    page = client.get(
        "/api/reconciliation/transactions",
        params={"company_id": str(company(db).id)},
        headers=comptable,
    ).json()

    assert statuses(db, en_cours, entry) == ["Non rapprochée", "Non rapprochée"]
    assert manuel.status_code == 409
    assert {op["id"] for op in page["operations"]} == {en_cours.id, definitive.id}
    assert page["en_cours"] == 1
    assert page["par_statut"]["Non rapprochée"] == 1


def test_the_wide_search_proposes_far_pairs_without_touching_the_engine(client, comptable, db):
    """20/10/2026 : bouton « Recherche élargie » ; ses propositions sont marquées, jamais
    cochées, et le moteur normal ne les remplace pas."""
    compte = account(db)
    cheque = operation(db, compte, montant="20000", libelle="ENCAISSEMENT CHEQUE N 0174512")
    loin = ecriture(
        db, compte, montant="-20000", jour=date(2026, 3, 1), libelle="EAR1° 174512 X",
        numero_piece="P1",
    )  # fmt: skip
    proche_tx = operation(db, compte, montant="700", libelle="VIR ABC")
    proche_ec = ecriture(db, compte, montant="-700", libelle="Virement ABC", numero_piece="P2")
    body = {"company_id": company(db).id, **PERIODE}

    refus = client.post("/api/reconciliation/run-elargie", json=body, headers={})
    response = client.post("/api/reconciliation/run-elargie", json=body, headers=comptable)

    assert refus.status_code == 401
    assert response.status_code == 200, response.text
    assert response.json()["nb_propositions"] == 1
    assert statuses(db, cheque, loin) == ["À vérifier", "À vérifier"]
    run(client, comptable, db)
    proposees = {
        p["operation"]["id"]: p for p in proposals(client, comptable, db).json()["correspondances"]
    }
    assert proposees[cheque.id]["elargie"] is True
    assert proposees[cheque.id]["ecart_jours"] == (JOUR - date(2026, 3, 1)).days
    assert proposees[proche_tx.id]["elargie"] is False
    assert proche_ec.id == proposees[proche_tx.id]["ecriture"]["id"]
