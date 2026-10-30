"""Rapports Excel (P16.4) : position du jour, état de rapprochement, permissions et audit."""

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import AuditLog
from tests.helpers import bearer, login, make_auth_user
from tests.test_reconciliation_api import account, company, ecriture, operation


@pytest.fixture
def responsable(client, reference):
    make_auth_user(reference, "RESPONSABLE", email="resp@example.com")
    return bearer(login(client, "resp@example.com"))


def classeur(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return load_workbook(BytesIO(response.content))


def test_the_daily_position_keeps_the_banks_table_structure(client, responsable, db):
    awb = account(db)
    awb.credit_autorise = Decimal("500000.00")
    db.flush()
    operation(db, awb, montant="-200000", jour=date(2026, 9, 1), solde=Decimal("-200000"))

    response = client.get(
        "/api/rapports/position",
        params={"company_id": str(company(db).id), "date": "2026-09-02"},
        headers=responsable,
    )

    sheet = classeur(response)["Banques"]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    en_tete = rows[2]
    assert en_tete[1:3] == ["AWB", "BMCE"] and en_tete[-2:] == ["TOTAL", "DEPASSEMENT"]
    assert rows[3][0] == "LIGNE" and rows[3][1] == 500000
    assert rows[4][0] == "facilité de caisse"
    # 01/09 : -200 000 + 500 000 ; 02/09 : dernier solde connu repris, en italique
    assert rows[5][1] == 300000 and rows[6][1] == 300000
    assert sheet.cell(row=7, column=2).font.italic is True
    assert rows[-1][0] == "Disponible Fc reel"
    assert "attachment" in response.headers["content-disposition"]
    audit = db.scalar(select(AuditLog).where(AuditLog.action == "export_rapport"))
    assert audit.nouvelle_valeur["rapport"] == "position_du_jour"


def test_the_reconciliation_state_counts_statuses_per_account(client, responsable, db):
    awb = account(db)
    operation(db, awb, montant="1000", jour=date(2026, 9, 2), statut="Rapprochée")
    operation(db, awb, montant="-250", jour=date(2026, 9, 3), libelle="CHQ 1")
    operation(db, awb, montant="400", jour=date(2026, 9, 4), statut="À vérifier")
    operation(db, awb, montant="-999", jour=date(2026, 10, 1), libelle="hors période")

    response = client.get(
        "/api/rapports/rapprochement",
        params={"company_id": str(company(db).id), "from": "2026-09-01", "to": "2026-09-30"},
        headers=responsable,
    )

    workbook = classeur(response)
    synthese = [[c.value for c in row] for row in workbook["Synthèse"].iter_rows(min_row=4)]
    assert len(synthese) == 1
    ligne = synthese[0]
    assert ligne[0].startswith("AWB · ")
    assert ligne[1:6] == [1, 1, 1, 0, 3]
    assert ligne[6] == pytest.approx(33.3)
    assert ligne[7:9] == [250, 400]
    detail = [[c.value for c in row] for row in workbook["À rapprocher"].iter_rows(min_row=2)]
    assert [(r[3], r[7]) for r in detail] == [
        ("CHQ 1", "Non rapprochée"),
        ("VIR ABC", "À vérifier"),
    ]


def test_reports_check_the_period_and_their_permissions(client, reference, db):
    make_auth_user(reference, "TRESORERIE", email="treso@example.com")
    treso = bearer(login(client, "treso@example.com"))
    societe = str(company(db).id)

    # Trésorerie consulte la position et le rapprochement (rôle de consultation)
    assert (
        client.get(
            "/api/rapports/position", params={"company_id": societe}, headers=treso
        ).status_code
        == 200
    )
    assert (
        client.get(
            "/api/rapports/rapprochement",
            params={"company_id": societe, "from": "2026-09-30", "to": "2026-09-01"},
            headers=treso,
        ).status_code
        == 409
    )

    make_auth_user(reference, email="sansrole@example.com")
    aucun = bearer(login(client, "sansrole@example.com"))
    assert (
        client.get(
            "/api/rapports/position", params={"company_id": societe}, headers=aucun
        ).status_code
        == 403
    )
    assert (
        client.get(
            "/api/rapports/rapprochement", params={"company_id": societe}, headers=aucun
        ).status_code
        == 403
    )


# --- Position bancaire en PDF (27/10/2026) --------------------------------------------------------


def test_the_daily_position_is_also_given_as_a_pdf(client, responsable, db):
    awb = account(db)
    awb.credit_autorise = Decimal("500000.00")
    db.flush()
    operation(db, awb, montant="-200000", jour=date(2026, 9, 1), solde=Decimal("-200000"))

    response = client.get(
        "/api/rapports/position/pdf",
        params={"company_id": str(company(db).id), "date": "2026-09-02"},
        headers=responsable,
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    assert 'filename="position-simtis-20260902.pdf"' in response.headers["content-disposition"]
    audit = db.scalar(select(AuditLog).where(AuditLog.action == "export_rapport"))
    assert audit.nouvelle_valeur["rapport"] == "position_du_jour_pdf"


def test_the_pdf_position_needs_the_position_permission(client, reference, db):
    make_auth_user(reference, email="sansrole@example.com")
    aucun = bearer(login(client, "sansrole@example.com"))

    response = client.get(
        "/api/rapports/position/pdf", params={"company_id": str(company(db).id)}, headers=aucun
    )

    assert response.status_code == 403


@pytest.mark.parametrize(
    ("valeur", "texte"),
    [
        (Decimal("-1234567.8"), "-1 234 567,80"),
        (Decimal("0"), "0,00"),
        (Decimal("950.5"), "950,50"),
        (None, ""),
    ],
)
def test_pdf_amounts_are_written_the_french_way(valeur, texte):
    from app.services.rapports_service import montant_pdf

    assert montant_pdf(valeur) == texte


# --- Lignes non rapprochées (30/10/2026) ----------------------------------------------------------


def lignes(feuille, colonne: int) -> list[str]:
    """Libellés d'un côté de la feuille (colonne 2 = relevé, 9 = Sage), sous l'en-tête."""
    valeurs = [row[colonne - 1] for row in feuille.iter_rows(min_row=4, values_only=True)]
    return [v for v in valeurs if v not in (None, "Total") and not isinstance(v, (int, float))]


def test_unreconciled_lines_side_by_side_for_one_bank_and_period(client, responsable, db):
    bp = account(db, bank_code="BP")
    awb = account(db, bank_code="AWB")
    operation(db, bp, montant="9000", jour=date(2026, 9, 30), libelle="VERSEMENT DEPLACE PAR X")
    operation(db, bp, montant="-25", jour=date(2026, 9, 29), libelle="COMMISSION")
    propose = operation(db, bp, montant="1200", jour=date(2026, 9, 28), libelle="VIR KARIM")
    propose.statut = "À vérifier"
    rapprochee = operation(db, bp, montant="500", jour=date(2026, 9, 27), libelle="VIR RAPPROCHE")
    rapprochee.statut = "Rapprochée"
    en_cours = operation(db, bp, montant="700", jour=date(2026, 9, 26), libelle="VIR SRBM (*)")
    en_cours.en_cours = True
    operation(db, bp, montant="800", jour=date(2026, 8, 15), libelle="HORS PERIODE")
    operation(db, awb, montant="900", jour=date(2026, 9, 30), libelle="AUTRE BANQUE")
    ecriture(db, bp, montant="-9000", jour=date(2026, 9, 30), libelle="ESP TT26272F48TV")
    db.flush()

    response = client.get(
        "/api/rapports/non-rapproches",
        params={
            "company_id": str(company(db).id),
            "from": "2026-09-01",
            "to": "2026-09-30",
            "bank_id": str(bp.bank_id),
        },
        headers=responsable,
    )

    workbook = classeur(response)
    assert len(workbook.sheetnames) == 1
    feuille = workbook[workbook.sheetnames[0]]
    assert feuille.cell(row=2, column=1).value == "Transactions bancaires (3)"
    assert feuille.cell(row=2, column=8).value == "Écritures comptables (1)"
    # Le plus récent d'abord, rapprochées, en cours et hors période exclues
    assert lignes(feuille, 2) == ["VERSEMENT DEPLACE PAR X", "COMMISSION", "VIR KARIM"]
    statuts = [row[5] for row in feuille.iter_rows(min_row=4, max_row=6, values_only=True)]
    assert statuts == ["Non rapprochée", "Non rapprochée · frais", "En proposition"]
    assert feuille.cell(row=4, column=9).value == "ESP TT26272F48TV"
    assert (
        "non-rapproches-simtis-bp-2026-09-01-2026-09-30.xlsx"
        in (response.headers["content-disposition"])
    )
    audit = db.scalar(select(AuditLog).where(AuditLog.action == "export_rapport"))
    assert audit.nouvelle_valeur["rapport"] == "non_rapproches"


def test_unreconciled_lines_need_the_reconciliation_permission(client, reference, db):
    make_auth_user(reference, email="sansrole@example.com")
    aucun = bearer(login(client, "sansrole@example.com"))

    response = client.get(
        "/api/rapports/non-rapproches", params={"company_id": str(company(db).id)}, headers=aucun
    )

    assert response.status_code == 403
