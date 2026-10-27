"""Lecture d'un classeur d'import, commune aux relevés bancaires et aux exports comptables.

Depuis le 27/10/2026 un relevé bancaire peut aussi être un PDF **généré par la banque** (texte, pas
une image scannée) : ses tableaux, page après page, deviennent les lignes d'une feuille unique
« PDF », puis tout se passe comme pour un Excel (en-tête, colonnes, aperçu, contrôles). Les en-têtes
répétés en haut de chaque page sont retirés. Les exports Sage restent en Excel.

Fichiers Excel `.xlsx` ou `.xls` (l'ancien format, accepté depuis le 08/10/2026), 20 Mo et
50 000 lignes au plus (limites relevées le 08/10/2026, 5 Mo et 5 000 lignes avant). Le format est reconnu au contenu du fichier, pas seulement à son nom ; les deux
donnent les mêmes valeurs (dates en `datetime`, nombres, textes, cellule vide = None). La ligne d'en-tête
est celle qui reconnaît le plus de champs parmi les 30 premières ; la correspondance colonnes /
champs se propose par synonymes d'en-têtes, ou se reprend d'un modèle mémorisé (par en-tête).
"""

import zipfile
from collections import Counter
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date
from io import BytesIO

import pdfplumber
import xlrd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from app.services.errors import ConflictError
from app.services.normalization_service import clean_text, is_blank, normalize_header

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50_000
HEADER_SCAN_ROWS = 30
# Limite écrite dans les messages, avec l'espace des milliers
MAX_ROWS_TEXTE = f"{MAX_ROWS:,}".replace(",", " ")
# Recherche des doublons par paquets : une requête ne porte jamais sur 50 000 empreintes
HASH_BATCH = 5000
SAMPLES = 3

Mapping = dict[str, int | None]


@dataclass(frozen=True)
class ImportField:
    code: str
    libelle: str
    obligatoire: bool
    # En-têtes reconnus, déjà normalisés (`normalize_header`)
    synonymes: tuple[str, ...]


@dataclass
class Column:
    index: int
    lettre: str
    entete: str
    exemples: list[str]


EXTENSIONS = (".xlsx", ".xls")
FORMATS_ACCEPTES = "Seuls les fichiers Excel .xlsx ou .xls sont acceptés."
FORMATS_ACCEPTES_PDF = "Seuls les fichiers Excel .xlsx ou .xls, ou PDF, sont acceptés."
# Signature d'un .xls (document OLE2 / BIFF) ; tout le reste est lu comme un .xlsx (archive zip)
_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_PDF = b"%PDF"
FEUILLE_PDF = "PDF"
MAX_PAGES_PDF = 500


def check_file(fichier_nom: str, content: bytes, *, pdf: bool = False) -> None:
    """`pdf` : le fichier peut aussi être un PDF (relevés bancaires seulement, 27/10/2026)."""
    extensions = (*EXTENSIONS, ".pdf") if pdf else EXTENSIONS
    if not fichier_nom.lower().endswith(extensions):
        raise ConflictError(FORMATS_ACCEPTES_PDF if pdf else FORMATS_ACCEPTES)
    if not content:
        raise ConflictError("Le fichier est vide.")
    if len(content) > MAX_FILE_BYTES:
        raise ConflictError("Fichier trop volumineux : 20 Mo au plus.")


def read_sheet(
    content: bytes, feuille: str | None, unite: str = "fichier", *, pdf: bool = False
) -> tuple[list[str], str, list[tuple]]:
    """Feuilles du classeur, feuille lue et ses lignes (sans les lignes vides de la fin).
    `unite` nomme ce qu'est le fichier dans le message de taille (« relevé », « fichier »).
    `pdf` : un PDF est lu comme une feuille unique « PDF » ; sinon il est refusé."""
    if content.startswith(_PDF):
        if not pdf:
            raise ConflictError(FORMATS_ACCEPTES)
        names, name, rows = _read_pdf(content, feuille, unite)
    elif content.startswith(_OLE2):
        names, name, rows = _read_xls(content, feuille, unite)
    else:
        names, name, rows = _read_xlsx(content, feuille, unite)
    while rows and all(is_blank(cell) for cell in rows[-1]):
        rows.pop()
    if not rows:
        raise ConflictError(f"La feuille « {name} » est vide.")
    return names, name, rows


_ILLISIBLE = "Fichier illisible : ce n'est pas un classeur Excel .xlsx ou .xls valide."


def _too_long(rows: list, unite: str) -> None:
    if len(rows) > MAX_ROWS + HEADER_SCAN_ROWS:
        raise ConflictError(f"Fichier trop long : {MAX_ROWS_TEXTE} lignes au plus par {unite}.")


def _sheet_name(names: list[str], feuille: str | None) -> str:
    name = feuille or names[0]
    if name not in names:
        raise ConflictError(f"Feuille introuvable dans le fichier : « {name} ».")
    return name


def _read_xlsx(content: bytes, feuille: str | None, unite: str) -> tuple[list[str], str, list]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        raise ConflictError(_ILLISIBLE) from error
    try:
        names = list(workbook.sheetnames)
        name = _sheet_name(names, feuille)
        rows: list[tuple] = []
        for row in workbook[name].iter_rows(values_only=True):
            rows.append(tuple(row))
            _too_long(rows, unite)
    finally:
        workbook.close()
    return names, name, rows


def _cellule_pdf(valeur: str | None) -> str | None:
    """Texte d'une cellule de tableau PDF : retours à la ligne et espaces multiples réduits."""
    if valeur is None:
        return None
    texte = " ".join(valeur.split())
    return texte or None


def _tableaux_page(page) -> list[list[tuple]]:
    """Lignes des tableaux tracés (avec des traits) d'une page."""
    return [
        tuple(_cellule_pdf(cellule) for cellule in brute)
        for tableau in page.extract_tables()
        if tableau and len(tableau[0]) > 1
        for brute in tableau
    ]


@dataclass(frozen=True)
class MotPdf:
    """Un mot d'une page PDF et sa position (points, origine en haut à gauche)."""

    x0: float
    x1: float
    top: float
    bottom: float
    texte: str


# Deux mots sont sur la même ligne si leurs hauts diffèrent de 3 points au plus ; ils sont dans la
# même cellule si l'écart qui les sépare est inférieur à la moitié de leur hauteur (une espace
# fait environ le quart de la hauteur, l'écart entre deux colonnes bien plus).
TOLERANCE_LIGNE = 3.0
ECART_CELLULE = 0.5
# Une ligne de moins de 3 cellules (titre, « Compte : … », pied de page) ne sert pas à placer
# les colonnes.
CELLULES_MIN_COLONNES = 3

_Cellule = tuple[float, float, str]


def _cellules_des_lignes(mots: Sequence[MotPdf]) -> list[list[_Cellule]]:
    """Mots d'une page → lignes de haut en bas, chacune en cellules (x0, x1, texte)."""
    lignes: list[list[MotPdf]] = []
    for mot in sorted(mots, key=lambda m: (m.top, m.x0)):
        if lignes and abs(mot.top - lignes[-1][0].top) <= TOLERANCE_LIGNE:
            lignes[-1].append(mot)
        else:
            lignes.append([mot])
    resultat = []
    for ligne in lignes:
        cellules: list[list] = []
        for mot in sorted(ligne, key=lambda m: m.x0):
            ecart_max = max(2.0, ECART_CELLULE * (mot.bottom - mot.top))
            if cellules and mot.x0 - cellules[-1][1] <= ecart_max:
                cellules[-1][1] = max(cellules[-1][1], mot.x1)
                cellules[-1][2] += " " + mot.texte
            else:
                cellules.append([mot.x0, mot.x1, mot.texte])
        resultat.append([(c[0], c[1], c[2]) for c in cellules])
    return resultat


def _colonnes(lignes: Sequence[Sequence[_Cellule]]) -> list[tuple[float, float]]:
    """Colonnes communes à tout le document : les cellules qui se chevauchent horizontalement
    appartiennent à la même colonne."""
    intervalles = sorted(
        (x0, x1) for ligne in lignes if len(ligne) >= CELLULES_MIN_COLONNES for x0, x1, _ in ligne
    )
    colonnes: list[list[float]] = []
    for x0, x1 in intervalles:
        if colonnes and x0 <= colonnes[-1][1]:
            colonnes[-1][1] = max(colonnes[-1][1], x1)
        else:
            colonnes.append([x0, x1])
    return [(c[0], c[1]) for c in colonnes]


def _colonne_de(cellule: _Cellule, colonnes: Sequence[tuple[float, float]]) -> int:
    """Première colonne que la cellule recouvre (un titre qui déborde reste dans la colonne où il
    commence), sinon la colonne la plus proche de son milieu."""
    x0, x1, _ = cellule
    for index, (c0, c1) in enumerate(colonnes):
        if min(x1, c1) > max(x0, c0):
            return index
    milieu = (x0 + x1) / 2
    return min(range(len(colonnes)), key=lambda i: abs(sum(colonnes[i]) / 2 - milieu))


def lignes_par_mots(pages: Sequence[Sequence[MotPdf]]) -> list[tuple]:
    """Lignes d'un relevé PDF sans traits de tableau, à partir de la position des mots (fonction
    pure) : mêmes colonnes pour toutes les pages, une cellule vide reste vide (None)."""
    par_page = [_cellules_des_lignes(mots) for mots in pages]
    colonnes = _colonnes([ligne for page in par_page for ligne in page])
    if not colonnes:
        return []
    pages_lignes = []
    for page in par_page:
        lignes = []
        for ligne in page:
            valeurs: list[str | None] = [None] * len(colonnes)
            for cellule in ligne:
                index = _colonne_de(cellule, colonnes)
                valeurs[index] = f"{valeurs[index]} {cellule[2]}" if valeurs[index] else cellule[2]
            lignes.append(tuple(valeurs))
        pages_lignes.append(lignes)
    return _sans_entetes_repetes(pages_lignes)


def lignes_pdf(pages: Sequence[Sequence[Sequence[str | None]]]) -> list[tuple]:
    """Lignes de la feuille « PDF » à partir des lignes des tableaux tracés de chaque page."""
    return _sans_entetes_repetes(
        [[tuple(_cellule_pdf(c) for c in ligne) for ligne in page] for page in pages]
    )


def _sans_entetes_repetes(pages: Sequence[Sequence[tuple]]) -> list[tuple]:
    """Pages mises bout à bout, sans lignes vides ; une ligne d'une page suivante identique à une
    ligne de la première page qui ne contient aucun chiffre (l'en-tête répété en haut de page) est
    retirée."""
    lignes: list[tuple] = []
    entetes: set[tuple] = set()
    for numero, page in enumerate(pages):
        for ligne in page:
            if all(cellule is None for cellule in ligne):
                continue
            cle = tuple(normalize_header(c) if c else "" for c in ligne)
            sans_chiffre = not any(ch.isdigit() for c in ligne if c for ch in c)
            if numero == 0:
                if sans_chiffre:
                    entetes.add(cle)
            elif cle in entetes:
                continue
            lignes.append(ligne)
    return lignes


def _mots(page) -> list[MotPdf]:
    return [
        MotPdf(m["x0"], m["x1"], m["top"], m["bottom"], m["text"])
        for m in page.extract_words(keep_blank_chars=False, use_text_flow=False)
    ]


def _read_pdf(content: bytes, feuille: str | None, unite: str) -> tuple[list[str], str, list]:
    """Tableaux tracés s'il y en a, sinon les colonnes déduites de la position des mots (beaucoup
    de relevés n'ont pas de traits verticaux)."""
    if feuille not in (None, FEUILLE_PDF):
        raise ConflictError(f"Feuille introuvable dans le fichier : « {feuille} ».")
    try:
        document = pdfplumber.open(BytesIO(content))
        with document:
            if len(document.pages) > MAX_PAGES_PDF:
                raise ConflictError(f"PDF trop long : {MAX_PAGES_PDF} pages au plus par {unite}.")
            mots = [_mots(page) for page in document.pages]
            tableaux = [_tableaux_page(page) for page in document.pages]
    except ConflictError:
        raise
    except Exception as error:  # pdfminer lève des erreurs variées sur un PDF abîmé
        raise ConflictError("Fichier illisible : ce n'est pas un PDF valide.") from error
    if not any(mots):
        raise ConflictError(
            "PDF sans texte (document scanné) : SIMTIS ne lit que les PDF générés par la banque. "
            "Demandez le relevé en Excel ou en PDF d'origine."
        )
    rows = lignes_pdf(tableaux) if any(tableaux) else lignes_par_mots(mots)
    if not any(sum(cellule is not None for cellule in row) > 1 for row in rows):
        raise ConflictError("Aucun tableau trouvé dans le PDF : demandez le relevé en Excel.")
    _too_long(rows, unite)
    return [FEUILLE_PDF], FEUILLE_PDF, rows


def _xls_value(cell: xlrd.sheet.Cell, datemode: int) -> object:
    """Valeur d'une cellule .xls comme openpyxl la donnerait pour un .xlsx."""
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
        except (xlrd.xldate.XLDateError, ValueError, OverflowError):
            return cell.value
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return None
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
        return int(cell.value)
    return cell.value


def _read_xls(content: bytes, feuille: str | None, unite: str) -> tuple[list[str], str, list]:
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    except (
        xlrd.XLRDError,
        xlrd.compdoc.CompDocError,
        OSError,
        ValueError,
        AssertionError,
    ) as error:
        raise ConflictError(_ILLISIBLE) from error
    try:
        names = list(workbook.sheet_names())
        name = _sheet_name(names, feuille)
        sheet = workbook.sheet_by_name(name)
        rows: list[tuple] = []
        for index in range(sheet.nrows):
            rows.append(tuple(_xls_value(cell, workbook.datemode) for cell in sheet.row(index)))
            _too_long(rows, unite)
    finally:
        workbook.release_resources()
    return names, name, rows


# Abréviations courantes des en-têtes de banques et d'exports (08/10/2026, ex. AWB « DT opération ») :
# un mot entier de l'en-tête normalisé est remplacé avant la comparaison aux synonymes
ABREVIATIONS = {
    "dt": "date",
    "val": "valeur",
    "lib": "libelle",
    "mnt": "montant",
    "mt": "montant",
    "ope": "operation",
    "oper": "operation",
    "deb": "debit",
    "cred": "credit",
    "ech": "echeance",
}


def expand_abbreviations(header: str) -> str:
    """« dt operation » → « date operation » : chaque mot abrégé connu est développé."""
    return " ".join(ABREVIATIONS.get(word, word) for word in header.split())


def field_for_header(header: str, fields: Sequence[ImportField]) -> tuple[str | None, int]:
    """Champ reconnu pour un en-tête, et la force de la correspondance (2 exacte, 1 début, 0 aucune).
    L'en-tête est lu tel quel, puis avec ses abréviations développées (« DT valeur »)."""
    if not header:
        return None, 0
    code, strength = _field_for(header, fields)
    expanded = expand_abbreviations(header)
    if strength < 2 and expanded != header:
        other, other_strength = _field_for(expanded, fields)
        if other_strength > strength:
            return other, other_strength
    return code, strength


def _field_for(header: str, fields: Sequence[ImportField]) -> tuple[str | None, int]:
    for item in fields:
        if header in item.synonymes:
            return item.code, 2
    best, best_length = None, 0
    for item in fields:
        for synonym in item.synonymes:
            if header.startswith(synonym + " ") and len(synonym) > best_length:
                best, best_length = item.code, len(synonym)
    return best, 1 if best else 0


def detect_header(rows: list[tuple], fields: Sequence[ImportField]) -> int:
    """Index de la ligne d'en-tête : celle qui reconnaît le plus de champs, parmi les premières."""
    best_index, best_score = None, 0
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        score = sum(1 for cell in row if field_for_header(normalize_header(cell), fields)[1])
        if score > best_score:
            best_index, best_score = index, score
    if best_index is not None and best_score >= 2:
        return best_index
    return next(i for i, row in enumerate(rows) if not all(is_blank(cell) for cell in row))


@dataclass
class IgnoredLine:
    """Ligne du fichier qui n'est ni importée ni en erreur, avec la raison (aperçu seulement)."""

    numero: int  # numéro de la ligne dans le fichier Excel
    raison: str
    cellules: list[str]


def ignored_line(numero: int, raison: str, row: tuple) -> IgnoredLine:
    """Ligne ignorée telle qu'elle s'affiche : cellules en texte, sans les vides de la fin."""
    cellules = [_display(value) for value in row]
    while cellules and not cellules[-1]:
        cellules.pop()
    return IgnoredLine(numero=numero, raison=raison, cellules=cellules)


def _display(value: object) -> str:
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return clean_text(value) or ""


def columns_of(header: tuple, data: list[tuple]) -> list[Column]:
    width = max([len(header), *(len(row) for row in data)])
    columns = []
    for index in range(width):
        values = [row[index] for row in data if index < len(row) and not is_blank(row[index])]
        title = clean_text(header[index]) if index < len(header) else None
        columns.append(
            Column(
                index=index,
                lettre=get_column_letter(index + 1),
                entete=title or "",
                exemples=[_display(value) for value in values[:SAMPLES]],
            )
        )
    return columns


def propose_mapping(columns: list[Column], fields: Sequence[ImportField]) -> Mapping:
    """Correspondance détectée : les correspondances exactes d'abord, puis les débuts d'en-tête."""
    mapping: Mapping = dict.fromkeys(item.code for item in fields)
    for strength in (2, 1):
        for column in columns:
            if column.index in mapping.values():
                continue
            code, found = field_for_header(normalize_header(column.entete), fields)
            if code and found == strength and mapping[code] is None:
                mapping[code] = column.index
    return mapping


def mapping_from_headers(
    saved: dict[str, str | None], columns: list[Column], fields: Sequence[ImportField]
) -> Mapping | None:
    """Modèle mémorisé (champ → en-tête), s'il retrouve toutes ses colonnes dans le fichier. Un champ
    que le modèle ne connaît pas est complété par la détection, si le fichier a cette colonne
    (09/10/2026 : une colonne Échéance ou N° pièce ajoutée à l'export n'est plus ignorée)."""
    by_header = {
        normalize_header(column.entete): column.index for column in columns if column.entete
    }
    mapping: Mapping = dict.fromkeys(item.code for item in fields)
    for code, header in saved.items():
        if code not in mapping or header is None:
            continue
        index = by_header.get(normalize_header(header))
        if index is None:
            return None
        mapping[code] = index
    used = {index for index in mapping.values() if index is not None}
    for code, index in propose_mapping(columns, fields).items():
        if mapping.get(code) is None and index is not None and index not in used:
            mapping[code] = index
            used.add(index)
    return mapping


def mapping_errors(
    mapping: Mapping,
    width: int,
    fields: Sequence[ImportField],
    *,
    amount_codes: tuple[str, str, str] = ("debit", "credit", "montant"),
) -> list[str]:
    labels = {item.code: item.libelle for item in fields}
    errors = []
    for code, index in mapping.items():
        if index is not None and not 0 <= index < width:
            errors.append(f"{labels[code]} : colonne inexistante dans le fichier.")
    used = Counter(index for index in mapping.values() if index is not None)
    for index, count in sorted(used.items()):
        if count > 1 and 0 <= index < width:
            errors.append(
                f"La colonne {get_column_letter(index + 1)} est associée à plusieurs champs."
            )
    for item in fields:
        if item.obligatoire and mapping.get(item.code) is None:
            errors.append(f"Colonne obligatoire non associée : {item.libelle}.")
    debit, credit, montant = amount_codes
    has_debit_credit = mapping.get(debit) is not None or mapping.get(credit) is not None
    if mapping.get(montant) is not None and has_debit_credit:
        errors.append("Choisissez Débit et Crédit, ou Montant signé, pas les deux.")
    elif mapping.get(montant) is None and not has_debit_credit:
        errors.append("Associez au moins une colonne de montant : Débit, Crédit ou Montant signé.")
    return errors


def headers_of(mapping: Mapping, columns: list[Column]) -> dict[str, str] | None:
    """En-têtes à mémoriser (champ → en-tête) ; `None` si une colonne associée n'a pas d'en-tête."""
    headers = {}
    for code, index in mapping.items():
        if index is None:
            continue
        if not columns[index].entete:
            return None
        headers[code] = columns[index].entete
    return headers
