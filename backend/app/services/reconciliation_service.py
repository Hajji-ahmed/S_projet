"""Rapprochement bancaire 1→1 (P11) : lancement du moteur, décisions des utilisateurs, lectures.

Le moteur (`reconciliation_scoring`) ne fait que PROPOSER. Seul un utilisateur autorisé valide,
rejette ou annule une correspondance, et chaque décision est tracée dans `audit_logs`.

Statuts des opérations et des écritures :
- proposée ou ambiguë → « À vérifier » ;
- validée par un utilisateur (ou rapprochée à la main) → « Rapprochée » ;
- rejetée ou annulée → « Non rapprochée ». Une paire rejetée n'est reproposée qu'au prochain « Lancer le rapprochement ».
« Écart » est réservé aux écarts (P12) : le moteur n'y touche pas.
"""

from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.exc import IntegrityError, OperationalError
from sqlalchemy.orm import Session

from app.models import (
    AccountingEntry,
    BankTransaction,
    Company,
    ReconciliationMatch,
    ReconciliationMatchItem,
)
from app.repositories import discrepancy_repository
from app.repositories import reconciliation_repository as repo
from app.services import audit_service
from app.services.errors import ConflictError, NotFoundError
from app.services.position_service import BUSINESS_TZ, business_today
from app.services.reconciliation_groupes import (
    ResultatGroupes,
    explication_ensemble,
    grouper,
    score_ensemble,
    type_ensemble,
)
from app.services.reconciliation_scoring import (
    AUTRE,
    GRILLE_PAR_DEFAUT,
    Ecriture,
    Grille,
    Operation,
    Resultat,
    Score,
    _candidat,
    _ecart_si_comparable,
    comparable,
    meme_montant,
    proposer,
    proposer_elargi,
    score,
    type_ecriture,
    type_operation,
)
from app.services.reconciliation_scoring import type_operation as _type_de

PAGE_SIZE = 50
# Garde-fou contre une année mal saisie (« 0026 ») : 5 ans au plus par lancement (08/10/2026)
PERIODE_MAX_JOURS = 1827
NB_CANDIDATS = 20
COMMENTAIRE_MAX = 500
ENTITE = "reconciliation_match"
CONCURRENT = "Le rapprochement vient d'être modifié par un autre utilisateur : rechargez la page."


# --- Grille ---------------------------------------------------------------------------------------


# Critères dont la règle désactivée vaut 0 point (et non sa valeur par défaut)
_CRITERES_REGLES = {
    "REFERENCE": "reference",
    "MONTANT": "montant",
    "MONTANT_ROND": "montant_rond",
    "DATE": "date",
    "LIBELLE": "libelle",
    "TIERS": "tiers",
    "BONUS_SANS_REFERENCE": "bonus_sans_reference",
}


def grille(db: Session) -> Grille:
    """Grille lue dans `reconciliation_rules`.

    Un critère désactivé (`actif = false`) vaut 0 point ; un seuil ou une fenêtre désactivé ou
    absent garde sa valeur par défaut, comme un critère absent de la table.
    """
    valeurs: dict[str, object] = {}
    for rule in repo.all_rules(db):
        if rule.code in _CRITERES_REGLES:
            valeurs[_CRITERES_REGLES[rule.code]] = rule.poids if rule.actif else Decimal("0")
            if rule.code == "DATE" and rule.actif and rule.tolerance is not None:
                valeurs["tolerance_jours"] = int(rule.tolerance)
            continue
        if not rule.actif:
            continue
        match rule.code:
            case "FENETRE":
                if rule.tolerance is not None:
                    valeurs["fenetre_jours"] = int(rule.tolerance)
            case "FENETRE_ELARGIE":
                if rule.tolerance is not None:
                    valeurs["fenetre_elargie_jours"] = int(rule.tolerance)
            case "FENETRE_ESPECES":
                if rule.tolerance is not None:
                    valeurs["fenetre_especes_jours"] = int(rule.tolerance)
            case "SEUIL_PROPOSITION" | "SEUIL_FORT" | "ECART_AMBIGUITE":
                valeurs[rule.code.lower()] = rule.poids
    return Grille(**{**GRILLE_PAR_DEFAUT.__dict__, **valeurs})


def _operation(tx: BankTransaction) -> Operation:
    return Operation(
        id=tx.id,
        date_operation=tx.date_operation,
        date_valeur=tx.date_valeur,
        libelle=tx.libelle,
        reference=tx.reference,
        montant=tx.montant,
        bank_account_id=tx.bank_account_id,
    )


def montant_rapprochement(entry: AccountingEntry) -> Decimal:
    """Montant d'une écriture face au relevé : en devise pour un compte EUR / USD (colonne
    « Devise » de Sage, 14/10/2026), sinon le montant MAD. Le montant devise n'est lu que pour un
    compte en devise : une écriture d'un compte MAD n'en a jamais."""
    return entry.montant_devise if entry.montant_devise is not None else entry.montant


def _ecriture(entry: AccountingEntry) -> Ecriture:
    return Ecriture(
        id=entry.id,
        date_ecriture=entry.date_ecriture,
        libelle=entry.libelle,
        reference=entry.reference,
        numero_piece=entry.numero_piece,
        tiers=entry.tiers,
        montant=montant_rapprochement(entry),
        bank_account_id=entry.bank_account_id,
        echeance=entry.echeance,
    )


def _detail_json(resultat: Score, second: Decimal | None = None) -> dict[str, str]:
    """Points par critère ; `second` : score du meilleur autre candidat (concurrente proche)."""
    detail = {code: str(value) for code, value in resultat.detail.items()}
    if second is not None:
        detail["second"] = str(second)
    return detail


def score_second(match: ReconciliationMatch) -> Decimal | None:
    """Score de la 2e meilleure écriture (ou opération) au moment de la proposition."""
    value = (match.detail_score or {}).get("second")
    return None if value is None else Decimal(str(value))


EN_COURS = (
    "Un rapprochement ou une validation est déjà en cours pour cette société : réessayez dans "
    "un instant."
)


def _company(
    db: Session, company_id: int, *, lock: bool = False, sans_attendre: bool = False
) -> Company:
    if lock and sans_attendre:
        # Un lancement ne s'empile pas derrière un autre (15/10/2026) : 409 aussitôt
        try:
            with db.begin_nested():
                company = repo.lock_company(db, company_id, nowait=True)
        except OperationalError as error:
            if getattr(error.orig, "sqlstate", None) == "55P03":
                raise ConflictError(EN_COURS) from error
            raise
    else:
        company = repo.lock_company(db, company_id) if lock else db.get(Company, company_id)
    if company is None or not company.actif:
        raise NotFoundError("Société introuvable.")
    return company


def _check_period(date_from: date | None, date_to: date | None) -> None:
    if date_from and date_to and date_from > date_to:
        raise ConflictError("La date de début doit précéder la date de fin.")


def _commit(db: Session) -> None:
    try:
        db.commit()
    except IntegrityError as error:  # une même ligne retenue en même temps par une autre action
        db.rollback()
        raise ConflictError(CONCURRENT) from error


# --- Lancement du moteur --------------------------------------------------------------------------


@dataclass
class RunResult:
    nb_operations: int
    nb_ecritures: int
    nb_propositions: int
    nb_fortes: int
    nb_operations_ambigues: int
    nb_ecritures_ambigues: int
    # Groupes N → 1 proposés (12/10/2026)
    nb_groupes: int = 0


def run(
    db: Session,
    *,
    company_id: int,
    bank_account_id: int | None,
    date_from: date,
    date_to: date,
    acteur_id: int,
    ip: str | None,
) -> RunResult:
    """Propose des correspondances pour les opérations de la période.

    Les propositions automatiques encore en attente de la période sont remplacées ; les
    correspondances validées ou manuelles ne sont jamais modifiées. Les rejets en vigueur de la
    société sont d'abord levés (28/10/2026) : une paire rejetée peut être reproposée, mais
    seulement à ce lancement, jamais par une relance après un rapprochement manuel.
    """
    _check_period(date_from, date_to)
    if (date_to - date_from).days >= PERIODE_MAX_JOURS:
        raise ConflictError("La période ne peut pas dépasser 5 ans.")
    company = _company(db, company_id, lock=True, sans_attendre=True)
    accounts = repo.company_accounts(db, company.id, bank_account_id)
    if bank_account_id is not None and not accounts:
        raise NotFoundError("Compte bancaire introuvable pour cette société.")
    account_ids = [account.id for account in accounts]
    rejets_leves = repo.lever_rejets(db, company.id)
    regles = grille(db)
    moteur = _engine(db, company, account_ids, date_from, date_to, regles)
    result = RunResult(
        nb_operations=moteur.nb_operations,
        nb_ecritures=moteur.nb_ecritures,
        nb_propositions=len(moteur.resultat.propositions),
        nb_fortes=sum(
            1 for p in moteur.resultat.propositions if p.score.total >= regles.seuil_fort
        ),
        nb_operations_ambigues=len(moteur.resultat.operations_ambigues),
        nb_ecritures_ambigues=len(moteur.resultat.ecritures_ambigues),
        nb_groupes=len(moteur.groupes.groupes),
    )
    audit_service.log(
        db,
        user_id=acteur_id,
        action="rapprochement_lance",
        entite="company",
        entite_id=company.id,
        apres={
            "compte": bank_account_id,
            "du": date_from,
            "au": date_to,
            "propositions_remplacees": moteur.nb_remplacees,
            "rejets_leves": rejets_leves,
            **result.__dict__,
        },
        ip=ip,
    )
    _commit(db)
    return result


@dataclass(frozen=True)
class RunElargiResult:
    nb_operations: int
    nb_ecritures: int
    nb_propositions: int
    fenetre_jours: int


def run_elargi(
    db: Session,
    *,
    company_id: int,
    bank_account_id: int | None,
    date_from: date,
    date_to: date,
    acteur_id: int,
    ip: str | None,
) -> RunElargiResult:
    """Recherche élargie (20/10/2026) : parmi les lignes encore libres de la période, les paires
    au montant exactement opposé et au même n° jusqu'à `fenetre_elargie_jours` d'écart
    (`proposer_elargi`). Seules ses propres propositions en attente sont refaites ; celles du
    moteur normal ne bougent pas. Jamais validée d'office."""
    _check_period(date_from, date_to)
    if (date_to - date_from).days >= PERIODE_MAX_JOURS:
        raise ConflictError("La période ne peut pas dépasser 5 ans.")
    company = _company(db, company_id, lock=True, sans_attendre=True)
    accounts = repo.company_accounts(db, company.id, bank_account_id)
    if bank_account_id is not None and not accounts:
        raise NotFoundError("Compte bancaire introuvable pour cette société.")
    account_ids = [account.id for account in accounts]
    regles = grille(db)
    db.flush()

    anciennes = [
        match
        for match in repo.automatic_proposals_of_period(db, account_ids, date_from, date_to)
        if (match.detail_score or {}).get("elargie")
    ]
    anciennes_tx = {
        i.bank_transaction_id for m in anciennes for i in m.items if i.bank_transaction_id
    }
    anciennes_ec = {
        i.accounting_entry_id for m in anciennes for i in m.items if i.accounting_entry_id
    }
    repo.delete_matches(db, [match.id for match in anciennes])
    repo.set_transaction_status(db, anciennes_tx, "Non rapprochée")
    repo.set_entry_status(db, anciennes_ec, "Non rapprochée")
    db.expire_all()

    lecture = timedelta(days=regles.fenetre_elargie_jours)
    transactions = repo.free_transactions(db, account_ids, date_from, date_to)
    frais = repo.fee_transaction_ids(db, [tx.id for tx in transactions])
    transactions = [tx for tx in transactions if tx.id not in frais]
    entries = repo.free_entries(db, company.id, account_ids, date_from - lecture, date_to + lecture)
    by_tx = {tx.id: tx for tx in transactions}
    by_entry = {entry.id: entry for entry in entries}
    propositions = proposer_elargi(
        [_operation(tx) for tx in transactions],
        [_ecriture(entry) for entry in entries],
        regles,
        repo.rejected_pairs(db, company.id),
    )
    for proposition in propositions:
        tx, entry = by_tx[proposition.operation_id], by_entry[proposition.ecriture_id]
        repo.add(
            db,
            ReconciliationMatch(
                company_id=company.id,
                type="1-1",
                score=proposition.score.total,
                statut="Proposée",
                origine="Automatique",
                detail_score={
                    **_detail_json(proposition.score),
                    "elargie": True,
                    "ecart_jours": proposition.ecart_jours,
                },
                items=[
                    ReconciliationMatchItem(
                        bank_transaction_id=tx.id, montant_affecte=abs(tx.montant), actif=True
                    ),
                    ReconciliationMatchItem(
                        accounting_entry_id=entry.id,
                        montant_affecte=abs(montant_rapprochement(entry)),
                        actif=True,
                    ),
                ],
            ),
        )
    repo.set_transaction_status(db, {p.operation_id for p in propositions}, "À vérifier")
    repo.set_entry_status(db, {p.ecriture_id for p in propositions}, "À vérifier")
    result = RunElargiResult(
        nb_operations=len(transactions),
        nb_ecritures=len(entries),
        nb_propositions=len(propositions),
        fenetre_jours=regles.fenetre_elargie_jours,
    )
    audit_service.log(
        db,
        user_id=acteur_id,
        action="rapprochement_elargi",
        entite="company",
        entite_id=company.id,
        apres={
            "compte": bank_account_id,
            "du": date_from,
            "au": date_to,
            "propositions_remplacees": len(anciennes),
            **result.__dict__,
        },
        ip=ip,
    )
    _commit(db)
    return result


@dataclass
class _Moteur:
    resultat: Resultat
    nb_operations: int
    nb_ecritures: int
    nb_remplacees: int
    groupes: ResultatGroupes = field(default_factory=ResultatGroupes)


def _engine(
    db: Session,
    company: Company,
    account_ids: list[int],
    date_from: date,
    date_to: date,
    regles: Grille,
) -> _Moteur:
    """Refait les propositions automatiques en attente des comptes sur la période (sans commit) :
    les correspondances validées, rejetées ou manuelles ne sont jamais modifiées."""
    db.flush()
    # Écritures lues jusqu'à la fenêtre des n° (chèques, effets), la plus large (12/10/2026)
    lecture = max(regles.fenetre_jours, regles.fenetre_numero_jours)
    entry_from = date_from - timedelta(days=lecture)
    entry_to = date_to + timedelta(days=lecture)

    # 1. Les propositions en attente de la période sont refaites
    anciennes = [
        match
        for match in repo.automatic_proposals_of_period(db, account_ids, date_from, date_to)
        # Une proposition de la recherche élargie se refait par son propre bouton (20/10/2026)
        if not (match.detail_score or {}).get("elargie")
    ]
    anciennes_tx = {
        i.bank_transaction_id for m in anciennes for i in m.items if i.bank_transaction_id
    }
    anciennes_ec = {
        i.accounting_entry_id for m in anciennes for i in m.items if i.accounting_entry_id
    }
    repo.delete_matches(db, [match.id for match in anciennes])
    repo.set_transaction_status(db, anciennes_tx, "Non rapprochée")
    repo.set_entry_status(db, anciennes_ec, "Non rapprochée")
    repo.reset_unmatched_to_check(
        db, company.id, account_ids, date_from, date_to, entry_from, entry_to
    )
    db.expire_all()

    # 2. Le moteur compare ce qui reste libre
    transactions = repo.free_transactions(db, account_ids, date_from, date_to)
    entries = repo.free_entries(db, company.id, account_ids, entry_from, entry_to)
    by_tx = {tx.id: tx for tx in transactions}
    by_entry = {entry.id: entry for entry in entries}
    # Les frais bancaires et leur TVA se comptabilisent dans SIMTIS (15/10/2026) : jamais
    # proposés seuls face à une écriture ; ils n'entrent que dans un groupe, en complément d'une
    # opération principale
    frais_ids = repo.fee_transaction_ids(db, [tx.id for tx in transactions])
    resultat = proposer(
        [_operation(tx) for tx in transactions if tx.id not in frais_ids],
        [_ecriture(entry) for entry in entries],
        regles,
        repo.rejected_pairs(db, company.id),
    )

    # 3. Groupes N → 1 (12/10/2026) sur ce que le 1 → 1 laisse libre : un groupe équilibré passe
    # avant une paire 1 → 1 aux montants différents (qui ne pourrait jamais être validée)
    equilibrees = [
        p
        for p in resultat.propositions
        if by_tx[p.operation_id].montant == -montant_rapprochement(by_entry[p.ecriture_id])
    ]
    bloque_tx = {p.operation_id for p in equilibrees} | resultat.operations_ambigues
    bloque_ec = {p.ecriture_id for p in equilibrees} | resultat.ecritures_ambigues
    groupes = grouper(
        [_operation(tx) for tx in transactions if tx.id not in bloque_tx],
        [_ecriture(entry) for entry in entries if entry.id not in bloque_ec],
        regles,
        repo.rejected_pairs(db, company.id),
        frais_ids,
    )
    groupe_tx = {i for g in groupes.groupes for i in g.operation_ids}
    groupe_ec = {g.ecriture_id for g in groupes.groupes}
    resultat.propositions = [
        p
        for p in resultat.propositions
        if p in equilibrees or (p.operation_id not in groupe_tx and p.ecriture_id not in groupe_ec)
    ]

    # 4. Chaque proposition est enregistrée, en attente d'une décision humaine
    for proposition in resultat.propositions:
        _enregistrer_proposition(
            db,
            company,
            by_tx[proposition.operation_id],
            by_entry[proposition.ecriture_id],
            proposition,
        )
    a_verifier_tx = {p.operation_id for p in resultat.propositions} | resultat.operations_ambigues
    a_verifier_ec = {p.ecriture_id for p in resultat.propositions} | resultat.ecritures_ambigues

    for groupe in groupes.groupes:
        entry = by_entry[groupe.ecriture_id]
        repo.add(
            db,
            ReconciliationMatch(
                company_id=company.id,
                type="N-1",
                score=groupe.score.total,
                statut="Proposée",
                origine="Automatique",
                detail_score={**_detail_json(groupe.score), "explication": groupe.explication},
                items=[
                    *(
                        ReconciliationMatchItem(
                            bank_transaction_id=tx_id,
                            montant_affecte=abs(by_tx[tx_id].montant),
                            actif=True,
                        )
                        for tx_id in groupe.operation_ids
                    ),
                    ReconciliationMatchItem(
                        accounting_entry_id=entry.id,
                        montant_affecte=abs(montant_rapprochement(entry)),
                        actif=True,
                    ),
                ],
            ),
        )
        a_verifier_tx.update(groupe.operation_ids)
        a_verifier_ec.add(entry.id)

    repo.set_transaction_status(db, a_verifier_tx, "À vérifier")
    repo.set_entry_status(db, a_verifier_ec, "À vérifier")
    db.flush()
    return _Moteur(resultat, len(transactions), len(entries), len(anciennes), groupes)


def _enregistrer_proposition(
    db: Session,
    company: Company,
    tx: BankTransaction,
    entry: AccountingEntry,
    proposition,
) -> None:
    repo.add(
        db,
        ReconciliationMatch(
            company_id=company.id,
            type="1-1",
            score=proposition.score.total,
            statut="Proposée",
            origine="Automatique",
            detail_score=_detail_json(proposition.score, proposition.second),
            items=[
                ReconciliationMatchItem(
                    bank_transaction_id=tx.id, montant_affecte=abs(tx.montant), actif=True
                ),
                ReconciliationMatchItem(
                    accounting_entry_id=entry.id,
                    montant_affecte=abs(montant_rapprochement(entry)),
                    actif=True,
                ),
            ],
        ),
    )


@dataclass
class _Voisinage:
    """Candidates d'un ensemble de lignes, au-dessus du seuil de proposition et hors paires
    rejetées : opérations et écritures, par id."""

    operations: dict[int, BankTransaction] = field(default_factory=dict)
    ecritures: dict[int, AccountingEntry] = field(default_factory=dict)


def _voisins(
    db: Session,
    tx_ids: set[int],
    entry_ids: set[int],
    regles: Grille,
    rejetees: set[tuple[int, int]],
) -> _Voisinage:
    voisins = _Voisinage()
    for tx_id in sorted(tx_ids):
        _, _, items, seuil, _ = candidates(db, tx_id, regles=regles, rejetees=rejetees)
        for item in items:
            if eligible(item, seuil):
                voisins.ecritures[item.entry.id] = item.entry
    for entry_id in sorted(entry_ids):
        _, items, seuil = candidates_of_entry(db, entry_id, regles=regles, rejetees=rejetees)
        for item in items:
            if not item.rejetee:
                voisins.operations[item.transaction.id] = item.transaction
    return voisins


def _relance_ciblee(
    db: Session,
    company: Company,
    tx_ids: set[int],
    entry_ids: set[int],
    regles: Grille,
    liberees: tuple[set[int], set[int]] = (set(), set()),
) -> Resultat:
    """Relance du moteur réduite aux lignes en concurrence avec celles qu'on vient de rapprocher
    (16/10/2026), pour le rapprochement depuis les ambiguës : la relance du compte entier sur ±
    la fenêtre prenait 13 s pour une seule opération.

    Les lignes concernées sont les candidates des lignes rapprochées encore « À vérifier » sans
    correspondance (elles étaient à égalité avec elles), et leurs propres candidates dans le même
    cas (la suite de l'égalité), ainsi que les lignes `liberees` (leur proposition contenait une
    ligne rapprochée et vient d'être rejetée). Elles sont comparées, avec les mêmes règles que le moteur, à leurs
    candidates libres (sans correspondance active ; une ligne déjà proposée ailleurs n'est jamais
    reprise) : une égalité les laisse ambiguës, une meilleure candidate est proposée, sinon elles
    repassent « Non rapprochée ». Les autres lignes du compte ne bougent pas."""
    db.flush()
    rejetees = repo.rejected_pairs(db, company.id)

    def libres_a_verifier(voisins: _Voisinage) -> tuple[set[int], set[int]]:
        actives_tx = repo.active_matches_by_transaction(db, list(voisins.operations))
        actives_ec = repo.active_matches_by_entry(db, list(voisins.ecritures))
        return (
            {
                i
                for i, tx in voisins.operations.items()
                if tx.statut == "À vérifier" and i not in actives_tx
            },
            {
                i
                for i, entry in voisins.ecritures.items()
                if entry.statut == "À vérifier" and i not in actives_ec
            },
        )

    premiers = _voisins(db, tx_ids, entry_ids, regles, rejetees)
    concernees_tx, concernees_ec = libres_a_verifier(premiers)
    concernees_tx |= liberees[0] - tx_ids
    concernees_ec |= liberees[1] - entry_ids
    seconds = _voisins(db, concernees_tx, concernees_ec, regles, rejetees)
    suite_tx, suite_ec = libres_a_verifier(seconds)
    concernees_tx |= suite_tx - tx_ids
    concernees_ec |= suite_ec - entry_ids
    if not concernees_tx and not concernees_ec:
        return Resultat()
    # Contexte : les lignes concernées et toutes leurs candidates libres
    contexte = _voisins(
        db, concernees_tx - set(premiers.operations), concernees_ec - set(premiers.ecritures),
        regles, rejetees,
    )  # fmt: skip
    operations = {**premiers.operations, **seconds.operations, **contexte.operations}
    ecritures = {**premiers.ecritures, **seconds.ecritures, **contexte.ecritures}
    for tx_id in concernees_tx - set(operations):
        found = repo.get_transaction(db, tx_id)
        if found is not None:
            operations[tx_id] = found[0]
    for entry_id in concernees_ec - set(ecritures):
        entry = repo.get_entry(db, entry_id)
        if entry is not None:
            ecritures[entry_id] = entry
    actives_tx = repo.active_matches_by_transaction(db, list(operations))
    actives_ec = repo.active_matches_by_entry(db, list(ecritures))
    frais = repo.fee_transaction_ids(db, list(operations))
    operations = {i: tx for i, tx in operations.items() if i not in actives_tx and i not in frais}
    ecritures = {i: entry for i, entry in ecritures.items() if i not in actives_ec}

    repo.set_transaction_status(db, concernees_tx, "Non rapprochée")
    repo.set_entry_status(db, concernees_ec, "Non rapprochée")
    resultat = proposer(
        [_operation(tx) for tx in operations.values()],
        [_ecriture(entry) for entry in ecritures.values()],
        regles,
        rejetees,
    )
    retenues = [
        p
        for p in resultat.propositions
        if p.operation_id in concernees_tx or p.ecriture_id in concernees_ec
    ]
    for proposition in retenues:
        _enregistrer_proposition(
            db,
            company,
            operations[proposition.operation_id],
            ecritures[proposition.ecriture_id],
            proposition,
        )
    resultat.propositions = retenues
    resultat.operations_ambigues &= concernees_tx
    resultat.ecritures_ambigues &= concernees_ec
    repo.set_transaction_status(
        db,
        {p.operation_id for p in retenues} | resultat.operations_ambigues,
        "À vérifier",
    )
    repo.set_entry_status(
        db,
        {p.ecriture_id for p in retenues} | resultat.ecritures_ambigues,
        "À vérifier",
    )
    db.flush()
    return resultat


# --- Décisions -----------------------------------------------------------------------------------


def _locked_match(db: Session, match_id: int) -> ReconciliationMatch:
    """Correspondance verrouillée, après le verrou de sa société (même ordre que `run`)."""
    match = repo.get_match(db, match_id)
    if match is None:
        raise NotFoundError("Correspondance introuvable.")
    _company(db, match.company_id, lock=True)
    match = repo.get_match(db, match_id, lock=True)
    if match is None:
        raise NotFoundError("Correspondance introuvable.")
    db.refresh(match)
    return match


def _members(match: ReconciliationMatch) -> tuple[set[int], set[int]]:
    tx_ids = {item.bank_transaction_id for item in match.items if item.bank_transaction_id}
    entry_ids = {item.accounting_entry_id for item in match.items if item.accounting_entry_id}
    return tx_ids, entry_ids


def _snapshot(match: ReconciliationMatch) -> dict:
    tx_ids, entry_ids = _members(match)
    return {
        "statut": match.statut,
        "origine": match.origine,
        "score": match.score,
        "code": match.code,
        "operations": sorted(tx_ids),
        "ecritures": sorted(entry_ids),
        "commentaire": match.commentaire,
    }


def _decide(match: ReconciliationMatch, statut: str, acteur_id: int) -> None:
    """Nouvelle décision humaine : son auteur et sa date restent dans l'historique."""
    match.statut = statut
    match.decide_par_id = acteur_id
    match.decide_le = datetime.now(UTC)


def _release(
    db: Session,
    match: ReconciliationMatch,
    statut: str,
    commentaire: str | None,
    acteur_id: int,
) -> None:
    """Rejette ou annule : les éléments sont libérés et redeviennent « Non rapprochée »."""
    _decide(match, statut, acteur_id)
    if commentaire is not None:
        match.commentaire = commentaire
    for item in match.items:
        item.actif = False
    tx_ids, entry_ids = _members(match)
    db.flush()
    repo.set_transaction_status(db, tx_ids, "Non rapprochée")
    repo.set_entry_status(db, entry_ids, "Non rapprochée")


MONTANTS_DIFFERENTS = (
    "Les montants de l'opération et de l'écriture diffèrent : ce rapprochement ne peut pas être "
    "validé. Signalez un écart « Montant différent »."
)


def _check_amounts(db: Session, match: ReconciliationMatch) -> None:
    """Un rapprochement validé s'équilibre : la somme des montants (banque + Sage) vaut zéro."""
    tx_ids, entry_ids = _members(match)
    total = sum((db.get(BankTransaction, i).montant for i in tx_ids), Decimal("0")) + sum(
        (montant_rapprochement(db.get(AccountingEntry, i)) for i in entry_ids), Decimal("0")
    )
    if total != 0:
        raise ConflictError(MONTANTS_DIFFERENTS)


def _validate(db: Session, match: ReconciliationMatch, acteur_id: int) -> None:
    _check_amounts(db, match)
    _decide(match, "Validée", acteur_id)
    match.valide_par_id = acteur_id
    match.valide_le = match.decide_le
    tx_ids, entry_ids = _members(match)
    if match.code is None:
        match.code = _new_code(db, min(tx_ids))
    db.flush()
    repo.set_transaction_status(db, tx_ids, "Rapprochée")
    repo.set_entry_status(db, entry_ids, "Rapprochée")


def _new_code(db: Session, transaction_id: int) -> str:
    """Code unique du rapprochement (13/10/2026) : « BP-26-000123 », un compteur par compte
    bancaire et par année de validation (date du Maroc)."""
    tx = db.get(BankTransaction, transaction_id)
    annee = business_today().year
    numero = repo.next_code_number(db, tx.bank_account_id, annee)
    banque = repo.bank_code_of_account(db, tx.bank_account_id)
    return f"{banque}-{annee % 100:02d}-{numero:06d}"


def create_validated(
    db: Session,
    *,
    company_id: int,
    transactions: list[BankTransaction],
    entry: AccountingEntry,
    origine: str,
    detail: dict,
    acteur_id: int,
) -> ReconciliationMatch:
    """Correspondance validée d'emblée entre des opérations et une écriture (comptabilisation des
    frais, 13/10/2026) : mêmes contrôles et même code qu'une validation humaine."""
    match = ReconciliationMatch(
        company_id=company_id,
        type="N-1" if len(transactions) > 1 else "1-1",
        statut="Proposée",
        origine=origine,
        detail_score=detail,
        items=[
            *(
                ReconciliationMatchItem(
                    bank_transaction_id=tx.id, montant_affecte=abs(tx.montant), actif=True
                )
                for tx in transactions
            ),
            ReconciliationMatchItem(
                accounting_entry_id=entry.id,
                montant_affecte=abs(montant_rapprochement(entry)),
                actif=True,
            ),
        ],
    )
    repo.add(db, match)
    db.flush()
    _validate(db, match, acteur_id)
    return match


ECRITURE_SIMTIS = (
    "Cette écriture a été comptabilisée par SIMTIS : supprimez-la depuis « À comptabiliser » "
    "(possible tant qu'elle n'est pas exportée pour Sage)."
)


def _clean_comment(commentaire: str | None) -> str | None:
    text = " ".join(commentaire.split()) if commentaire else ""
    if len(text) > COMMENTAIRE_MAX:
        raise ConflictError(f"Le commentaire ne peut pas dépasser {COMMENTAIRE_MAX} caractères.")
    return text or None


def validate(db: Session, match_id: int, *, acteur_id: int, ip: str | None) -> ReconciliationMatch:
    """Un utilisateur confirme une proposition du moteur."""
    match = _locked_match(db, match_id)
    if match.statut != "Proposée":
        raise ConflictError(
            f"Seule une correspondance proposée se valide (statut : {match.statut})."
        )
    avant = _snapshot(match)
    _validate(db, match, acteur_id)
    audit_service.log(
        db,
        user_id=acteur_id,
        action="validation_rapprochement",
        entite=ENTITE,
        entite_id=match.id,
        avant=avant,
        apres=_snapshot(match),
        ip=ip,
    )
    _commit(db)
    return match


def validate_batch(
    db: Session, match_ids: list[int], *, acteur_id: int, ip: str | None
) -> list[ReconciliationMatch]:
    """Valide en une fois les propositions choisies par l'utilisateur, toutes ou aucune.

    Décision du 07/10/2026 : la sélection peut contenir des propositions faibles, cochées une à
    une par l'utilisateur (l'écran ne coche d'office que les fortes). Chacune doit être encore
    « Proposée » et s'équilibrer (même montant, sens opposé) ; sinon rien n'est validé.
    """
    ids = sorted(set(match_ids))
    found = repo.get_matches(db, ids)
    if len(found) != len(ids):
        raise NotFoundError("Correspondance introuvable.")
    companies = {match.company_id for match in found}
    if len(companies) != 1:
        raise ConflictError("Une validation en lot ne porte que sur une seule société.")
    _company(db, companies.pop(), lock=True)
    matches = repo.get_matches(db, ids, lock=True)
    seuil = grille(db).seuil_fort
    for match in matches:
        db.refresh(match)
        if match.statut != "Proposée":
            raise ConflictError(
                f"La correspondance n° {match.id} n'est plus proposée (statut : {match.statut}) : "
                "rien n'a été validé, rechargez la liste."
            )
        try:
            _check_amounts(db, match)
        except ConflictError as error:
            raise ConflictError(
                f"Correspondance n° {match.id} : {error.message} Rien n'a été validé."
            ) from error
    for match in matches:
        avant = _snapshot(match)
        _validate(db, match, acteur_id)
        audit_service.log(
            db,
            user_id=acteur_id,
            action="validation_rapprochement",
            entite=ENTITE,
            entite_id=match.id,
            avant=avant,
            apres={
                **_snapshot(match),
                "en_lot": True,
                "forte": match.score is not None and match.score >= seuil,
            },
            ip=ip,
        )
    _commit(db)
    return matches


def reject(
    db: Session, match_id: int, *, commentaire: str | None, acteur_id: int, ip: str | None
) -> ReconciliationMatch:
    """Un utilisateur refuse une proposition : la paire ne sera plus proposée avant le prochain lancement du rapprochement."""
    match = _locked_match(db, match_id)
    if match.statut != "Proposée":
        raise ConflictError(
            f"Seule une correspondance proposée se rejette (statut : {match.statut})."
        )
    avant = _snapshot(match)
    _release(db, match, "Rejetée", _clean_comment(commentaire), acteur_id)
    audit_service.log(
        db,
        user_id=acteur_id,
        action="rejet_rapprochement",
        entite=ENTITE,
        entite_id=match.id,
        avant=avant,
        apres=_snapshot(match),
        ip=ip,
    )
    _commit(db)
    return match


def reject_batch(
    db: Session,
    match_ids: list[int],
    *,
    commentaire: str | None,
    acteur_id: int,
    ip: str | None,
) -> int:
    """Rejette en une fois les propositions choisies (15/10/2026), toutes ou aucune : chacune doit
    être encore « Proposée ». Comme un rejet un par un, la paire n'est plus proposée avant le prochain lancement et ses
    lignes repassent « Non rapprochée »."""
    texte = _clean_comment(commentaire)
    ids = sorted(set(match_ids))
    found = repo.get_matches(db, ids)
    if len(found) != len(ids):
        raise NotFoundError("Correspondance introuvable.")
    companies = {match.company_id for match in found}
    if len(companies) != 1:
        raise ConflictError("Un rejet en lot ne porte que sur une seule société.")
    _company(db, companies.pop(), lock=True)
    matches = repo.get_matches(db, ids, lock=True)
    for match in matches:
        db.refresh(match)
        if match.statut != "Proposée":
            raise ConflictError(
                f"La correspondance n° {match.id} n'est plus proposée (statut : {match.statut}) : "
                "rien n'a été rejeté, rechargez la liste."
            )
    for match in matches:
        avant = _snapshot(match)
        _release(db, match, "Rejetée", texte, acteur_id)
        audit_service.log(
            db,
            user_id=acteur_id,
            action="rejet_rapprochement",
            entite=ENTITE,
            entite_id=match.id,
            avant=avant,
            apres={**_snapshot(match), "en_lot": True},
            ip=ip,
        )
    _commit(db)
    return len(matches)


def cancel(
    db: Session, match_id: int, *, motif: str, acteur_id: int, ip: str | None
) -> ReconciliationMatch:
    """Annule un rapprochement validé ; le motif est obligatoire et la trace conservée."""
    texte = _clean_comment(motif)
    if texte is None:
        raise ConflictError("Indiquez le motif de l'annulation.")
    match = _locked_match(db, match_id)
    if match.statut != "Validée":
        raise ConflictError(
            f"Seul un rapprochement validé s'annule (statut : {match.statut}) ; "
            "une proposition se rejette."
        )
    _, entry_ids = _members(match)
    if any(db.get(AccountingEntry, i).origine == "SIMTIS" for i in entry_ids):
        raise ConflictError(ECRITURE_SIMTIS)
    avant = _snapshot(match)
    _release(db, match, "Annulée", texte, acteur_id)
    audit_service.log(
        db,
        user_id=acteur_id,
        action="annulation_rapprochement",
        entite=ENTITE,
        entite_id=match.id,
        avant=avant,
        apres=_snapshot(match),
        ip=ip,
    )
    _commit(db)
    return match


def release_pending_matches(
    db: Session,
    *,
    transaction_id: int | None,
    entry_id: int | None,
    motif: str,
    acteur_id: int,
    ip: str | None,
) -> list[int]:
    """Rejette (et trace) les propositions en attente qui contiennent l'opération ou l'écriture.

    Un rapprochement validé n'est jamais défait ici : 409, à annuler d'abord. Ne valide pas la
    transaction ; retourne les identifiants des propositions rejetées.
    """
    actives = repo.active_matches_of(db, transaction_id=transaction_id, entry_id=entry_id)
    for active in actives:
        if active.statut == "Validée":
            tx_ids, _ = _members(active)
            quoi = "L'opération" if transaction_id in tx_ids else "L'écriture"
            raise ConflictError(f"{quoi} est déjà rapprochée : annulez d'abord ce rapprochement.")
    remplacees = []
    for active in actives:
        avant = _snapshot(active)
        _release(db, active, "Rejetée", motif, acteur_id)
        audit_service.log(
            db,
            user_id=acteur_id,
            action="rejet_rapprochement",
            entite=ENTITE,
            entite_id=active.id,
            avant=avant,
            apres=_snapshot(active),
            ip=ip,
        )
        remplacees.append(active.id)
    return remplacees


# Au plus 50 lignes de chaque côté dans un rapprochement manuel (P13, 15/10/2026)
MANUEL_MAX = 50


def match_manually(
    db: Session,
    *,
    transaction_ids: list[int],
    entry_ids: list[int],
    commentaire: str | None,
    acteur_id: int,
    ip: str | None,
    relance_ciblee: bool = False,
) -> ReconciliationMatch:
    """Rapprochement choisi par l'utilisateur, validé d'emblée (origine « Manuelle ») : une ou
    plusieurs opérations avec une ou plusieurs écritures (1-1, 1-N, N-1, N-N ; P13, 15/10/2026).

    Exige un seul compte bancaire et des totaux égaux en sens opposé (montant devise pour un
    compte EUR / USD), des lignes libres (une proposition en attente qui contient l'une d'elles est
    rejetée et tracée ; un rapprochement validé ne l'est jamais) et des types compatibles. Tout est
    vérifié avant la première écriture.
    """
    texte = _clean_comment(commentaire)
    tx_ids, ecr_ids = sorted(set(transaction_ids)), sorted(set(entry_ids))
    if not tx_ids or not ecr_ids:
        raise ConflictError("Choisissez au moins une opération bancaire et une écriture comptable.")
    if len(tx_ids) > MANUEL_MAX or len(ecr_ids) > MANUEL_MAX:
        raise ConflictError(
            f"Un rapprochement manuel compte au plus {MANUEL_MAX} lignes de chaque côté."
        )

    lignes = []
    for tx_id in tx_ids:
        found = repo.get_transaction(db, tx_id)
        if found is None:
            raise NotFoundError("Opération bancaire introuvable.")
        lignes.append(found)
    comptes = {account.id for _, account in lignes}
    if len(comptes) > 1:
        raise ConflictError("Les opérations choisies sont sur plusieurs comptes bancaires.")
    account = lignes[0][1]
    company = _company(db, account.company_id, lock=True)
    operations = []
    for tx, _ in lignes:
        db.refresh(tx)
        operations.append(tx)
    entries = []
    for entry_id in ecr_ids:
        entry = repo.get_entry(db, entry_id)
        if entry is None:
            raise NotFoundError("Écriture comptable introuvable.")
        db.refresh(entry)
        entries.append(entry)

    if any(entry.company_id != company.id for entry in entries):
        raise ConflictError(
            "L'opération et l'écriture appartiennent à deux sociétés différentes : "
            "elles ne se rapprochent jamais."
        )
    if any(entry.bank_account_id != account.id for entry in entries):
        raise ConflictError(
            "L'écriture n'est pas passée dans le journal Sage de ce compte bancaire."
        )
    if any(tx.montant == 0 for tx in operations):
        raise ConflictError("Une opération à zéro ne se rapproche pas.")
    if any(tx.en_cours for tx in operations):
        raise ConflictError(
            "Une opération en cours (« (*) ») ne se rapproche pas : attendez sa définitive."
        )
    total_banque = sum((tx.montant for tx in operations), Decimal("0"))
    total_sage = sum((montant_rapprochement(entry) for entry in entries), Decimal("0"))
    if total_banque != -total_sage:
        raise ConflictError(
            "Les montants doivent être égaux, en sens opposé (un crédit en banque correspond à "
            f"un débit dans Sage) : écart de {abs(total_banque + total_sage):.2f}."
        )
    if any("Écart" == ligne.statut for ligne in [*operations, *entries]):
        raise ConflictError("Une ligne en écart se traite dans l'écran Écarts.")
    types = {type_operation(tx.libelle) for tx in operations} | {
        type_ecriture(entry.libelle) for entry in entries
    }
    types.discard(None)
    if len(types) > 1:
        if len(operations) == 1 and len(entries) == 1:
            raise ConflictError(
                f"Une opération {type_operation(operations[0].libelle)} ne se rapproche pas "
                f"d'une écriture {type_ecriture(entries[0].libelle)}."
            )
        raise ConflictError(
            f"Les lignes choisies mélangent des types différents ({', '.join(sorted(types))})."
        )

    remplacees: list[int] = []
    for tx in operations:
        remplacees += release_pending_matches(
            db,
            transaction_id=tx.id,
            entry_id=None,
            motif="Remplacée par un rapprochement manuel.",
            acteur_id=acteur_id,
            ip=ip,
        )
    for entry in entries:
        remplacees += release_pending_matches(
            db,
            transaction_id=None,
            entry_id=entry.id,
            motif="Remplacée par un rapprochement manuel.",
            acteur_id=acteur_id,
            ip=ip,
        )

    regles = grille(db)
    ops = [_operation(tx) for tx in operations]
    ecritures = [_ecriture(entry) for entry in entries]
    if len(ops) == 1 and len(ecritures) == 1:
        resultat = score(ops[0], ecritures[0], regles)
        detail = _detail_json(resultat)
    else:
        resultat = score_ensemble(ops, ecritures, regles)
        detail = {**_detail_json(resultat), "explication": explication_ensemble(ops, ecritures)}
    match = ReconciliationMatch(
        company_id=company.id,
        type=type_ensemble(len(ops), len(ecritures)),
        score=resultat.total,
        statut="Proposée",
        origine="Manuelle",
        detail_score=detail,
        commentaire=texte,
        items=[
            *(
                ReconciliationMatchItem(
                    bank_transaction_id=tx.id, montant_affecte=abs(tx.montant), actif=True
                )
                for tx in operations
            ),
            *(
                ReconciliationMatchItem(
                    accounting_entry_id=entry.id,
                    montant_affecte=abs(montant_rapprochement(entry)),
                    actif=True,
                )
                for entry in entries
            ),
        ],
    )
    repo.add(db, match)
    db.flush()
    _validate(db, match, acteur_id)
    snapshot = _snapshot(match)
    # Les lignes qui convoitaient les mêmes opérations ou écritures sont départagées à nouveau
    # (08/10/2026) : relance du moteur sur ce compte, autour des lignes rapprochées
    # Depuis les ambiguës (16/10/2026) : seulement les lignes qui étaient en concurrence
    if relance_ciblee:
        liberees: tuple[set[int], set[int]] = (set(), set())
        for match_id in remplacees:
            remplacee = repo.get_match(db, match_id)
            if remplacee is not None:
                membres = _members(remplacee)
                liberees[0].update(membres[0])
                liberees[1].update(membres[1])
        resultat = _relance_ciblee(
            db, company, {tx.id for tx in operations}, {e.id for e in entries}, regles, liberees
        )
    else:
        marge = timedelta(days=regles.fenetre_jours)
        dates = [tx.date_operation for tx in operations] + [e.date_ecriture for e in entries]
        resultat = _engine(
            db, company, [account.id], min(dates) - marge, max(dates) + marge, regles
        ).resultat
    audit_service.log(
        db,
        user_id=acteur_id,
        action="rapprochement_manuel",
        entite=ENTITE,
        entite_id=match.id,
        apres={
            **snapshot,
            "propositions_remplacees": sorted(set(remplacees)),
            "relance": {
                "ciblee": relance_ciblee,
                "propositions": len(resultat.propositions),
                "operations_ambigues": len(resultat.operations_ambigues),
            },
        },
        ip=ip,
    )
    _commit(db)
    return match


# --- Lectures ------------------------------------------------------------------------------------


@dataclass
class MatchView:
    match: ReconciliationMatch
    transaction: BankTransaction
    bank_code: str
    entry: AccountingEntry
    entry_bank_code: str | None
    valide_par: str | None
    forte: bool
    decide_par: str | None = None
    # Critères de la grille actuelle : le détail affiche ceux-ci, plus ceux qui ont rapporté des points
    actifs: tuple[str, ...] = ()
    # Une autre écriture (ou opération) à moins de `ecart_ambiguite` points (08/10/2026)
    score_second: Decimal | None = None
    concurrente_proche: bool = False
    # Recherche élargie (20/10/2026) : écart en jours ; jamais cochée d'office
    elargie: bool = False
    ecart_jours: int | None = None
    # Toutes les opérations (une seule en 1 → 1, plusieurs pour un groupe N → 1), avec leur banque
    transactions: list[tuple[BankTransaction, str]] = field(default_factory=list)
    # Toutes les écritures (plusieurs en 1 → N et N → N, P13), la plus grosse en tête
    entries: list[tuple[AccountingEntry, str | None]] = field(default_factory=list)
    # Nombres et totaux de TOUTES les lignes : `transactions` et `entries` n'en gardent que
    # `LIGNES_AFFICHEES` (une reprise en compte des milliers, 27/10/2026)
    nb_operations: int = 1
    nb_ecritures: int = 1
    total_operations: Decimal | None = None
    total_ecritures: Decimal | None = None


@dataclass
class TransactionsPage:
    total: int
    page: int
    taille: int
    par_statut: dict[str, int]
    operations: list[tuple[BankTransaction, str, ReconciliationMatch | None]]
    # Écart ouvert de chaque opération de la page qui en a un (P12)
    ecarts: dict[int, int]
    # Frais bancaires non rapprochés du filtre : nombre et total (12/10/2026)
    a_comptabiliser: int = 0
    total_a_comptabiliser: Decimal = Decimal(0)
    # Opérations en cours du filtre (20/10/2026)
    en_cours: int = 0


@dataclass
class Candidate:
    entry: AccountingEntry
    bank_code: str | None
    score: Score
    rejetee: bool
    proposee_ailleurs: bool
    actifs: tuple[str, ...] = ()
    # Montant exactement opposé : seule une telle candidate peut être proposée (28/10/2026)
    meme_montant: bool = True


def eligible(item: "Candidate | CandidateOperation", seuil: Decimal) -> bool:
    """Candidate que le moteur pourrait proposer : au-dessus du seuil, non rejetée et au même
    montant (28/10/2026 : un montant différent n'est jamais proposé ni rendu ambigu)."""
    return item.score.total >= seuil and not item.rejetee and item.meme_montant


def list_transactions(
    db: Session,
    company_id: int,
    *,
    bank_account_id: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    statut: str | None = None,
    q: str | None = None,
    page: int = 1,
    sans_a_verifier: bool = False,
    a_comptabiliser: bool = False,
    type_operation: str | None = None,
    sens: str | None = None,
) -> TransactionsPage:
    """Volet « Transactions bancaires » : 50 par page ; le décompte par statut porte sur tout le
    filtre, hors filtre de statut.

    `sans_a_verifier` : les opérations « À vérifier » ne sont pas listées (elles se traitent dans
    l'onglet Propositions, décision du 07/10/2026) ; elles restent comptées dans `par_statut`.
    """
    company = _company(db, company_id)
    _check_period(date_from, date_to)
    query = repo.transactions_query(
        company.id,
        bank_account_id=bank_account_id,
        date_from=date_from,
        date_to=date_to,
        q=q,
        type_operation=type_operation,
        sens=sens,
    )
    # Opérations en cours (20/10/2026) : listées avec leur badge, comptées à part
    en_cours = repo.count_rows(db, query.where(BankTransaction.en_cours.is_(True)))
    comptees = query.where(BankTransaction.en_cours.is_(False))
    par_statut = repo.count_by_status(db, comptees)
    # Frais bancaires non rapprochés : « À comptabiliser », comptés à part (12/10/2026)
    nb_frais, total_frais = repo.a_comptabiliser(db, comptees)
    par_statut["Non rapprochée"] = par_statut.get("Non rapprochée", 0) - nb_frais
    listed = repo.without_status(query, "À vérifier") if sans_a_verifier else query
    if a_comptabiliser:
        listed = listed.where(BankTransaction.statut == "Non rapprochée", repo.frais_bancaire())
        statut = None
    elif statut == "Non rapprochée":
        listed = listed.where(~repo.frais_bancaire())
    total, rows = repo.page_of_transactions(
        db, listed, statut=statut, offset=(page - 1) * PAGE_SIZE, limit=PAGE_SIZE
    )
    matches = repo.active_match_by_transaction(db, [tx.id for tx, _ in rows])
    return TransactionsPage(
        total=total,
        page=page,
        taille=PAGE_SIZE,
        par_statut=par_statut,
        a_comptabiliser=nb_frais,
        total_a_comptabiliser=total_frais,
        en_cours=en_cours,
        operations=[(tx, code, matches.get(tx.id)) for tx, code in rows],
        ecarts=discrepancy_repository.open_discrepancy_by_transaction(
            db, [tx.id for tx, _ in rows]
        ),
    )


# Lignes envoyées par côté pour un rapprochement : une reprise en compte des milliers (27/10/2026)
LIGNES_AFFICHEES = 20


def _views(db: Session, matches: list[ReconciliationMatch], regles: Grille) -> list[MatchView]:
    tx_ids = {i.bank_transaction_id for m in matches for i in m.items if i.bank_transaction_id}
    entry_ids = {i.accounting_entry_id for m in matches for i in m.items if i.accounting_entry_id}
    transactions = repo.transactions_with_bank(db, tx_ids)
    entries = repo.entries_with_bank(db, entry_ids)
    names = repo.user_names(
        db,
        {m.valide_par_id for m in matches if m.valide_par_id}
        | {m.decide_par_id for m in matches if m.decide_par_id},
    )
    views = []
    for match in matches:
        second = score_second(match)
        # Groupe N → 1 : toutes ses opérations ; la principale (la plus grosse) est affichée en tête
        lignes = [transactions[i.bank_transaction_id] for i in match.items if i.bank_transaction_id]
        lignes.sort(key=lambda ligne: (-abs(ligne[0].montant), ligne[0].id))
        tx, code = lignes[0]
        ecritures = [entries[i.accounting_entry_id] for i in match.items if i.accounting_entry_id]
        ecritures.sort(key=lambda ligne: (-abs(ligne[0].montant), ligne[0].id))
        entry, entry_code = ecritures[0]
        views.append(
            MatchView(
                match=match,
                transaction=tx,
                transactions=lignes[:LIGNES_AFFICHEES],
                entries=ecritures[:LIGNES_AFFICHEES],
                nb_operations=len(lignes),
                nb_ecritures=len(ecritures),
                total_operations=sum((t.montant for t, _ in lignes), Decimal("0")),
                total_ecritures=sum((montant_rapprochement(e) for e, _ in ecritures), Decimal("0")),
                bank_code=code,
                entry=entry,
                entry_bank_code=entry_code,
                valide_par=names.get(match.valide_par_id) if match.valide_par_id else None,
                decide_par=names.get(match.decide_par_id) if match.decide_par_id else None,
                forte=match.score is not None and match.score >= regles.seuil_fort,
                actifs=regles.criteres_actifs,
                score_second=second,
                concurrente_proche=second is not None
                and match.score is not None
                and match.score - second < regles.ecart_ambiguite,
                elargie=bool((match.detail_score or {}).get("elargie")),
                ecart_jours=(match.detail_score or {}).get("ecart_jours"),
            )
        )
    return views


def list_matches(
    db: Session,
    company_id: int,
    *,
    statut: str = "Proposée",
    bank_account_id: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    type_operation: str | None = None,
) -> tuple[Decimal, list[MatchView]]:
    """Correspondances 1→1 de la société (par défaut celles en attente), avec le seuil « forte »."""
    company = _company(db, company_id)
    _check_period(date_from, date_to)
    regles = grille(db)
    matches = repo.matches_of_period(
        db,
        company.id,
        statut=statut,
        bank_account_id=bank_account_id,
        date_from=date_from,
        date_to=date_to,
    )
    views = _views(db, matches, regles)
    if type_operation is not None:
        # Type de l'opération principale (la seule en 1 → 1, la plus grosse d'un groupe)
        voulu = None if type_operation == AUTRE else type_operation
        views = [view for view in views if _type_de(view.transaction.libelle) == voulu]
    return regles.seuil_fort, views


@dataclass
class HistoryPage:
    total: int
    page: int
    taille: int
    par_statut: dict[str, int]
    seuil_fort: Decimal
    decisions: list[MatchView]


def list_history(
    db: Session,
    company_id: int,
    *,
    statut: str | None = None,
    bank_account_id: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    page: int = 1,
) -> HistoryPage:
    """Historique des décisions (validées, rejetées, annulées), de la plus récente à la plus
    ancienne, 50 par page ; le décompte par statut porte sur tout le filtre, hors statut."""
    company = _company(db, company_id)
    _check_period(date_from, date_to)
    regles = grille(db)
    query = repo.decisions_query(
        company.id, bank_account_id=bank_account_id, date_from=date_from, date_to=date_to
    )
    total, matches = repo.page_of_decisions(
        db, query, statut=statut, offset=(page - 1) * PAGE_SIZE, limit=PAGE_SIZE
    )
    return HistoryPage(
        total=total,
        page=page,
        taille=PAGE_SIZE,
        par_statut=repo.count_decisions_by_status(db, query),
        seuil_fort=regles.seuil_fort,
        decisions=_views(db, matches, regles),
    )


def get_match_view(db: Session, match_id: int) -> MatchView:
    match = repo.get_match(db, match_id)
    if match is None:
        raise NotFoundError("Correspondance introuvable.")
    return _views(db, [match], grille(db))[0]


def candidates(
    db: Session,
    transaction_id: int,
    *,
    regles: Grille | None = None,
    rejetees: set[tuple[int, int]] | None = None,
) -> tuple[BankTransaction, str, list[Candidate], Decimal, Decimal]:
    """Écritures qui pourraient correspondre à l'opération, de la meilleure à la moins bonne.

    Même compte, sens opposé, dans la fenêtre de dates ; une écriture déjà rapprochée n'apparaît
    pas. Retourne aussi le seuil de proposition et le seuil « forte ».
    """
    found = repo.get_transaction(db, transaction_id)
    if found is None:
        raise NotFoundError("Opération bancaire introuvable.")
    tx, account = found
    regles = regles or grille(db)
    jours = [tx.date_operation] + ([tx.date_valeur] if tx.date_valeur else [])
    courte = timedelta(days=regles.fenetre_jours)
    large = timedelta(days=max(regles.fenetre_jours, regles.fenetre_numero_jours))
    entries = repo.candidate_entries(
        db,
        tx,
        account.company_id,
        min(jours) - courte,
        max(jours) + courte,
        large_from=min(jours) - large,
        large_to=max(jours) + large,
    )
    operation = _operation(tx)
    # Seules les écritures comparables sont lues plus loin : la fenêtre de 180 jours (même numéro)
    # ramène des centaines d'écritures par opération (13/10/2026, 20 s pour 50 ambiguës)
    comparables = [
        (entry, ecriture)
        for entry in entries
        if comparable(operation, ecriture := _ecriture(entry), regles)
    ]
    if rejetees is None:
        rejetees = repo.rejected_pairs(db, account.company_id)
    actives = repo.active_matches_by_entry(db, [entry.id for entry, _ in comparables])
    codes = repo.entries_with_bank(db, {entry.id for entry, _ in comparables})
    scored = []
    for entry, ecriture in comparables:
        active = actives.get(entry.id)
        scored.append(
            Candidate(
                entry=entry,
                bank_code=codes[entry.id][1],
                score=score(operation, ecriture, regles),
                rejetee=(tx.id, entry.id) in rejetees,
                proposee_ailleurs=active is not None and tx.id not in _members(active)[0],
                actifs=regles.criteres_actifs,
                meme_montant=meme_montant(operation, ecriture),
            )
        )
    scored.sort(key=lambda item: (-item.score.total, item.entry.date_ecriture, item.entry.id))
    bank_code = repo.transactions_with_bank(db, {tx.id})[tx.id][1]
    return tx, bank_code, scored[:NB_CANDIDATS], regles.seuil_proposition, regles.seuil_fort


def a_egalite(items: list, seuil: Decimal) -> list:
    """Candidates d'une ligne ambiguë (16/10/2026) : celles au même meilleur score, au-dessus du
    seuil et non rejetées. Une ligne n'est ambiguë que par cette égalité ; « Aucune
    correspondance » n'écarte qu'elles."""
    eligibles = [item for item in items if eligible(item, seuil)]
    if not eligibles:
        return []
    meilleur = max(item.score.total for item in eligibles)
    return [item for item in eligibles if item.score.total == meilleur]


# Opérations ambiguës par page ; le total porte sur tout le filtre (08/10/2026)
PAGE_AMBIGUES = 50


@dataclass
class AmbiguousView:
    transaction: BankTransaction
    bank_code: str
    candidats: list[Candidate]


def list_ambiguous(
    db: Session,
    company_id: int,
    *,
    bank_account_id: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    page: int = 1,
    type_operation: str | None = None,
) -> tuple[Decimal, int, list[AmbiguousView]]:
    """Opérations ambiguës de la période (« À vérifier » sans proposition), chacune avec ses
    écritures candidates au-dessus du seuil de proposition, la meilleure d'abord (hors paires déjà
    rejetées). L'utilisateur en choisit une : rapprochement manuel. Retourne aussi le seuil
    « forte »."""
    company = _company(db, company_id)
    _check_period(date_from, date_to)
    regles = grille(db)
    views = []
    total, transactions = repo.ambiguous_transactions(
        db,
        company.id,
        bank_account_id=bank_account_id,
        date_from=date_from,
        date_to=date_to,
        offset=(page - 1) * PAGE_AMBIGUES,
        limit=PAGE_AMBIGUES,
        type_operation=type_operation,
    )
    rejetees = repo.rejected_pairs(db, company.id)
    for tx in transactions:
        _, bank_code, items, seuil_proposition, _ = candidates(
            db, tx.id, regles=regles, rejetees=rejetees
        )
        views.append(
            AmbiguousView(
                transaction=tx,
                bank_code=bank_code,
                candidats=a_egalite(items, seuil_proposition),
            )
        )
    return regles.seuil_fort, total, views


# Au plus 500 opérations ou écritures par « aucune correspondance » en lot (15/10/2026)
AMBIGUES_LOT_MAX = 500


def _paire_rejetee(
    company_id: int,
    tx: BankTransaction,
    entry: AccountingEntry,
    resultat: Score,
    commentaire: str,
    acteur_id: int,
) -> ReconciliationMatch:
    """Paire écartée par l'utilisateur : une correspondance « Rejetée » (éléments inactifs), le
    mécanisme des rejets ; la paire n'est reproposée qu'au prochain lancement."""
    match = ReconciliationMatch(
        company_id=company_id,
        type="1-1",
        score=resultat.total,
        statut="Proposée",
        origine="Manuelle",
        detail_score=_detail_json(resultat),
        commentaire=commentaire,
        items=[
            ReconciliationMatchItem(
                bank_transaction_id=tx.id, montant_affecte=abs(tx.montant), actif=False
            ),
            ReconciliationMatchItem(
                accounting_entry_id=entry.id,
                montant_affecte=abs(montant_rapprochement(entry)),
                actif=False,
            ),
        ],
    )
    _decide(match, "Rejetée", acteur_id)
    return match


def _ecarts_operation(db: Session, tx: BankTransaction) -> list[Candidate]:
    """Écritures candidates à écarter pour l'opération ; vérifie qu'elle s'y prête (sans écrire)."""
    db.refresh(tx)
    if tx.statut in ("Rapprochée", "Écart") or repo.active_matches_of(db, transaction_id=tx.id):
        raise ConflictError(
            f"L'opération « {tx.libelle} » a déjà un rapprochement ou une proposition : "
            "traitez-les d'abord."
        )
    _, _, items, seuil_proposition, _ = candidates(db, tx.id)
    # Plus aucune candidate (déjà écartées depuis l'autre côté) : la ligne repasse simplement
    # « Non rapprochée » (15/10/2026)
    return a_egalite(items, seuil_proposition)


def _liberer_ecritures_sans_candidate(db: Session, entry_ids: set[int]) -> list[int]:
    """Écritures « À vérifier » sans correspondance active qui n'ont plus aucune opération
    candidate (toutes écartées) : elles repassent « Non rapprochée » (15/10/2026). Sans cela, elles
    restaient dans « Écritures ambiguës » sans rien à choisir."""
    if not entry_ids:
        return []
    db.flush()
    regles = grille(db)
    rejetees: set[tuple[int, int]] | None = None
    liberees = []
    for entry_id in sorted(entry_ids):
        entry = repo.get_entry(db, entry_id)
        if entry is None or entry.statut != "À vérifier":
            continue
        if repo.active_matches_of(db, entry_id=entry.id):
            continue
        if rejetees is None:
            rejetees = repo.rejected_pairs(db, entry.company_id)
        _, items, seuil = candidates_of_entry(db, entry.id, regles=regles, rejetees=rejetees)
        if not any(eligible(item, seuil) for item in items):
            liberees.append(entry.id)
    repo.set_entry_status(db, set(liberees), "Non rapprochée")
    return liberees


def _liberer_operations_sans_candidate(db: Session, transaction_ids: set[int]) -> list[int]:
    """Le symétrique pour les opérations « À vérifier » qui n'ont plus d'écriture candidate."""
    if not transaction_ids:
        return []
    db.flush()
    regles = grille(db)
    rejetees: set[tuple[int, int]] | None = None
    liberees = []
    for tx_id in sorted(transaction_ids):
        found = repo.get_transaction(db, tx_id)
        if found is None or found[0].statut != "À vérifier":
            continue
        tx, account = found
        if repo.active_matches_of(db, transaction_id=tx.id):
            continue
        if rejetees is None:
            rejetees = repo.rejected_pairs(db, account.company_id)
        _, _, items, seuil, _ = candidates(db, tx.id, regles=regles, rejetees=rejetees)
        if not any(eligible(item, seuil) for item in items):
            liberees.append(tx.id)
    repo.set_transaction_status(db, set(liberees), "Non rapprochée")
    return liberees


def _ecarter_pour_operation(
    db: Session,
    company: Company,
    tx: BankTransaction,
    ecartees: list[Candidate],
    texte: str | None,
    acteur_id: int,
    ip: str | None,
    *,
    en_lot: bool,
) -> list[int]:
    for item in ecartees:
        repo.add(
            db,
            _paire_rejetee(
                company.id,
                tx,
                item.entry,
                item.score,
                texte or "Aucune écriture ne correspond à l'opération.",
                acteur_id,
            ),
        )
    db.flush()
    repo.set_transaction_status(db, {tx.id}, "Non rapprochée")
    ids = sorted(item.entry.id for item in ecartees)
    liberees = _liberer_ecritures_sans_candidate(db, set(ids))
    audit_service.log(
        db,
        user_id=acteur_id,
        action="rejet_ambiguite",
        entite="bank_transaction",
        entite_id=tx.id,
        avant={"statut": "À vérifier"},
        apres={
            "statut": "Non rapprochée",
            "ecritures_ecartees": ids,
            "ecritures_liberees": liberees,
            "commentaire": texte,
            "en_lot": en_lot,
        },
        ip=ip,
    )
    return ids


def _operations_d_une_societe(
    db: Session, transaction_ids: list[int]
) -> tuple[Company, list[BankTransaction]]:
    lignes = []
    for tx_id in sorted(set(transaction_ids)):
        found = repo.get_transaction(db, tx_id)
        if found is None:
            raise NotFoundError("Opération bancaire introuvable.")
        lignes.append(found)
    societes = {account.company_id for _, account in lignes}
    if len(societes) != 1:
        raise ConflictError("Un traitement en lot ne porte que sur une seule société.")
    company = _company(db, societes.pop(), lock=True)
    return company, [tx for tx, _ in lignes]


def rejeter_ambiguite(
    db: Session,
    transaction_id: int,
    *,
    commentaire: str | None,
    acteur_id: int,
    ip: str | None,
) -> list[int]:
    """« Aucune écriture ne correspond » (15/10/2026) : toutes les écritures candidates d'une
    opération ambiguë sont écartées d'un coup. Chaque paire est enregistrée comme rejetée (le
    mécanisme des rejets : reproposée au prochain lancement seulement), l'opération repasse « Non rapprochée ». Une
    écriture importée plus tard reste comparée normalement. Retourne les écritures écartées."""
    texte = _clean_comment(commentaire)
    company, [tx] = _operations_d_une_societe(db, [transaction_id])
    ecartees = _ecarts_operation(db, tx)
    ids = _ecarter_pour_operation(db, company, tx, ecartees, texte, acteur_id, ip, en_lot=False)
    _commit(db)
    return ids


def rejeter_ambiguites(
    db: Session,
    transaction_ids: list[int],
    *,
    commentaire: str | None,
    acteur_id: int,
    ip: str | None,
) -> int:
    """« Aucune correspondance » pour plusieurs opérations ambiguës, toutes ou aucune : tout est
    vérifié avant la première écriture. Retourne le nombre d'opérations traitées."""
    if not transaction_ids or len(set(transaction_ids)) > AMBIGUES_LOT_MAX:
        raise ConflictError(f"Choisissez de 1 à {AMBIGUES_LOT_MAX} opérations.")
    texte = _clean_comment(commentaire)
    company, operations = _operations_d_une_societe(db, transaction_ids)
    a_ecarter = [(tx, _ecarts_operation(db, tx)) for tx in operations]
    for tx, ecartees in a_ecarter:
        _ecarter_pour_operation(db, company, tx, ecartees, texte, acteur_id, ip, en_lot=True)
    _commit(db)
    return len(a_ecarter)


# --- Écritures ambiguës (15/10/2026) --------------------------------------------------------------


@dataclass
class CandidateOperation:
    transaction: BankTransaction
    bank_code: str
    score: Score
    rejetee: bool
    proposee_ailleurs: bool
    actifs: tuple[str, ...] = ()
    # Toujours vrai : `candidates_of_entry` ne garde que les candidates de `_candidat`
    meme_montant: bool = True


@dataclass
class AmbiguousEntryView:
    entry: AccountingEntry
    bank_code: str | None
    candidats: list[CandidateOperation]


def candidates_of_entry(
    db: Session,
    entry_id: int,
    *,
    regles: Grille | None = None,
    rejetees: set[tuple[int, int]] | None = None,
) -> tuple[AccountingEntry, list[CandidateOperation], Decimal]:
    """Opérations qui pourraient correspondre à l'écriture, au-dessus du seuil de proposition, de
    la meilleure à la moins bonne : le symétrique de `candidates` (même compte, sens opposé,
    fenêtre de dates des types). Retourne aussi le seuil de proposition."""
    entry = repo.get_entry(db, entry_id)
    if entry is None:
        raise NotFoundError("Écriture comptable introuvable.")
    regles = regles or grille(db)
    jours = [entry.date_ecriture] + ([entry.echeance] if entry.echeance else [])
    courte = timedelta(days=regles.fenetre_jours)
    large = timedelta(days=max(regles.fenetre_jours, regles.fenetre_numero_jours))
    transactions = repo.candidate_transactions(
        db,
        entry,
        min(jours) - courte,
        max(jours) + courte,
        large_from=min(jours) - large,
        large_to=max(jours) + large,
        montant=-montant_rapprochement(entry),
    )
    ecriture = _ecriture(entry)
    # Seules les opérations au-dessus du seuil de proposition sont gardées (tous les appelants n'en
    # retiennent pas d'autres) : le tri rapide du moteur écarte les autres sans calculer leur
    # libellé (15/10/2026)
    retenues = []
    for tx in transactions:
        operation = _operation(tx)
        ecart = _ecart_si_comparable(operation, ecriture, regles)
        if ecart is None:
            continue
        resultat = _candidat(operation, ecriture, regles, ecart)
        if resultat is not None:
            retenues.append((tx, resultat))
    if rejetees is None:
        rejetees = repo.rejected_pairs(db, entry.company_id)
    actives = repo.active_matches_by_transaction(db, [tx.id for tx, _ in retenues])
    codes = repo.transactions_with_bank(db, {tx.id for tx, _ in retenues})
    scored = []
    for tx, resultat in retenues:
        active = actives.get(tx.id)
        scored.append(
            CandidateOperation(
                transaction=tx,
                bank_code=codes[tx.id][1],
                score=resultat,
                rejetee=(tx.id, entry.id) in rejetees,
                proposee_ailleurs=active is not None and entry.id not in _members(active)[1],
                actifs=regles.criteres_actifs,
            )
        )
    scored.sort(
        key=lambda item: (-item.score.total, item.transaction.date_operation, item.transaction.id)
    )
    return entry, scored[:NB_CANDIDATS], regles.seuil_proposition


def _ecritures_a_egalite(
    db: Session,
    entries: list[AccountingEntry],
    regles: Grille,
    rejetees: set[tuple[int, int]],
) -> set[int]:
    """Parmi ces écritures, celles qui ont au moins deux opérations candidates au même meilleur
    score (au-dessus du seuil de proposition, hors paires rejetées) : la règle de
    `candidates_of_entry` + `a_egalite`, avec les opérations lues en une fois par compte au lieu
    d'une lecture par écriture (16/10/2026)."""
    courte = timedelta(days=regles.fenetre_jours)
    large = timedelta(days=max(regles.fenetre_jours, regles.fenetre_numero_jours))
    par_compte: dict[int, list[AccountingEntry]] = {}
    for entry in entries:
        par_compte.setdefault(entry.bank_account_id, []).append(entry)
    resultat: set[int] = set()
    for compte_id, ecritures in par_compte.items():
        bornes = [
            jour
            for entry in ecritures
            for jour in (entry.date_ecriture, entry.echeance)
            if jour is not None
        ]
        transactions = repo.transactions_comparables(
            db, compte_id, min(bornes) - large, max(bornes) + large
        )
        operations = {tx.id: _operation(tx) for tx in transactions}
        par_jour: dict[date, list[BankTransaction]] = {}
        par_montant: dict[Decimal, list[BankTransaction]] = {}
        for tx in transactions:
            for jour in {tx.date_operation, tx.date_valeur} - {None}:
                par_jour.setdefault(jour, []).append(tx)
            par_montant.setdefault(tx.montant, []).append(tx)
        for entry in ecritures:
            jours = [entry.date_ecriture] + ([entry.echeance] if entry.echeance else [])
            debut, fin = min(jours) - courte, max(jours) + courte
            large_debut, large_fin = min(jours) - large, max(jours) + large
            # Les opérations que lirait `repo.candidate_transactions` pour cette écriture
            vues: dict[int, BankTransaction] = {}
            jour = debut
            while jour <= fin:
                for tx in par_jour.get(jour, ()):
                    vues[tx.id] = tx
                jour += timedelta(days=1)
            for tx in par_montant.get(-montant_rapprochement(entry), ()):
                if any(
                    j is not None and large_debut <= j <= large_fin
                    for j in (tx.date_operation, tx.date_valeur)
                ):
                    vues[tx.id] = tx
            ecriture = _ecriture(entry)
            scores = []
            for tx in vues.values():
                if (tx.id, entry.id) in rejetees:
                    continue
                operation = operations[tx.id]
                ecart = _ecart_si_comparable(operation, ecriture, regles)
                if ecart is None:
                    continue
                score = _candidat(operation, ecriture, regles, ecart)
                if score is not None and score.total >= regles.seuil_proposition:
                    scores.append(score.total)
            if len(scores) >= 2 and scores.count(max(scores)) >= 2:
                resultat.add(entry.id)
    return resultat


def list_ambiguous_entries(
    db: Session,
    company_id: int,
    *,
    bank_account_id: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    page: int = 1,
    type_operation: str | None = None,
) -> tuple[int, list[AmbiguousEntryView]]:
    """Écritures ambiguës de la période, chacune avec ses opérations candidates au-dessus du
    seuil de proposition (hors paires rejetées). L'utilisateur en coche une (1 → 1) ou plusieurs
    dont la somme vaut l'écriture (N → 1).

    16/10/2026 : seule une écriture qui a au moins deux opérations au même meilleur score est
    listée. Une écriture « À vérifier » parce qu'une opération hésite entre elle et une autre
    écriture n'a rien à choisir de son côté : elle se traite depuis « Opérations ambiguës ». Le
    total et les pages portent sur ces seules écritures."""
    company = _company(db, company_id)
    _check_period(date_from, date_to)
    regles = grille(db)
    rejetees = repo.rejected_pairs(db, company.id)
    _, a_verifier = repo.ambiguous_entries(
        db,
        company.id,
        bank_account_id=bank_account_id,
        date_from=date_from,
        date_to=date_to,
        offset=0,
        limit=None,
        type_operation=type_operation,
    )
    egalites = _ecritures_a_egalite(db, a_verifier, regles, rejetees)
    ambigues = [entry for entry in a_verifier if entry.id in egalites]
    total = len(ambigues)
    entries = ambigues[(page - 1) * PAGE_AMBIGUES : page * PAGE_AMBIGUES]
    codes = repo.entries_with_bank(db, {entry.id for entry in entries})
    views = []
    for entry in entries:
        _, items, seuil = candidates_of_entry(db, entry.id, regles=regles, rejetees=rejetees)
        views.append(
            AmbiguousEntryView(
                entry=entry,
                bank_code=codes[entry.id][1],
                candidats=a_egalite(items, seuil),
            )
        )
    return total, views


def _ecarts_ecriture(db: Session, entry: AccountingEntry) -> list[CandidateOperation]:
    """Opérations candidates à écarter pour l'écriture ; vérifie qu'elle s'y prête (sans écrire)."""
    db.refresh(entry)
    if entry.statut in ("Rapprochée", "Écart") or repo.active_matches_of(db, entry_id=entry.id):
        raise ConflictError(
            f"L'écriture « {entry.libelle} » a déjà un rapprochement ou une proposition : "
            "traitez-les d'abord."
        )
    _, items, seuil = candidates_of_entry(db, entry.id)
    return a_egalite(items, seuil)


def _ecarter_pour_ecriture(
    db: Session,
    company: Company,
    entry: AccountingEntry,
    ecartees: list[CandidateOperation],
    texte: str | None,
    acteur_id: int,
    ip: str | None,
    *,
    en_lot: bool,
) -> list[int]:
    for item in ecartees:
        repo.add(
            db,
            _paire_rejetee(
                company.id,
                item.transaction,
                entry,
                item.score,
                texte or "Aucune opération ne correspond à l'écriture.",
                acteur_id,
            ),
        )
    db.flush()
    repo.set_entry_status(db, {entry.id}, "Non rapprochée")
    ids = sorted(item.transaction.id for item in ecartees)
    liberees = _liberer_operations_sans_candidate(db, set(ids))
    audit_service.log(
        db,
        user_id=acteur_id,
        action="rejet_ambiguite_ecriture",
        entite="accounting_entry",
        entite_id=entry.id,
        avant={"statut": "À vérifier"},
        apres={
            "statut": "Non rapprochée",
            "operations_ecartees": ids,
            "operations_liberees": liberees,
            "commentaire": texte,
            "en_lot": en_lot,
        },
        ip=ip,
    )
    return ids


def _ecritures_d_une_societe(
    db: Session, entry_ids: list[int]
) -> tuple[Company, list[AccountingEntry]]:
    entries = []
    for entry_id in sorted(set(entry_ids)):
        entry = repo.get_entry(db, entry_id)
        if entry is None:
            raise NotFoundError("Écriture comptable introuvable.")
        entries.append(entry)
    societes = {entry.company_id for entry in entries}
    if len(societes) != 1:
        raise ConflictError("Un traitement en lot ne porte que sur une seule société.")
    return _company(db, societes.pop(), lock=True), entries


def rejeter_ambiguite_ecriture(
    db: Session,
    entry_id: int,
    *,
    commentaire: str | None,
    acteur_id: int,
    ip: str | None,
) -> list[int]:
    """« Aucune opération ne correspond » : le symétrique de `rejeter_ambiguite`. Les opérations
    candidates sont écartées (paires rejetées jusqu'au prochain lancement), l'écriture repasse « Non
    rapprochée ». Retourne les opérations écartées."""
    texte = _clean_comment(commentaire)
    company, [entry] = _ecritures_d_une_societe(db, [entry_id])
    ecartees = _ecarts_ecriture(db, entry)
    ids = _ecarter_pour_ecriture(db, company, entry, ecartees, texte, acteur_id, ip, en_lot=False)
    _commit(db)
    return ids


def rejeter_ambiguites_ecritures(
    db: Session,
    entry_ids: list[int],
    *,
    commentaire: str | None,
    acteur_id: int,
    ip: str | None,
) -> int:
    """« Aucune correspondance » pour plusieurs écritures ambiguës, toutes ou aucune."""
    if not entry_ids or len(set(entry_ids)) > AMBIGUES_LOT_MAX:
        raise ConflictError(f"Choisissez de 1 à {AMBIGUES_LOT_MAX} écritures.")
    texte = _clean_comment(commentaire)
    company, entries = _ecritures_d_une_societe(db, entry_ids)
    a_ecarter = [(entry, _ecarts_ecriture(db, entry)) for entry in entries]
    for entry, ecartees in a_ecarter:
        _ecarter_pour_ecriture(db, company, entry, ecartees, texte, acteur_id, ip, en_lot=True)
    _commit(db)
    return len(a_ecarter)


# --- Export du lettrage (14/10/2026) ----------------------------------------------------------------

LETTRAGE_HEADERS = (
    "Code",
    "Côté",
    "Banque",
    "Date",
    "Journal",
    "N° pièce",
    "Libellé",
    "Débit",
    "Crédit",
    "Type",
    "Origine",
    "Score",
    "Validé par",
    "Validé le",
    # Compte en devise (14/10/2026) : le montant devise de la ligne Sage, celui rapproché
    "Montant devise",
)
_LARGEURS = (16, 8, 8, 12, 9, 16, 44, 14, 14, 6, 12, 7, 20, 17, 14)


def _heure_locale(moment: datetime | None) -> datetime | None:
    """Excel ne connaît pas les fuseaux : l'heure du Maroc, sans fuseau."""
    return moment.astimezone(BUSINESS_TZ).replace(tzinfo=None) if moment else None


def export_lettrage(
    db: Session,
    company_id: int,
    *,
    bank_account_id: int | None,
    date_from: date | None,
    date_to: date | None,
    acteur_id: int,
    ip: str | None,
) -> tuple[str, bytes, int]:
    """Lettrage du filtre de l'historique en classeur Excel : tous les rapprochements VALIDÉS dont
    une opération est dans la période (et du compte, s'il est donné). Une ligne par élément,
    regroupées par code : les opérations du relevé, puis les écritures Sage ; chaque code
    s'équilibre (débit d'un côté = crédit de l'autre). Ligne de total en bas."""
    company = _company(db, company_id)
    _check_period(date_from, date_to)
    if date_from and date_to and (date_to - date_from).days > PERIODE_MAX_JOURS:
        raise ConflictError("La période ne peut pas dépasser 5 ans.")
    query = repo.decisions_query(
        company.id, bank_account_id=bank_account_id, date_from=date_from, date_to=date_to
    )
    matches = repo.validated_matches(db, query)
    tx_ids = {i.bank_transaction_id for m in matches for i in m.items if i.bank_transaction_id}
    entry_ids = {i.accounting_entry_id for m in matches for i in m.items if i.accounting_entry_id}
    transactions = repo.transactions_with_bank(db, tx_ids)
    entries = repo.entries_with_bank(db, entry_ids)
    noms = repo.user_names(db, {m.valide_par_id for m in matches if m.valide_par_id})

    def premiere_date(match: ReconciliationMatch) -> date:
        return min(
            transactions[i.bank_transaction_id][0].date_operation
            for i in match.items
            if i.bank_transaction_id
        )

    matches.sort(key=lambda m: (premiere_date(m), m.code or "", m.id))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lettrage"
    sheet.append(LETTRAGE_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    total_debit = total_credit = Decimal("0.00")
    for match in matches:
        commun = [
            match.type,
            match.origine,
            match.score,
            noms.get(match.valide_par_id),
            _heure_locale(match.valide_le),
        ]
        operations = sorted(
            (transactions[i.bank_transaction_id] for i in match.items if i.bank_transaction_id),
            key=lambda row: (row[0].date_operation, row[0].id),
        )
        ecritures = sorted(
            (entries[i.accounting_entry_id] for i in match.items if i.accounting_entry_id),
            key=lambda row: (row[0].date_ecriture, row[0].id),
        )
        for tx, banque in operations:
            sheet.append(
                [match.code, "Banque", banque, tx.date_operation, None, tx.reference, tx.libelle,
                 tx.debit or None, tx.credit or None, *commun]
            )  # fmt: skip
            total_debit += tx.debit
            total_credit += tx.credit
        for entry, banque in ecritures:
            sheet.append(
                [match.code, "Sage", banque, entry.date_ecriture, entry.journal,
                 entry.numero_piece, entry.libelle, entry.debit or None, entry.credit or None,
                 *commun, entry.montant_devise]
            )  # fmt: skip
            total_debit += entry.debit
            total_credit += entry.credit
    sheet.append(
        [f"Total : {len(matches)} rapprochement{'s' if len(matches) > 1 else ''}", None, None,
         None, None, None, None, total_debit, total_credit]
    )  # fmt: skip
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        row[3].number_format = "DD/MM/YYYY"
        row[7].number_format = row[8].number_format = "#,##0.00"
        row[13].number_format = "DD/MM/YYYY HH:MM"
        row[14].number_format = "#,##0.00"
    for index, largeur in enumerate(_LARGEURS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = largeur
    sheet.freeze_panes = "A2"
    buffer = BytesIO()
    workbook.save(buffer)

    audit_service.log(
        db,
        user_id=acteur_id,
        action="export_lettrage",
        entite="company",
        entite_id=company.id,
        apres={
            "compte": bank_account_id,
            "du": date_from,
            "au": date_to,
            "rapprochements": len(matches),
        },
        ip=ip,
    )
    db.commit()
    banque = (
        repo.bank_code_of_account(db, bank_account_id) if bank_account_id is not None else "tous"
    )
    nom = f"lettrage-{banque}-{date_from or 'debut'}-{date_to or business_today()}.xlsx"
    return nom, buffer.getvalue(), len(matches)
